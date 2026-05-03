from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import check_workbook_integrity, resolve_static_pool_paths

DEFAULT_CANDIDATES = "deliveries/archive/milestones/milestone17_writeback_admission/milestone17_writeback_admission_package_v1.json"
DEFAULT_RUN_SUMMARY = "deliveries/archive/milestones/milestone17_writeback_admission/milestone17_writeback_admission_run_summary_v1.json"
DEFAULT_OUTPUT_JSON = "deliveries/archive/milestones/milestone18_l3_operationalization/milestone18_l3_operational_package_v1.json"
DEFAULT_OUTPUT_MD = "docs/03-执行与校验/Milestone 18-L3可消费池运营化复盘-v1.md"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M18 post-writeback operational package for L3 candidates.")
    parser.add_argument("--candidate-package", default=DEFAULT_CANDIDATES)
    parser.add_argument("--run-summary-file", default=DEFAULT_RUN_SUMMARY)
    parser.add_argument("--output-json", default=DEFAULT_OUTPUT_JSON)
    parser.add_argument("--output-md", default=DEFAULT_OUTPUT_MD)
    return parser


def _clean(value: object) -> str:
    return str(value or "").strip()


def _load_json(path: str) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write_json(path: str, payload: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _rows_by(workbook_path: Path, sheet_name: str, key_header: str) -> dict[str, dict[str, Any]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: dict[str, dict[str, Any]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        key = _clean(row.get(key_header))
        if key:
            rows[key] = row
    return rows


def _group_rows(workbook_path: Path, sheet_name: str, key_header: str) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        key = _clean(row.get(key_header))
        if key:
            grouped.setdefault(key, []).append(row)
    return grouped


def _is_ready(main: dict[str, Any], profile: dict[str, Any], shared: dict[str, Any], evidence_rows: list[dict[str, Any]]) -> tuple[bool, list[str]]:
    gaps: list[str] = []
    if _clean(main.get("静态潜客记录成熟度")) != "L3":
        gaps.append("main_not_l3")
    if _clean(profile.get("静态潜客记录成熟度")) != "L3":
        gaps.append("profile_not_l3")
    if _clean(shared.get("静态潜客记录成熟度")) != "L3":
        gaps.append("shared_not_l3")
    if _clean(profile.get("profile_status")) != "standard_ready":
        gaps.append("profile_status_not_standard_ready")
    for field in ("公司产品与服务概述", "商业模式概述", "admission_reason_summary", "validation_gap"):
        if not _clean(main.get(field)):
            gaps.append(f"main_missing_{field}")
    if len(evidence_rows) < 3:
        gaps.append("evidence_less_than_3")
    return not gaps, gaps


def _item_next_action(gaps: list[str]) -> str:
    if not gaps:
        return "进入首批 L3 可消费清单；可用于销售/研究/运营侧查看和后续 L2 补强。"
    return "暂不对外消费，先处理：" + ",".join(gaps)


def _render_markdown(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 18-L3可消费池运营化复盘-v1",
        "",
        "## 摘要",
        "",
        f"- 复核对象：`{summary['account_count']}`",
        f"- L3 可消费对象：`{summary['share_ready_count']}`",
        f"- 仍需治理对象：`{summary['needs_fix_count']}`",
        f"- 工作簿完整性：`{summary['workbook_integrity_ok']}`",
        f"- 写回状态：`promoted={summary['m17_promoted']} / skipped={summary['m17_skipped']}`",
        "",
        "## 首批 L3 可消费对象",
        "",
    ]
    for item in payload["items"]:
        lines.extend(
            [
                f"### {item['account_name']}（{item['account_id']}）",
                "",
                f"- 主线/画像：`{item['track']}` / `{item['persona']}`",
                f"- 产品服务：{item['product_summary']}",
                f"- 为什么值得看：{item['why_worth_review']}",
                f"- evidence 数：`{item['evidence_count']}`",
                f"- 队列状态：`{item['open_queue_count']} open / {item['resolved_queue_count']} resolved`",
                f"- 消费状态：`{item['share_status']}`",
                f"- 下一步：{item['next_action']}",
                "",
            ]
        )
    lines.extend(
        [
            "## 运营规则",
            "",
            "1. 本包不再执行写回，只作为写回后消费和运营清单。",
            "2. L3 对象可进入共享消费，但 L2 仍需补收入、利润、增长等财报口径字段。",
            "3. 后续扩容前继续使用 M14.3 preflight，避免未补字段候选直接进入执行批次。",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    pool = resolve_static_pool_paths()
    candidate_package = _load_json(args.candidate_package)
    run_summary = _load_json(args.run_summary_file)
    candidates = candidate_package.get("candidates") or []
    account_ids = [_clean(item.get("account_id")) for item in candidates if isinstance(item, dict) and _clean(item.get("account_id"))]

    main_by_id = _rows_by(pool["main"], "accounts_main", "account_id")
    profile_by_id = _rows_by(pool["profile"], "account_profiles", "account_id")
    shared_by_name = _rows_by(pool["main_shared"], "全量主表", "公司主体")
    evidence_by_id = _group_rows(pool["governance"], "evidence_log", "account_id")
    queue_by_id = _group_rows(pool["governance"], "review_queue", "account_id")
    integrity = check_workbook_integrity([pool["main"], pool["profile"], pool["governance"]], deep_scan=True)

    items = []
    for account_id in account_ids:
        main = main_by_id.get(account_id, {})
        profile = profile_by_id.get(account_id, {})
        account_name = _clean(main.get("account_canonical_name") or profile.get("account_canonical_name"))
        shared = shared_by_name.get(account_name, {})
        evidence_rows = evidence_by_id.get(account_id, [])
        queue_rows = queue_by_id.get(account_id, [])
        ready, gaps = _is_ready(main, profile, shared, evidence_rows)
        items.append(
            {
                "account_id": account_id,
                "account_name": account_name,
                "track": _clean(main.get("primary_track")),
                "persona": _clean(main.get("persona_tag")),
                "main_level": _clean(main.get("静态潜客记录成熟度")),
                "profile_level": _clean(profile.get("静态潜客记录成熟度")),
                "shared_level": _clean(shared.get("静态潜客记录成熟度")),
                "profile_status": _clean(profile.get("profile_status")),
                "product_summary": _clean(main.get("公司产品与服务概述")),
                "business_model_summary": _clean(main.get("商业模式概述")),
                "why_worth_review": _clean(main.get("admission_reason_summary")),
                "validation_gap": _clean(main.get("validation_gap")),
                "evidence_count": len(evidence_rows),
                "open_queue_count": len([row for row in queue_rows if _clean(row.get("status")) in {"", "open", "in_progress"}]),
                "resolved_queue_count": len([row for row in queue_rows if _clean(row.get("status")) == "resolved"]),
                "share_status": "l3_share_ready" if ready else "needs_post_writeback_fix",
                "quality_gaps": gaps,
                "next_action": _item_next_action(gaps),
            }
        )

    status_counts = Counter(item["share_status"] for item in items)
    promote_write = ((run_summary.get("promote") or {}).get("write_back") or {})
    payload = {
        "batch_id": "milestone18_l3_operational_package_v1",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_files": {
            "candidate_package": args.candidate_package,
            "run_summary_file": args.run_summary_file,
        },
        "static_pool_root": str(pool["root"]),
        "summary": {
            "account_count": len(items),
            "share_status_counts": dict(status_counts),
            "share_ready_count": int(status_counts.get("l3_share_ready") or 0),
            "needs_fix_count": int(status_counts.get("needs_post_writeback_fix") or 0),
            "workbook_integrity_ok": bool(integrity.get("ok")),
            "m17_promoted": int(promote_write.get("promoted") or 0),
            "m17_skipped": int(promote_write.get("skipped") or 0),
        },
        "items": items,
        "workbook_integrity": integrity,
    }
    _write_json(args.output_json, payload)
    _write_text(args.output_md, _render_markdown(payload))
    print(json.dumps({"output_json": args.output_json, "output_md": args.output_md, "summary": payload["summary"]}, ensure_ascii=False, indent=2))
    return 0 if payload["summary"]["needs_fix_count"] == 0 and payload["summary"]["workbook_integrity_ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
