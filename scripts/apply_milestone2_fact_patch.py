from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
ROOT = Path.home()
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MAIN_SHARED_XLSX = VAULT / "内部运营-静态潜客池-共享版.xlsx"
GOV_XLSX = VAULT / "治理与证据.xlsx"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Apply Milestone 2 fact patch into workbook fact sources.")
    parser.add_argument("--fact-patch-file", required=True, help="Fact patch JSON payload.")
    parser.add_argument("--output-file", help="Optional JSON summary output.")
    return parser


def _clean(value: object) -> str:
    return str(value or "").strip()


def backup_once(path: Path, suffix: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = path.with_name(f"{path.stem}.{suffix}_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    return str(backup_path)


def build_header_index(ws) -> dict[str, int]:
    headers = [cell.value for cell in ws[1]]
    return {str(value): idx + 1 for idx, value in enumerate(headers) if value}


def ensure_column(ws, header: str) -> int:
    header_index = build_header_index(ws)
    if header in header_index:
        return header_index[header]
    target_col = ws.max_column + 1
    ws.cell(1, target_col).value = header
    return target_col


def build_row_index(ws, key_header: str) -> tuple[dict[str, int], dict[str, int]]:
    header_index = build_header_index(ws)
    if key_header not in header_index:
        raise KeyError(f"missing key header: {key_header}")
    row_index: dict[str, int] = {}
    key_col = header_index[key_header]
    for row in range(2, ws.max_row + 1):
        value = _clean(ws.cell(row, key_col).value)
        if value:
            row_index[value] = row
    return header_index, row_index


def set_if_header(ws, header_index: dict[str, int], row: int, header: str, value: object) -> None:
    if header in header_index:
        ws.cell(row, header_index[header]).value = value


def upsert_evidence_row(ws, header_index: dict[str, int], account_id: str, row_data: dict[str, object]) -> bool:
    evidence_id = _clean(row_data.get("evidence_id"))
    if not evidence_id:
        raise ValueError(f"missing evidence_id for {account_id}")
    evidence_col = header_index["evidence_id"]
    target_row = None
    for row in range(2, ws.max_row + 1):
        if _clean(ws.cell(row, evidence_col).value) == evidence_id:
            target_row = row
            break
    if target_row is None:
        target_row = ws.max_row + 1
    set_if_header(ws, header_index, target_row, "evidence_id", evidence_id)
    set_if_header(ws, header_index, target_row, "account_id", account_id)
    set_if_header(ws, header_index, target_row, "evidence_type", row_data.get("evidence_type"))
    set_if_header(ws, header_index, target_row, "source_locator", row_data.get("source_locator"))
    set_if_header(ws, header_index, target_row, "evidence_strength", row_data.get("evidence_strength"))
    set_if_header(ws, header_index, target_row, "supports_dimension", row_data.get("supports_dimension"))
    set_if_header(ws, header_index, target_row, "summary", row_data.get("summary"))
    set_if_header(ws, header_index, target_row, "checked_by", "codex")
    set_if_header(ws, header_index, target_row, "checked_at", datetime.now().strftime("%Y-%m-%d"))
    set_if_header(ws, header_index, target_row, "field_name", row_data.get("field_name"))
    set_if_header(ws, header_index, target_row, "field_value", row_data.get("field_value"))
    return target_row == ws.max_row


def main() -> int:
    args = build_parser().parse_args()
    payload = json.loads(Path(args.fact_patch_file).read_text(encoding="utf-8"))

    backups = {
        "profile": backup_once(PROFILE_XLSX, "milestone2_fact_patch_backup"),
        "main_shared": backup_once(MAIN_SHARED_XLSX, "milestone2_fact_patch_backup"),
        "governance": backup_once(GOV_XLSX, "milestone2_fact_patch_backup"),
    }

    profile_wb = load_workbook(PROFILE_XLSX)
    profile_ws = profile_wb["account_profiles"]
    main_wb = load_workbook(MAIN_SHARED_XLSX)
    main_ws = main_wb["全量主表"]
    gov_wb = load_workbook(GOV_XLSX)
    evidence_ws = gov_wb["evidence_log"]

    for header in (
        "primary_source_types",
        "primary_source_refs",
        "official_source_count",
        "high_confidence_source_count",
        "产品与服务长摘录",
        "商业模式长摘录",
        "客户客群长摘录",
    ):
        ensure_column(profile_ws, header)
    profile_headers, profile_rows = build_row_index(profile_ws, "account_id")
    main_headers, main_rows = build_row_index(main_ws, "公司主体")
    evidence_headers, _ = build_row_index(evidence_ws, "evidence_id")

    profile_updates = 0
    main_updates = 0
    evidence_updates = 0

    for account in payload.get("accounts") or []:
        account_id = _clean(account.get("account_id"))
        account_name = _clean(account.get("account_name"))
        if account_id not in profile_rows:
            raise KeyError(f"profile account_id not found: {account_id}")
        if account_name not in main_rows:
            raise KeyError(f"main shared account name not found: {account_name}")

        profile_row = profile_rows[account_id]
        main_row = main_rows[account_name]

        for field, value in (account.get("profile_fields") or {}).items():
            set_if_header(profile_ws, profile_headers, profile_row, field, value)
        set_if_header(profile_ws, profile_headers, profile_row, "primary_source_types", account.get("primary_source_types"))
        set_if_header(profile_ws, profile_headers, profile_row, "primary_source_refs", account.get("primary_source_refs"))
        set_if_header(profile_ws, profile_headers, profile_row, "official_source_count", account.get("official_source_count"))
        set_if_header(profile_ws, profile_headers, profile_row, "high_confidence_source_count", account.get("high_confidence_source_count"))
        set_if_header(profile_ws, profile_headers, profile_row, "last_profiled_at", datetime.now().strftime("%Y-%m-%d"))
        profile_updates += 1

        for field, value in (account.get("main_fields") or {}).items():
            set_if_header(main_ws, main_headers, main_row, field, value)
        if "待验证项" in main_headers and account.get("profile_fields", {}).get("validation_gap"):
            set_if_header(main_ws, main_headers, main_row, "待验证项", account["profile_fields"]["validation_gap"])
        main_updates += 1

        for evidence in account.get("evidence_rows") or []:
            created = upsert_evidence_row(evidence_ws, evidence_headers, account_id, evidence)
            evidence_updates += 1 if created else 0

    profile_wb.save(PROFILE_XLSX)
    main_wb.save(MAIN_SHARED_XLSX)
    gov_wb.save(GOV_XLSX)

    result = {
        "batch_id": _clean(payload.get("batch_id")),
        "backups": backups,
        "profile_updates": profile_updates,
        "main_shared_updates": main_updates,
        "evidence_rows_created": evidence_updates,
        "account_count": len(payload.get("accounts") or []),
    }
    if args.output_file:
        Path(args.output_file).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
