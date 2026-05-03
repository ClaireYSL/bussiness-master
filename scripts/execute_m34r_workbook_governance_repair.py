from __future__ import annotations

import argparse
import json
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import check_workbook_integrity, resolve_static_pool_paths, workbook_write_lock
from shared.static_pool.promotion_writeback import backup_once, build_header_index


DEFAULT_ADMISSION = "deliveries/archive/milestones/milestone34r_workbook_governance_repair_admission/milestone34r_workbook_governance_repair_admission_package_v1.json"
DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone34r_workbook_governance_true_repair"
DEFAULT_REVIEW_MD = "docs/03-执行与校验/Milestone 34R-三表治理真实修复复盘-v1.md"
REQUIRED_MAIN_SHEET = "accounts_main"
REQUIRED_PROFILE_SHEET = "account_profiles"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Execute M34R workbook governance repair after explicit user confirmation.")
    parser.add_argument("--admission-package", default=DEFAULT_ADMISSION)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--review-md", default=DEFAULT_REVIEW_MD)
    parser.add_argument("--lock-timeout-seconds", type=float, default=5.0)
    return parser


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(value: object) -> str:
    return str(value or "").strip()


def _write_json(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str | Path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _count_main_duplicates(main_path: Path) -> dict[str, Any]:
    wb = load_workbook(main_path, read_only=True, data_only=True)
    try:
        ws = wb[REQUIRED_MAIN_SHEET]
        headers = [cell.value for cell in ws[1]]
        id_col = headers.index("account_id")
        level_col = headers.index("静态潜客记录成熟度")
        counts: Counter[str] = Counter()
        level_counts: Counter[str] = Counter()
        row_count = 0
        for values in ws.iter_rows(min_row=2, values_only=True):
            row_count += 1
            account_id = _clean(values[id_col])
            level = _clean(values[level_col])
            if account_id:
                counts[account_id] += 1
            if level:
                level_counts[level] += 1
        duplicate_groups = {account_id: count for account_id, count in counts.items() if count > 1}
        return {
            "row_count": row_count,
            "unique_account_ids": len(counts),
            "duplicate_group_count": len(duplicate_groups),
            "duplicate_rows_extra_count": sum(count - 1 for count in duplicate_groups.values()),
            "level_counts": dict(level_counts),
        }
    finally:
        wb.close()


def _count_profile(profile_path: Path) -> dict[str, Any]:
    wb = load_workbook(profile_path, read_only=True, data_only=True)
    try:
        ws = wb[REQUIRED_PROFILE_SHEET]
        headers = [cell.value for cell in ws[1]]
        id_col = headers.index("account_id")
        counts: Counter[str] = Counter()
        row_count = 0
        for values in ws.iter_rows(min_row=2, values_only=True):
            row_count += 1
            account_id = _clean(values[id_col])
            if account_id:
                counts[account_id] += 1
        duplicate_groups = {account_id: count for account_id, count in counts.items() if count > 1}
        return {
            "row_count": row_count,
            "unique_account_ids": len(counts),
            "duplicate_group_count": len(duplicate_groups),
        }
    finally:
        wb.close()


def _shared_counts(shared_path: Path) -> dict[str, Any]:
    wb = load_workbook(shared_path, read_only=True, data_only=True)
    try:
        counts: dict[str, Any] = {"sheet_names": list(wb.sheetnames)}
        if "全量主表" in wb.sheetnames:
            ws = wb["全量主表"]
            headers = [cell.value for cell in ws[1]]
            level_counts: Counter[str] = Counter()
            names: set[str] = set()
            if "静态潜客记录成熟度" in headers:
                level_col = headers.index("静态潜客记录成熟度")
            else:
                level_col = -1
            if "公司主体" in headers:
                name_col = headers.index("公司主体")
            else:
                name_col = -1
            row_count = 0
            for values in ws.iter_rows(min_row=2, values_only=True):
                row_count += 1
                if name_col >= 0 and _clean(values[name_col]):
                    names.add(_clean(values[name_col]))
                if level_col >= 0 and _clean(values[level_col]):
                    level_counts[_clean(values[level_col])] += 1
            counts.update({"row_count": row_count, "unique_company_names": len(names), "level_counts": dict(level_counts)})
        return counts
    finally:
        wb.close()


def _main_profile_gap(main_path: Path, profile_path: Path) -> dict[str, int]:
    main_ids: set[str] = set()
    profile_ids: set[str] = set()
    main_wb = load_workbook(main_path, read_only=True, data_only=True)
    profile_wb = load_workbook(profile_path, read_only=True, data_only=True)
    try:
        main_ws = main_wb[REQUIRED_MAIN_SHEET]
        profile_ws = profile_wb[REQUIRED_PROFILE_SHEET]
        main_headers = [cell.value for cell in main_ws[1]]
        profile_headers = [cell.value for cell in profile_ws[1]]
        main_id_col = main_headers.index("account_id")
        profile_id_col = profile_headers.index("account_id")
        for values in main_ws.iter_rows(min_row=2, values_only=True):
            if _clean(values[main_id_col]):
                main_ids.add(_clean(values[main_id_col]))
        for values in profile_ws.iter_rows(min_row=2, values_only=True):
            if _clean(values[profile_id_col]):
                profile_ids.add(_clean(values[profile_id_col]))
        return {
            "main_not_in_profile_count": len(main_ids - profile_ids),
            "profile_not_in_main_count": len(profile_ids - main_ids),
        }
    finally:
        main_wb.close()
        profile_wb.close()


def _preflight(admission: dict[str, Any], pool: dict[str, Path]) -> dict[str, Any]:
    summary = admission.get("summary") or {}
    main = _count_main_duplicates(pool["main"])
    profile = _count_profile(pool["profile"])
    shared = _shared_counts(pool["main_shared"])
    gaps = _main_profile_gap(pool["main"], pool["profile"])
    checks = {
        "main_duplicate_group_matches_admission": main["duplicate_group_count"] == int(summary.get("duplicate_group_count") or -1),
        "main_duplicate_extra_matches_admission": main["duplicate_rows_extra_count"] == int(summary.get("duplicate_rows_to_remove_count") or -1),
        "profile_gap_matches_admission": gaps["main_not_in_profile_count"] == int(summary.get("missing_profile_patch_count") or -1),
        "shared_preview_exists": Path(summary.get("shared_preview_xlsx") or "").exists(),
    }
    return {
        "main": main,
        "profile": profile,
        "shared": shared,
        "main_profile_gap": gaps,
        "checks": checks,
        "ok": all(checks.values()),
    }


def _delete_duplicate_rows(main_path: Path, duplicate_plan: list[dict[str, Any]]) -> int:
    rows_to_delete: set[int] = set()
    for item in duplicate_plan:
        for row in item.get("recommended_remove") or []:
            excel_row = int(row.get("excel_row") or 0)
            if excel_row > 1:
                rows_to_delete.add(excel_row)
    wb = load_workbook(main_path)
    ws = wb[REQUIRED_MAIN_SHEET]
    for excel_row in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(excel_row, 1)
    wb.save(main_path)
    wb.close()
    return len(rows_to_delete)


def _append_profile_rows(profile_path: Path, profile_patch: dict[str, Any]) -> int:
    accounts = profile_patch.get("accounts") or []
    wb = load_workbook(profile_path)
    ws = wb[REQUIRED_PROFILE_SHEET]
    header_index = build_header_index(ws)
    headers = [None] * len(header_index)
    for header, idx in header_index.items():
        headers[idx - 1] = header
    existing_ids = {
        _clean(ws.cell(row, header_index["account_id"]).value)
        for row in range(2, ws.max_row + 1)
        if _clean(ws.cell(row, header_index["account_id"]).value)
    }
    appended = 0
    for account in accounts:
        account_id = _clean(account.get("account_id"))
        if not account_id or account_id in existing_ids:
            continue
        ws.append([account.get(header, "") for header in headers])
        existing_ids.add(account_id)
        appended += 1
    wb.save(profile_path)
    wb.close()
    return appended


def _replace_shared_workbook(shared_path: Path, preview_path: Path) -> dict[str, Any]:
    shutil.copy2(preview_path, shared_path)
    return _shared_counts(shared_path)


def _render_md(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 34R-三表治理真实修复复盘-v1",
        "",
        "## 摘要",
        "",
        f"- 执行状态：`{payload['status']}`",
        f"- 删除主表重复行：`{summary['duplicate_rows_deleted']}`",
        f"- 补写档案库 profile stub：`{summary['profile_rows_appended']}`",
        f"- 共享版重建行数：`{summary['shared_rebuilt_row_count']}`",
        f"- 修复后主表行数：`{summary['post_main_row_count']}`，唯一 account_id：`{summary['post_main_unique_account_ids']}`，重复组：`{summary['post_main_duplicate_group_count']}`",
        f"- 修复后档案库行数：`{summary['post_profile_row_count']}`，唯一 account_id：`{summary['post_profile_unique_account_ids']}`",
        f"- 主表有档案库无：`{summary['post_main_not_in_profile_count']}`",
        f"- 工作簿完整性：`{summary['post_workbook_integrity_ok']}`",
        f"- 锁状态：`acquired={summary['workbook_lock_acquired']}`，wait_seconds=`{summary['workbook_lock_wait_seconds']}`",
        "",
        "## 边界",
        "",
        "- 本轮只治理主表、档案库、共享版三表口径。",
        "- 未写入正式知识资产、未修改画像注册表。",
        "- 已在真实写入前创建 M34R 专用备份。",
    ]
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    pool = resolve_static_pool_paths()
    admission = _load_json(args.admission_package)
    summary = admission.get("summary") or {}
    preview_path = Path(summary.get("shared_preview_xlsx") or "")
    output_dir = Path(args.output_dir)

    preflight = _preflight(admission, pool)
    if not preflight["ok"]:
        payload = {
            "batch_id": "milestone34r_workbook_governance_true_repair_v1",
            "generated_at": _now(),
            "status": "blocked_preflight_failed",
            "static_pool_root": str(pool["root"]),
            "preflight": preflight,
        }
        _write_json(output_dir / "milestone34r_workbook_governance_true_repair_summary_v1.json", payload)
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 1

    with workbook_write_lock(timeout_seconds=args.lock_timeout_seconds) as lock_meta:
        backups = {
            "main": backup_once(pool["main"], "m34r_repair_backup"),
            "profile": backup_once(pool["profile"], "m34r_repair_backup"),
            "main_shared": backup_once(pool["main_shared"], "m34r_repair_backup"),
            "governance": backup_once(pool["governance"], "m34r_repair_backup"),
        }
        duplicate_rows_deleted = _delete_duplicate_rows(pool["main"], admission.get("duplicate_resolution_plan") or [])
        profile_rows_appended = _append_profile_rows(pool["profile"], admission.get("missing_profile_patch") or {})
        rebuilt_shared = _replace_shared_workbook(pool["main_shared"], preview_path)
        post_integrity = check_workbook_integrity([pool["main"], pool["profile"], pool["main_shared"], pool["governance"]], deep_scan=True)

    post_main = _count_main_duplicates(pool["main"])
    post_profile = _count_profile(pool["profile"])
    post_gaps = _main_profile_gap(pool["main"], pool["profile"])
    post_shared = _shared_counts(pool["main_shared"])
    result_summary = {
        "duplicate_rows_deleted": duplicate_rows_deleted,
        "profile_rows_appended": profile_rows_appended,
        "shared_rebuilt_row_count": int(rebuilt_shared.get("row_count") or 0),
        "post_main_row_count": post_main["row_count"],
        "post_main_unique_account_ids": post_main["unique_account_ids"],
        "post_main_duplicate_group_count": post_main["duplicate_group_count"],
        "post_profile_row_count": post_profile["row_count"],
        "post_profile_unique_account_ids": post_profile["unique_account_ids"],
        "post_profile_duplicate_group_count": post_profile["duplicate_group_count"],
        "post_main_not_in_profile_count": post_gaps["main_not_in_profile_count"],
        "post_profile_not_in_main_count": post_gaps["profile_not_in_main_count"],
        "post_shared_row_count": post_shared.get("row_count", 0),
        "post_shared_unique_company_names": post_shared.get("unique_company_names", 0),
        "post_workbook_integrity_ok": bool(post_integrity.get("ok")),
        "workbook_lock_acquired": bool(lock_meta.get("acquired")),
        "workbook_lock_wait_seconds": lock_meta.get("wait_seconds"),
        "true_repair_executed": True,
        "knowledge_asset_write_enabled": False,
        "persona_registry_write_enabled": False,
    }
    payload = {
        "batch_id": "milestone34r_workbook_governance_true_repair_v1",
        "generated_at": _now(),
        "status": "success",
        "static_pool_root": str(pool["root"]),
        "preflight": preflight,
        "backups": backups,
        "summary": result_summary,
        "post_counts": {
            "main": post_main,
            "profile": post_profile,
            "shared": post_shared,
            "main_profile_gap": post_gaps,
        },
        "workbook_integrity": post_integrity,
    }
    _write_json(output_dir / "milestone34r_workbook_governance_true_repair_summary_v1.json", payload)
    _write_text(args.review_md, _render_md(payload))
    print(json.dumps({"output_json": str(output_dir / "milestone34r_workbook_governance_true_repair_summary_v1.json"), "review_md": args.review_md, "summary": result_summary}, ensure_ascii=False, indent=2))

    ok = (
        duplicate_rows_deleted == int(summary.get("duplicate_rows_to_remove_count") or -1)
        and profile_rows_appended == int(summary.get("missing_profile_patch_count") or -1)
        and result_summary["post_main_duplicate_group_count"] == 0
        and result_summary["post_main_not_in_profile_count"] == 0
        and result_summary["post_workbook_integrity_ok"]
    )
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
