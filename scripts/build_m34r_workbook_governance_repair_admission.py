from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import check_workbook_integrity, resolve_static_pool_paths


DEFAULT_M33 = "deliveries/archive/milestones/milestone33r_workbook_governance_audit/milestone33r_workbook_governance_audit_package_v1.json"
DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone34r_workbook_governance_repair_admission"
DEFAULT_REVIEW_MD = "docs/03-执行与校验/Milestone 34R-三表治理安全修复准入包-v1.md"

HIGH_QUALITY_LEVELS = {"L1", "L2", "L3"}
LEVEL_RANK = {"L1": 5, "L2": 4, "L3": 3, "L4": 2, "L5": 1}
CORE_FIELDS = [
    "account_canonical_name",
    "primary_track",
    "persona_tag",
    "admission_reason_summary",
    "公司产品与服务概述",
    "商业模式概述",
    "validation_gap",
    "knowledge_asset_refs",
]
PROFILE_HEADERS = [
    "account_id",
    "account_canonical_name",
    "brand_name",
    "group_name",
    "primary_track",
    "industry_l1",
    "industry_l2",
    "business_model",
    "persona_tag",
    "management_persona_tags",
    "static_maturity_level",
    "static_priority",
    "信息扎实度",
    "ICP匹配概率",
    "静态潜客记录成熟度",
    "admission_reason_summary",
    "current_business_problem",
    "primary_jtbd",
    "secondary_jtbd",
    "validation_gap",
    "archive_status",
    "公司产品与服务概述",
    "商业模式概述",
    "核心客户客群",
    "收入规模区间",
    "利润状态概述",
    "营收增长概述",
    "已上线系统概况",
    "数字化项目动态",
    "相似客户线索",
    "主要竞品概述",
    "招聘代表岗位",
    "近一年重大事件",
    "产品与服务长摘录",
    "商业模式长摘录",
    "客户客群长摘录",
    "系统与数字化长摘录",
    "重大事件长摘录",
    "primary_source_types",
    "primary_source_refs",
    "official_source_count",
    "high_confidence_source_count",
    "last_profiled_at",
    "profile_owner",
    "profile_status",
    "knowledge_asset_refs",
    "talk_track_refs",
    "representative_for_persona",
    "representative_for_track",
    "should_have_focus_note",
    "focus_note_path",
    "share_status",
    "share_batch_id",
    "share_last_marked_at",
    "share_note",
    "secondary_persona_tags",
]
SHARED_HEADERS = [
    "公司主体",
    "品牌名",
    "主线",
    "业务形态画像",
    "管理诉求画像",
    "信息扎实度",
    "ICP匹配概率",
    "静态潜客记录成熟度",
    "公司产品与服务概述",
    "商业模式概述",
    "核心客户客群",
    "收入规模区间",
    "营收增长概述",
    "已上线系统概况",
    "数字化项目动态",
    "近一年重大事件",
    "一话入池理由",
    "主要知识资产引用",
    "主要切入话术引用",
    "待验证项",
]


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M34R workbook governance repair admission package without mutating workbooks.")
    parser.add_argument("--m33-package", default=DEFAULT_M33)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--review-md", default=DEFAULT_REVIEW_MD)
    return parser


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _today() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def _clean(value: object) -> str:
    return str(value or "").strip()


def _write_json(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str | Path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _rows(path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = dict(zip(headers, values))
            row["_excel_row"] = row_number
            rows.append(row)
        return [str(header) for header in headers if header], rows
    finally:
        wb.close()


def _profile_by_id(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {_clean(row.get("account_id")): row for row in rows if _clean(row.get("account_id"))}


def _row_quality_score(row: dict[str, Any], profile_by_id: dict[str, dict[str, Any]]) -> tuple[int, int, int, int]:
    level = _clean(row.get("静态潜客记录成熟度"))
    filled_core = sum(1 for field in CORE_FIELDS if _clean(row.get(field)))
    has_profile = 1 if _clean(row.get("account_id")) in profile_by_id else 0
    row_number = int(row.get("_excel_row") or 999999)
    return (LEVEL_RANK.get(level, 0), filled_core, has_profile, -row_number)


def _duplicate_resolution_plan(main_rows: list[dict[str, Any]], profile_by_id: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], set[int]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in main_rows:
        account_id = _clean(row.get("account_id"))
        if account_id:
            grouped[account_id].append(row)
    plans: list[dict[str, Any]] = []
    remove_rows: set[int] = set()
    for account_id, rows in sorted(grouped.items()):
        if len(rows) <= 1:
            continue
        ranked = sorted(rows, key=lambda row: _row_quality_score(row, profile_by_id), reverse=True)
        keep = ranked[0]
        remove = ranked[1:]
        for row in remove:
            remove_rows.add(int(row["_excel_row"]))
        plans.append(
            {
                "account_id": account_id,
                "duplicate_count": len(rows),
                "recommended_keep": _row_brief(keep),
                "recommended_remove": [_row_brief(row) for row in remove],
                "rule": "保留层级更高、核心字段更完整、且有档案承接的行；若仍并列，保留靠前行。",
                "writeback_status": "not_executed",
            }
        )
    return plans, remove_rows


def _row_brief(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "excel_row": int(row.get("_excel_row") or 0),
        "account_id": _clean(row.get("account_id")),
        "account_name": _clean(row.get("account_canonical_name")),
        "level": _clean(row.get("静态潜客记录成熟度")),
        "track": _clean(row.get("primary_track")),
        "persona": _clean(row.get("persona_tag")),
        "filled_core_fields": sum(1 for field in CORE_FIELDS if _clean(row.get(field))),
    }


def _deduped_main_rows(main_rows: list[dict[str, Any]], remove_rows: set[int]) -> list[dict[str, Any]]:
    return [row for row in main_rows if int(row.get("_excel_row") or 0) not in remove_rows]


def _missing_profile_patch(deduped_main: list[dict[str, Any]], profile_by_id: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    patch: list[dict[str, Any]] = []
    today = _today()
    for row in deduped_main:
        account_id = _clean(row.get("account_id"))
        if not account_id or account_id in profile_by_id:
            continue
        level = _clean(row.get("静态潜客记录成熟度"))
        profile_row = {header: "" for header in PROFILE_HEADERS}
        copy_fields = [
            "account_id",
            "account_canonical_name",
            "brand_name",
            "group_name",
            "primary_track",
            "industry_l1",
            "industry_l2",
            "business_model",
            "persona_tag",
            "static_priority",
            "信息扎实度",
            "ICP匹配概率",
            "静态潜客记录成熟度",
            "admission_reason_summary",
            "current_business_problem",
            "primary_jtbd",
            "secondary_jtbd",
            "validation_gap",
            "公司产品与服务概述",
            "商业模式概述",
            "核心客户客群",
            "收入规模区间",
            "利润状态概述",
            "营收增长概述",
            "已上线系统概况",
            "数字化项目动态",
            "相似客户线索",
            "主要竞品概述",
            "招聘代表岗位",
            "近一年重大事件",
            "knowledge_asset_refs",
        ]
        for field in copy_fields:
            if field in profile_row:
                profile_row[field] = _clean(row.get(field))
        profile_row["static_maturity_level"] = level
        profile_row["archive_status"] = "profile_stub_from_main"
        profile_row["profile_status"] = "stub_needs_enrichment" if level not in HIGH_QUALITY_LEVELS else "standard_ready"
        profile_row["last_profiled_at"] = today
        profile_row["profile_owner"] = "codex"
        profile_row["share_status"] = "pending_shared_projection"
        profile_row["share_note"] = "M34R 候选补档案行；真实写入前需用户确认。"
        patch.append(profile_row)
    return patch


def _shared_row(main: dict[str, Any], profile: dict[str, Any] | None) -> dict[str, Any]:
    profile = profile or {}
    return {
        "公司主体": _clean(main.get("account_canonical_name")),
        "品牌名": _clean(main.get("brand_name")),
        "主线": _clean(main.get("primary_track")),
        "业务形态画像": _clean(main.get("persona_tag")),
        "管理诉求画像": _clean(profile.get("management_persona_tags") or profile.get("secondary_persona_tags")),
        "信息扎实度": _clean(main.get("信息扎实度")),
        "ICP匹配概率": _clean(main.get("ICP匹配概率")),
        "静态潜客记录成熟度": _clean(main.get("静态潜客记录成熟度")),
        "公司产品与服务概述": _clean(main.get("公司产品与服务概述") or profile.get("公司产品与服务概述")),
        "商业模式概述": _clean(main.get("商业模式概述") or profile.get("商业模式概述")),
        "核心客户客群": _clean(main.get("核心客户客群") or profile.get("核心客户客群")),
        "收入规模区间": _clean(main.get("收入规模区间") or profile.get("收入规模区间")),
        "营收增长概述": _clean(main.get("营收增长概述") or profile.get("营收增长概述")),
        "已上线系统概况": _clean(main.get("已上线系统概况") or profile.get("已上线系统概况")),
        "数字化项目动态": _clean(main.get("数字化项目动态") or profile.get("数字化项目动态")),
        "近一年重大事件": _clean(main.get("近一年重大事件") or profile.get("近一年重大事件")),
        "一话入池理由": _clean(main.get("admission_reason_summary") or profile.get("admission_reason_summary")),
        "主要知识资产引用": _clean(main.get("knowledge_asset_refs") or profile.get("knowledge_asset_refs")),
        "主要切入话术引用": _clean(profile.get("talk_track_refs")),
        "待验证项": _clean(main.get("validation_gap") or profile.get("validation_gap")),
    }


def _shared_projection(deduped_main: list[dict[str, Any]], profile_by_id: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for main in deduped_main:
        account_id = _clean(main.get("account_id"))
        rows.append(_shared_row(main, profile_by_id.get(account_id)))
    return rows


def _write_shared_preview_xlsx(path: Path, shared_rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "全量主表"
    ws.append(SHARED_HEADERS)
    for row in shared_rows:
        ws.append([row.get(header, "") for header in SHARED_HEADERS])

    ws_summary = wb.create_sheet("主线汇总")
    ws_summary.append(["主线", "账户数", "L1", "L2", "L3", "L4", "L5"])
    tracks = sorted({_clean(row.get("主线")) for row in shared_rows if _clean(row.get("主线"))})
    for track in tracks:
        rows = [row for row in shared_rows if _clean(row.get("主线")) == track]
        levels = Counter(_clean(row.get("静态潜客记录成熟度")) for row in rows)
        ws_summary.append([track, len(rows), levels.get("L1", 0), levels.get("L2", 0), levels.get("L3", 0), levels.get("L4", 0), levels.get("L5", 0)])

    high = wb.create_sheet("高质量层")
    high.append(SHARED_HEADERS)
    for row in shared_rows:
        if _clean(row.get("静态潜客记录成熟度")) in HIGH_QUALITY_LEVELS:
            high.append([row.get(header, "") for header in SHARED_HEADERS])

    ext = wb.create_sheet("L4_L5扩展层")
    ext.append(SHARED_HEADERS)
    for row in shared_rows:
        if _clean(row.get("静态潜客记录成熟度")) not in HIGH_QUALITY_LEVELS:
            ext.append([row.get(header, "") for header in SHARED_HEADERS])

    boundary = wb.create_sheet("边界主体")
    boundary.append(SHARED_HEADERS)

    persona = wb.create_sheet("画像汇总")
    persona.append(["业务形态画像", "管理诉求画像", "账户数", "代表账户"])
    grouped: dict[tuple[str, str], list[str]] = defaultdict(list)
    for row in shared_rows:
        grouped[(_clean(row.get("业务形态画像")), _clean(row.get("管理诉求画像")))].append(_clean(row.get("公司主体")))
    for (biz_persona, mgmt_persona), names in sorted(grouped.items()):
        persona.append([biz_persona, mgmt_persona, len(names), "、".join(names[:5])])

    wb.save(path)


def _render_md(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 34R-三表治理安全修复准入包-v1",
        "",
        "## 摘要",
        "",
        f"- 重复 account_id 组：`{summary['duplicate_group_count']}`，建议移除重复行：`{summary['duplicate_rows_to_remove_count']}`",
        f"- 缺档案 profile patch：`{summary['missing_profile_patch_count']}`",
        f"- 共享版重建预览行数：`{summary['shared_projection_row_count']}`",
        f"- 预期高质量层：`{summary['expected_high_quality_count']}`",
        f"- 预期扩展层：`{summary['expected_extended_count']}`",
        f"- 工作簿完整性：`{summary['workbook_integrity_ok']}`",
        "",
        "## 结论",
        "",
        "- M34R 已形成可执行前修复准入材料，但没有真实修改工作簿。",
        "- 推荐顺序：先去重，再补 profile，再重建共享版。",
        "- 真实修复涉及删除重复行和覆盖共享版，必须单独确认。",
        "",
        "## 修复顺序",
        "",
        "1. 备份主表、档案库、共享版、治理与证据。",
        "2. 按 duplicate resolution plan 处理 7 组重复 account_id。",
        "3. 按 missing profile patch 补建 78 条 profile stub。",
        "4. 用共享版预览重建共享版，而不是手工逐行修。",
        "5. 重跑 M33R 审计，要求重复和缺口显著下降。",
    ]
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    pool = resolve_static_pool_paths()
    m33 = _load_json(args.m33_package)
    main_headers, main_rows = _rows(pool["main"], "accounts_main")
    _profile_headers, profile_rows = _rows(pool["profile"], "account_profiles")
    profile_by_id = { _clean(row.get("account_id")): row for row in profile_rows if _clean(row.get("account_id")) }
    integrity = check_workbook_integrity([pool["main"], pool["profile"], pool["main_shared"], pool["governance"]], deep_scan=True)

    duplicate_plan, remove_rows = _duplicate_resolution_plan(main_rows, profile_by_id)
    deduped_main = _deduped_main_rows(main_rows, remove_rows)
    missing_profile_patch = _missing_profile_patch(deduped_main, profile_by_id)
    shared_rows = _shared_projection(deduped_main, profile_by_id)
    output_dir = Path(args.output_dir)
    preview_xlsx = output_dir / "milestone34r_shared_view_rebuild_preview_v1.xlsx"
    _write_shared_preview_xlsx(preview_xlsx, shared_rows)

    level_counts = Counter(_clean(row.get("静态潜客记录成熟度")) for row in deduped_main if _clean(row.get("静态潜客记录成熟度")))
    expected_high = sum(count for level, count in level_counts.items() if level in HIGH_QUALITY_LEVELS)
    expected_extended = len(deduped_main) - expected_high
    repair_sequence = {
        "sequence_id": "milestone34r_repair_sequence_v1",
        "mode": "admission_only_no_write",
        "steps": [
            {"step": 1, "action": "backup_workbooks", "requires_confirmation": True},
            {"step": 2, "action": "resolve_duplicate_account_ids", "candidate_count": len(duplicate_plan), "requires_confirmation": True},
            {"step": 3, "action": "append_missing_profile_rows", "candidate_count": len(missing_profile_patch), "requires_confirmation": True},
            {"step": 4, "action": "rebuild_shared_view_from_projection", "candidate_count": len(shared_rows), "requires_confirmation": True},
            {"step": 5, "action": "rerun_m33r_audit_and_integrity", "requires_confirmation": False},
        ],
    }
    summary = {
        "source_m33_status": m33.get("summary", {}),
        "duplicate_group_count": len(duplicate_plan),
        "duplicate_rows_to_remove_count": len(remove_rows),
        "deduped_main_row_count": len(deduped_main),
        "missing_profile_patch_count": len(missing_profile_patch),
        "shared_projection_row_count": len(shared_rows),
        "expected_level_counts_after_dedupe": dict(level_counts),
        "expected_high_quality_count": expected_high,
        "expected_extended_count": expected_extended,
        "shared_preview_xlsx": str(preview_xlsx),
        "workbook_integrity_ok": bool(integrity.get("ok")),
        "true_writeback_enabled": False,
        "destructive_action_enabled": False,
        "shared_overwrite_enabled": False,
    }
    payload = {
        "batch_id": "milestone34r_workbook_governance_repair_admission_package_v1",
        "generated_at": _now(),
        "static_pool_root": str(pool["root"]),
        "policy": {
            "true_writeback_enabled": False,
            "destructive_action_enabled": False,
            "shared_overwrite_enabled": False,
            "requires_user_confirmation_for_real_repair": True,
        },
        "summary": summary,
        "duplicate_resolution_plan": duplicate_plan,
        "missing_profile_patch": {
            "batch_id": "milestone34r_missing_profile_patch_v1",
            "mode": "candidate_only_no_write",
            "accounts": missing_profile_patch,
        },
        "shared_view_projection": {
            "batch_id": "milestone34r_shared_view_projection_v1",
            "mode": "preview_only_no_write",
            "headers": SHARED_HEADERS,
            "rows": shared_rows,
        },
        "repair_sequence": repair_sequence,
        "workbook_integrity": integrity,
    }
    _write_json(output_dir / "milestone34r_workbook_governance_repair_admission_package_v1.json", payload)
    _write_json(output_dir / "milestone34r_duplicate_resolution_plan_v1.json", {"batch_id": "milestone34r_duplicate_resolution_plan_v1", "items": duplicate_plan})
    _write_json(output_dir / "milestone34r_missing_profile_patch_v1.json", payload["missing_profile_patch"])
    _write_json(output_dir / "milestone34r_shared_view_projection_v1.json", payload["shared_view_projection"])
    _write_json(output_dir / "milestone34r_repair_sequence_v1.json", repair_sequence)
    _write_text(args.review_md, _render_md(payload))
    print(json.dumps({"output_json": str(output_dir / "milestone34r_workbook_governance_repair_admission_package_v1.json"), "review_md": args.review_md, "summary": summary}, ensure_ascii=False, indent=2))
    ok = (
        summary["duplicate_group_count"] == 7
        and summary["missing_profile_patch_count"] == 78
        and summary["shared_projection_row_count"] == summary["deduped_main_row_count"]
        and summary["workbook_integrity_ok"]
        and not summary["true_writeback_enabled"]
        and not summary["destructive_action_enabled"]
    )
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
