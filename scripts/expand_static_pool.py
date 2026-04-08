from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
ROOT = Path.home()
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from scripts.expand_l5_consumer_personas_20260331 import run_expand_provider

VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"
TRACK_PERSONA_XLSX = VAULT / "主线与画像注册表.xlsx"
RETAIL_PROVIDER = WORKSPACE / "scripts/expand_l5_consumer_personas_20260331.py"

SUPPORTED_PERSONA_PROVIDERS = {
    "retail_brand_beauty": str(RETAIL_PROVIDER),
    "retail_brand_maternal_pet": str(RETAIL_PROVIDER),
    "retail_fashion_group": str(RETAIL_PROVIDER),
    "retail_multi_store": str(RETAIL_PROVIDER),
    "retail_high_sku_brand": str(RETAIL_PROVIDER),
}

TRACK_NAME_TO_ID = {
    "零售消费": "retail_consumer",
    "跨境电商": "cross_border_ecommerce",
    "先进制造": "advanced_manufacturing",
}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Route expand requests by active track/persona provider coverage.")
    parser.add_argument("--track", required=True, help="Track name, for example 零售消费.")
    parser.add_argument("--persona-id", help="Optional primary persona to route.")
    parser.add_argument("--candidate-source", default="", help="Optional source description for this batch.")
    parser.add_argument("--account-list-file", help="Optional candidate list file for rule-driven fallback.")
    parser.add_argument("--limit", type=int, default=20, help="Maximum number of candidates to return or execute for this batch.")
    parser.add_argument("--output-file", required=True, help="Where to write the expand routing result package.")
    parser.add_argument("--report-only", action="store_true", help="Only build routing report.")
    parser.add_argument("--write-back", action="store_true", help="Execute the matched provider when available.")
    return parser


def _clean(value: object) -> str:
    return str(value or "").strip()


def load_active_personas_by_track() -> dict[str, list[str]]:
    wb = load_workbook(TRACK_PERSONA_XLSX, read_only=True, data_only=True)
    ws = wb["personas"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {header: pos for pos, header in enumerate(headers)}
    result: dict[str, list[str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        persona_id = _clean(row[idx["persona_id"]])
        track_id = _clean(row[idx["track_id"]])
        status = _clean(row[idx["status"]])
        if not persona_id or status != "active" or persona_id.startswith("mgmt_"):
            continue
        result.setdefault(track_id, []).append(persona_id)
    return result


def load_candidate_list(path: str | None) -> list[str]:
    if not path:
        return []
    payload = Path(path).read_text(encoding="utf-8").strip()
    if not payload:
        return []
    if payload.startswith("["):
        values = json.loads(payload)
        return [str(item).strip() for item in values if str(item).strip()]
    return [line.strip() for line in payload.splitlines() if line.strip()]


def execute_provider(persona_id: str, limit: int, report_only: bool, output_file: str) -> dict[str, object]:
    if persona_id not in SUPPORTED_PERSONA_PROVIDERS:
        raise RuntimeError(f"当前 persona `{persona_id}` 没有可执行 provider。")
    if SUPPORTED_PERSONA_PROVIDERS[persona_id] != str(RETAIL_PROVIDER):
        raise RuntimeError(f"当前 provider `{SUPPORTED_PERSONA_PROVIDERS[persona_id]}` 还未接入统一 expand 执行。")
    return run_expand_provider(
        persona_ids=[persona_id],
        limit=limit,
        report_only=report_only,
        output_file=output_file,
    )


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    track_id = TRACK_NAME_TO_ID.get(args.track, args.track)
    active_by_track = load_active_personas_by_track()
    active_personas = sorted(active_by_track.get(track_id, []))
    persona_id = args.persona_id or (active_personas[0] if len(active_personas) == 1 else "")
    provider = SUPPORTED_PERSONA_PROVIDERS.get(persona_id, "")
    candidates = load_candidate_list(args.account_list_file)
    supported = bool(provider)
    provider_executed = False
    provider_result: dict[str, object] | None = None

    payload = {
        "batch_id": Path(args.output_file).stem,
        "track": args.track,
        "track_id": track_id,
        "persona_id": persona_id,
        "candidate_source": args.candidate_source,
        "report_only": args.report_only or not supported,
        "write_back": bool(args.write_back and supported and not args.report_only),
        "routing": {
            "supported_provider": supported,
            "provider_script": provider,
            "active_personas_for_track": active_personas,
            "mode": "provider_script" if supported else "rule_driven_fallback",
        },
        "candidate_list": candidates,
        "formal_candidate_count": 0,
        "observation_count": 0,
        "unsupported_provider": not supported,
        "provider_executed": False,
        "provider_result": None,
        "next_step": "",
    }
    if supported:
        if args.report_only or not args.write_back:
            payload["next_step"] = f"当前可接专题 provider：{provider}。如需真实扩池，使用 `--write-back` 进入统一 expand 执行。"
        else:
            provider_result = execute_provider(
                persona_id=persona_id,
                limit=args.limit,
                report_only=False,
                output_file=args.output_file,
            )
            provider_executed = True
            payload["provider_executed"] = True
            payload["provider_result"] = provider_result
            summary = provider_result.get("summary") or {}
            candidate_counter = summary.get("candidate_type_counter") or {}
            payload["formal_candidate_count"] = int(candidate_counter.get("formal_candidate", 0))
            payload["observation_count"] = int(candidate_counter.get("observation", 0))
            payload["next_step"] = "已通过统一 expand 入口执行匹配 provider，并产出 provider 结果包。"
    else:
        payload["next_step"] = "当前无匹配专题脚本，进入规则驱动模式：先补最小事实、最小 evidence，再交 enrich/promote 链路处理。"

    if not provider_executed:
        output_path = Path(args.output_file)
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "output_file": args.output_file,
                "routing": payload["routing"],
                "candidate_count": len(candidates),
                "provider_executed": payload["provider_executed"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
