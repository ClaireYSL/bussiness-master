from __future__ import annotations

import argparse
import json
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
import re
import sys
import tempfile

import akshare as ak
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from shared.static_pool import classify_l5_candidate, render_validation_summary, resolve_static_pool_paths


TODAY = "2026-03-31"
TARGET_NEW_L5 = 220

POOL_DIR = resolve_static_pool_paths()["root"]
MAIN_WB = POOL_DIR / "静态潜客主表.xlsx"
MAIN_SHARED_WB = POOL_DIR / "内部运营-静态潜客池-共享版.xlsx"
PROFILE_WB = POOL_DIR / "潜客档案库.xlsx"
GOV_WB = POOL_DIR / "治理与证据.xlsx"
HOME_MD = POOL_DIR / "潜客池-首页.md"
TOTAL_MD = POOL_DIR / "05-汇总与状态/01-总览/内部总览-静态潜客池-总池汇总.md"
MILESTONE_MD = POOL_DIR / "05-汇总与状态/01-总览/内部状态-静态潜客池-Milestone状态总览.md"
EXEC_MD = ROOT / "docs/03-执行与校验/静态潜客池-消费品新画像扩池专项-v1.md"
MEMORY_MD = ROOT / "memory/2026-03-31.md"


@dataclass(frozen=True)
class Candidate:
    code: str
    name: str
    board_name: str
    persona_tag: str
    board_code: str


BOARD_PERSONA = [
    ("品牌化妆品", "retail_brand_beauty"),
    ("化妆品", "retail_brand_beauty"),
    ("美容护理", "retail_brand_beauty"),
    ("洗护用品", "retail_brand_beauty"),
    ("个护用品", "retail_brand_beauty"),
    ("个护小家电", "retail_brand_beauty"),
    ("医疗美容", "retail_brand_beauty"),
    ("宠物食品", "retail_brand_maternal_pet"),
    ("小家电", "retail_brand_maternal_pet"),
    ("家居用品", "retail_brand_maternal_pet"),
    ("其他家居用品", "retail_brand_maternal_pet"),
    ("成品家居", "retail_brand_maternal_pet"),
    ("定制家居", "retail_brand_maternal_pet"),
    ("厨房小家电", "retail_brand_maternal_pet"),
    ("厨卫电器", "retail_brand_maternal_pet"),
    ("家纺", "retail_brand_maternal_pet"),
    ("服装家纺", "retail_fashion_group"),
    ("非运动服装", "retail_fashion_group"),
    ("运动服装", "retail_fashion_group"),
    ("鞋帽及其他", "retail_fashion_group"),
    ("钟表珠宝", "retail_fashion_group"),
    ("饰品", "retail_fashion_group"),
    ("其他饰品", "retail_fashion_group"),
    ("白色家电", "retail_brand_maternal_pet"),
    ("黑色家电", "retail_brand_maternal_pet"),
    ("清洁小家电", "retail_brand_maternal_pet"),
    ("文娱用品", "retail_fashion_group"),
    ("专业连锁Ⅱ", "retail_multi_store"),
    ("专业连锁Ⅲ", "retail_multi_store"),
    ("一般零售", "retail_multi_store"),
    ("多业态零售", "retail_multi_store"),
    ("百货", "retail_multi_store"),
    ("超市", "retail_multi_store"),
    ("旅游零售Ⅱ", "retail_multi_store"),
    ("旅游零售Ⅲ", "retail_multi_store"),
    ("零食", "retail_high_sku_brand"),
    ("休闲食品", "retail_high_sku_brand"),
    ("烘焙食品", "retail_high_sku_brand"),
    ("饮料乳品", "retail_high_sku_brand"),
    ("软饮料", "retail_high_sku_brand"),
    ("食品饮料", "retail_high_sku_brand"),
    ("食品加工", "retail_high_sku_brand"),
    ("预加工食品", "retail_high_sku_brand"),
]

PERSONA_META = {
    "retail_brand_beauty": {
        "industry_l2_default": "美妆个护",
        "business_model": "品牌消费品",
        "complexity_tag": "高SKU多渠道",
        "current_business_problem": "新品和营销活动频繁，多渠道分货、补货和库存协同压力大。",
        "primary_jtbd": "商品 / 渠道 / 库存协同",
        "secondary_jtbd": "总部经营透视与营销复盘",
        "admission_reason": "美妆个护品牌与多渠道经营特征成立，适合作为高 SKU 品牌消费品细分画像候选。",
        "product": "美妆个护、护肤、彩妆或相关消费品品牌业务，SKU 丰富且活动频繁。",
        "model": "以品牌消费品经营为主，依赖线上线下或多渠道协同与商品动销管理。",
        "customer": "大众消费客群、女性消费群体或细分个护用户。",
        "similar": "贝泰妮集团股份有限公司,水羊集团股份有限公司,敷尔佳",
        "competition": "美妆个护品牌、功效护肤与日化消费品公司。",
        "knowledge_refs": "ka_insight_beauty_personalcare_brand_v1,ka_narrative_sephora_distribution_v1",
        "talk_refs": "ka_talktrack_inventory_replenishment_retail_v1",
        "static_priority": "A",
    },
    "retail_brand_maternal_pet": {
        "industry_l2_default": "母婴宠物与耐用品",
        "business_model": "品牌消费品",
        "complexity_tag": "高SKU供应链协同",
        "current_business_problem": "SKU 生命周期、补货和履约协同复杂，复购与供应链平衡压力高。",
        "primary_jtbd": "库存与补货协同",
        "secondary_jtbd": "会员与生命周期经营分析",
        "admission_reason": "母婴宠物或家居耐用品属性明确，复购、履约和供应链协同压力适合纳入细分消费品画像。",
        "product": "母婴、宠物、家居耐用品或小家电等消费品业务，兼具产品矩阵与渠道协同复杂度。",
        "model": "以品牌消费品经营为主，强调品类管理、履约效率和供应链协同。",
        "customer": "母婴家庭、宠物养护用户、家庭耐用品消费人群。",
        "similar": "孩子王儿童用品股份有限公司,爱婴室,乖宝宠物",
        "competition": "母婴零售、宠物消费、家居耐用品与小家电品牌。",
        "knowledge_refs": "ka_case_jijia_pet_fullchain_v1,ka_case_qiutianmanman_ops_v1",
        "talk_refs": "ka_talktrack_inventory_replenishment_retail_v1",
        "static_priority": "A",
    },
    "retail_fashion_group": {
        "industry_l2_default": "时尚鞋服与珠宝配饰",
        "business_model": "品牌零售集团",
        "complexity_tag": "多品牌多渠道",
        "current_business_problem": "多品牌、多渠道和企划节奏并存，总部经营透视、会员运营与供应链协同复杂。",
        "primary_jtbd": "总部经营透视与区域差异分析",
        "secondary_jtbd": "商品企划与会员生命周期分析",
        "admission_reason": "时尚鞋服、珠宝配饰和品牌集团属性明显，适合作为消费品细分画像中的集团经营类候选。",
        "product": "时尚鞋服、珠宝配饰、潮玩或相关品牌零售业务，通常具备多品牌或集团化经营特征。",
        "model": "以品牌零售集团经营为主，依赖商品企划、会员运营与渠道协同。",
        "customer": "时尚消费人群、鞋服客群、珠宝配饰消费人群。",
        "similar": "比音勒芬服饰股份有限公司,赢家时尚控股有限公司,周大生珠宝股份有限公司",
        "competition": "时尚鞋服集团、珠宝零售与多品牌零售集团。",
        "knowledge_refs": "ka_case_dazzle_member_lifecycle_v1,ka_case_apparel_multi_brand_ops_v1",
        "talk_refs": "ka_talktrack_hq_visibility_retail_v1",
        "static_priority": "A",
    },
    "retail_multi_store": {
        "industry_l2_default": "多业态零售",
        "business_model": "连锁零售",
        "complexity_tag": "多门店多区域",
        "current_business_problem": "门店网络、区域经营和直营网 / 加盟结构并存，总部经营透视与区域差异分析压力大。",
        "primary_jtbd": "总部经营透视与区域差异分析",
        "secondary_jtbd": "门店经营健康度与商品结构分析",
        "admission_reason": "多业态零售、专业连锁或百货商超网络特征明确，适合作为多门店零售画像候选。",
        "product": "百货、商超、专业连锁或多业态零售业务，具备门店网络与区域经营复杂度。",
        "model": "以多门店连锁零售经营为主，强调总部到区域再到门店的经营穿透。",
        "customer": "区域消费人群、大众零售客群和线下门店消费用户。",
        "similar": "孩子王儿童用品股份有限公司,名创优品（广州）有限责任公司,周大生珠宝股份有限公司",
        "competition": "多业态零售、百货商超和专业连锁公司。",
        "knowledge_refs": "ka_case_xianfeng_retail_v1,ka_solution_hq_visibility_retail_v1",
        "talk_refs": "ka_talktrack_hq_visibility_retail_v1",
        "static_priority": "B",
    },
    "retail_high_sku_brand": {
        "industry_l2_default": "品牌消费品",
        "business_model": "品牌消费品",
        "complexity_tag": "高SKU多渠道",
        "current_business_problem": "商品结构复杂、渠道动销和库存平衡压力大，经营分析需要更细颗粒度。",
        "primary_jtbd": "商品与渠道动销分析",
        "secondary_jtbd": "库存与补货协同",
        "admission_reason": "品牌消费品属性明确，适合作为高 SKU 品牌消费品扩池候选。",
        "product": "食品饮料或品牌消费品业务，多品类、多渠道经营特征明显。",
        "model": "以品牌消费品经营为主，依赖商品、渠道与供应链协同。",
        "customer": "大众消费、家庭消费和品牌零售客群。",
        "similar": "良品铺子股份有限公司,三只松鼠股份有限公司,洽洽食品股份有限公司",
        "competition": "食品饮料与高 SKU 品牌消费品公司。",
        "knowledge_refs": "ka_case_wangxiaolu_growth_v1,ka_case_mixue_bi_warehouse_v1",
        "talk_refs": "ka_talktrack_inventory_replenishment_retail_v1",
        "static_priority": "B",
    },
}

SUPPORTED_PERSONA_IDS = sorted(PERSONA_META.keys())


def normalize_name(value: str | None) -> str:
    if not value:
        return ""
    text = str(value).strip().replace(" ", "")
    text = re.sub(r"[（）()·•\-_/]", "", text)
    for token in [
        "集团股份有限公司",
        "股份有限公司",
        "集团有限公司",
        "有限责任公司",
        "控股有限公司",
        "国际控股有限公司",
        "有限公司",
        "股份公司",
        "集团股份",
        "集团",
        "股份",
        "控股",
        "中国",
    ]:
        text = text.replace(token, "")
    return text


def load_board_codes() -> dict[str, str]:
    df = ak.stock_board_industry_name_em()
    return {str(row["板块名称"]).strip(): str(row["板块代码"]).strip() for _, row in df.iterrows()}


def gather_candidates(board_persona_pairs: list[tuple[str, str]], board_codes: dict[str, str]) -> list[Candidate]:
    seen_codes: set[str] = set()
    gathered: list[Candidate] = []
    for board_name, persona_tag in board_persona_pairs:
        df = ak.stock_board_industry_cons_em(symbol=board_name)
        board_code = board_codes.get(board_name, "")
        for _, row in df.iterrows():
            code = str(row["代码"]).strip()
            name = str(row["名称"]).replace(" ", "").strip()
            if not code or not name:
                continue
            if code.startswith(("200", "900")):
                continue
            if "ST" in name or "退" in name:
                continue
            if code in seen_codes:
                continue
            seen_codes.add(code)
            gathered.append(Candidate(code=code, name=name, board_name=board_name, persona_tag=persona_tag, board_code=board_code))
    return gathered


def load_existing_sets():
    write_target = "main"
    try:
        main_wb = load_workbook(MAIN_WB)
        main_ws = main_wb[main_wb.sheetnames[0]]
    except Exception:
        main_wb = load_workbook(MAIN_SHARED_WB)
        main_ws = main_wb["全量主表"]
        write_target = "shared"
    rows = list(main_ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {key: i for i, key in enumerate(header)}

    current_names = set()
    current_norms = set()
    current_ids = set()
    for row in rows[1:]:
        if "account_id" in idx:
            account_id = row[idx["account_id"]]
            name = row[idx["account_canonical_name"]]
            brand = row[idx["brand_name"]]
            if account_id:
                current_ids.add(str(account_id))
        else:
            name = row[idx["公司主体"]]
            brand = row[idx["品牌名"]]
        for value in [name, brand]:
            if value:
                current_names.add(str(value).strip())
                current_norms.add(normalize_name(value))

    gov_wb = load_workbook(GOV_WB)
    legacy_ws = gov_wb["legacy_customer_registry"]
    alias_ws = gov_wb["alias_registry"]

    legacy_names = set()
    legacy_norms = set()
    for row in legacy_ws.iter_rows(min_row=2, values_only=True):
        for value in row[1:4]:
            if value:
                legacy_names.add(str(value).strip())
                legacy_norms.add(normalize_name(value))

    alias_names = set()
    alias_norms = set()
    for row in alias_ws.iter_rows(min_row=2, values_only=True):
        value = row[1]
        if value:
            alias_names.add(str(value).strip())
            alias_norms.add(normalize_name(value))

    return main_wb, main_ws, idx, current_ids, current_names, current_norms, legacy_names, legacy_norms, alias_names, alias_norms, write_target


def filter_new_candidates(candidates: list[Candidate], current_norms: set[str], legacy_norms: set[str], alias_norms: set[str]) -> list[Candidate]:
    result: list[Candidate] = []
    for candidate in candidates:
        normalized = normalize_name(candidate.name)
        if normalized in current_norms or normalized in legacy_norms or normalized in alias_norms:
            continue
        result.append(candidate)
    return result


def build_main_row(candidate: Candidate) -> dict[str, str]:
    meta = PERSONA_META[candidate.persona_tag]
    board_note = f"东方财富行业板块成分股：{candidate.board_name}（{candidate.board_code}）"
    return {
        "account_id": f"acc_l5_{candidate.code}",
        "account_canonical_name": candidate.name,
        "brand_name": candidate.name,
        "group_name": "",
        "primary_track": "零售消费",
        "industry_l1": "零售消费",
        "industry_l2": candidate.board_name or meta["industry_l2_default"],
        "business_model": meta["business_model"],
        "persona_tag": candidate.persona_tag,
        "secondary_persona_tags": "",
        "company_scale_band": "待补公开财报口径",
        "complexity_tag": meta["complexity_tag"],
        "primary_jtbd": meta["primary_jtbd"],
        "secondary_jtbd": meta["secondary_jtbd"],
        "transformation_stage_tag": "待补公开披露",
        "current_business_problem": meta["current_business_problem"],
        "admission_reason_summary": meta["admission_reason"],
        "case_type_match": "待补案例映射",
        "existing_customer_reference": meta["similar"],
        "solution_match": meta["primary_jtbd"],
        "knowledge_asset_refs": meta["knowledge_refs"],
        "信息扎实度": "中低",
        "ICP匹配概率": "中",
        "静态潜客记录成熟度": "L5",
        "static_priority": meta["static_priority"],
        "dedupe_status": "canonical_new",
        "legacy_customer_check_status": "passed",
        "review_status": "pending_review",
        "source_note": f"{board_note}；消费品新画像扩池 2026-03-31",
        "validation_gap": "待确认官网/年报/IR中的主营产品、商业模式与关键证据。",
        "last_verified_at": TODAY,
        "公司产品与服务概述": meta["product"],
        "商业模式概述": meta["model"],
        "核心客户客群": meta["customer"],
        "收入规模区间": "待补公开财报口径",
        "利润状态概述": "待补公开财报口径",
        "营收增长概述": "待补公开财报口径",
        "已上线系统概况": "待补官网 / 年报 / IR 披露",
        "数字化项目动态": "待补公开披露",
        "相似客户线索": meta["similar"],
        "主要竞品概述": meta["competition"],
        "招聘代表岗位": "待补官网招聘页",
        "近一年重大事件": "上市公司 / 公开主体，待补最近定期报告和经营披露",
    }


def append_main_rows(main_ws, header_idx: dict[str, int], payloads: list[dict[str, str]]) -> None:
    headers = list(header_idx.keys())
    for payload in payloads:
        row = []
        for col in headers:
            if col == "公司主体":
                row.append(payload.get("account_canonical_name", ""))
            elif col == "品牌名":
                row.append(payload.get("brand_name", ""))
            elif col == "主线":
                row.append(payload.get("primary_track", ""))
            elif col == "业务形态画像":
                row.append(payload.get("persona_tag", ""))
            elif col == "管理诉求画像":
                row.append("")
            elif col == "一话入池理由":
                row.append(payload.get("admission_reason_summary", ""))
            elif col == "主要知识资产引用":
                row.append(payload.get("knowledge_asset_refs", ""))
            elif col == "主要切入话术引用":
                row.append(payload.get("talk_track_refs", ""))
            elif col == "待验证项":
                row.append(payload.get("validation_gap", ""))
            else:
                row.append(payload.get(col, ""))
        main_ws.append(row)


def append_profile_rows(candidates: list[Candidate], payloads: dict[str, dict[str, str]]) -> None:
    wb = load_workbook(PROFILE_WB)
    profiles = wb["account_profiles"]
    coverage = wb["profile_coverage"]
    profiles_header = {cell.value: i for i, cell in enumerate(profiles[1], start=1)}
    coverage_header = {cell.value: i for i, cell in enumerate(coverage[1], start=1)}

    for candidate in candidates:
        payload = payloads[candidate.code]
        meta = PERSONA_META[candidate.persona_tag]
        profile_row = {key: "" for key in profiles_header}
        for key in profile_row:
            if key in payload:
                profile_row[key] = payload[key]
        profile_row.update(
            {
                "secondary_persona_tags": "",
                "management_persona_tags": "",
                "static_maturity_level": "L5",
                "archive_status": "有效",
                "产品与服务长摘录": f"{candidate.name} 当前仅基于 {candidate.board_name} 板块归类进入首轮 L5 校验，产品与服务结论仍需官网/年报来源压实。",
                "商业模式长摘录": f"{candidate.name} 的商业模式当前仍是板块级和画像级起草判断，未达到公司级事实完成状态。",
                "客户客群长摘录": f"{candidate.name} 当前客群判断主要来自行业归类与画像映射，后续需补官方来源确认。",
                "系统与数字化长摘录": "当前仅保留为阅读层起草，不得视为已完成公司级系统与数字化事实确认。",
                "重大事件长摘录": "当前仅完成主体与板块层确认，近一年重大事件仍待最近定期报告和官方披露补强。",
                "primary_source_types": "东方财富行业板块成分股",
                "primary_source_refs": f"东方财富行业板块：{candidate.board_name}（{candidate.board_code}） / 股票代码 {candidate.code}",
                "official_source_count": 0,
                "high_confidence_source_count": 1,
                "last_profiled_at": TODAY,
                "profile_owner": "Codex 结构化沉淀",
                "profile_status": "minimum_ready",
                "talk_track_refs": meta["talk_refs"],
                "representative_for_persona": "conditional",
                "representative_for_track": "no",
                "should_have_focus_note": "no",
                "focus_note_path": "",
                "share_status": "未分享",
                "share_batch_id": "",
                "share_last_marked_at": "",
                "share_note": "L5 扩池候选，当前仅完成最小起草，不代表已完成公司级事实确认。",
            }
        )
        profiles.append([profile_row.get(cell.value, "") for cell in profiles[1]])

        coverage_row = {key: "" for key in coverage_header}
        coverage_row.update(
            {
                "account_id": f"acc_l5_{candidate.code}",
                "target_scope": "L5",
                "required_field_count": 4,
                "filled_field_count": 4,
                "coverage_ratio": 1,
                "official_source_ready": "no",
                "high_confidence_ready": "yes",
                "profile_complete_status": "minimum_ready",
                "missing_core_fields": "官网,年报,IR,财报口径,字段级evidence",
                "next_action": "先补官网、年报、IR 与字段级 evidence，再决定是否维持正式候选或进入上移视野。",
            }
        )
        coverage.append([coverage_row.get(cell.value, "") for cell in coverage[1]])

    wb.save(PROFILE_WB)


def append_gov_rows(candidates: list[Candidate], validation_by_code: dict[str, object]) -> None:
    wb = load_workbook(GOV_WB)
    queue = wb["review_queue"]
    evidence = wb["evidence_log"]
    for candidate in candidates:
        account_id = f"acc_l5_{candidate.code}"
        validation = validation_by_code[candidate.code]
        queue.append(
            [
                f"q_l5_{candidate.code}",
                validation.required_queue_type,
                account_id,
                "中",
                "open",
                "codex_llm",
                f"消费品新画像扩池候选；板块={candidate.board_name}；画像={candidate.persona_tag}；{validation.summary}",
                TODAY,
                "",
            ]
        )
        evidence.append(
            [
                f"ev_{account_id}_l5_intake",
                account_id,
                "industry_board_scan",
                f"东方财富行业板块：{candidate.board_name}（{candidate.board_code}） / 股票代码 {candidate.code}",
                "B",
                "主线,画像,入池",
                f"{candidate.name} 已出现在 {candidate.board_name} 板块成分股中，可作为 {candidate.persona_tag} 的 L5 静态候选。",
                "codex_llm",
                TODAY,
                PERSONA_META[candidate.persona_tag]["knowledge_refs"],
                "industry_l2",
                candidate.board_name,
            ]
        )
    wb.save(GOV_WB)


def replace_counts(path: Path, counts: dict[str, int]) -> None:
    text = path.read_text()
    replacements = [
        (r"`L1=\d+ / L2=\d+ / L3=\d+ / L4=\d+ / L5=\d+`", f"`L1={counts['L1']} / L2={counts['L2']} / L3={counts['L3']} / L4={counts['L4']} / L5={counts['L5']}`"),
        (r"`L1=\d+、L2=\d+、L3=\d+、L4=\d+、L5=\d+`", f"`L1={counts['L1']}、L2={counts['L2']}、L3={counts['L3']}、L4={counts['L4']}、L5={counts['L5']}`"),
        (r"`L5=\d+`", f"`L5={counts['L5']}`"),
        (r"`L3\+ = \d+`", f"`L3+ = {counts['L3_plus']}`"),
        (r"`L3\+=\d+`", f"`L3+={counts['L3_plus']}`"),
        (r"`L3\+` 已经有 `\d+` 个", f"`L3+` 已经有 `{counts['L3_plus']}` 个"),
        (r"当前事实口径：`总主体=\d+ / L1=\d+ / L2=\d+ / L3=\d+ / L4=\d+ / L5=\d+`", f"当前事实口径：`总主体={counts['total']} / L1={counts['L1']} / L2={counts['L2']} / L3={counts['L3']} / L4={counts['L4']} / L5={counts['L5']}`"),
        (r"总主体数：`?\d+`?", f"总主体数：`{counts['total']}`"),
        (r"唯一主体数：`?\d+`?", f"唯一主体数：`{counts['total']}`"),
        (r"当前分层：`L1=\d+ / L2=\d+ / L3=\d+ / L4=\d+ / L5=\d+`", f"当前分层：`L1={counts['L1']} / L2={counts['L2']} / L3={counts['L3']} / L4={counts['L4']} / L5={counts['L5']}`"),
        (r"`L3\+` 总量变成：`?\d+`?", f"`L3+` 总量变成：`{counts['L3_plus']}`"),
        (r"`L3\+` 总量是：`?\d+`?", f"`L3+` 总量是：`{counts['L3_plus']}`"),
    ]
    for pattern, repl in replacements:
        text = re.sub(pattern, repl, text)
    path.write_text(text)


def write_exec_doc(candidates: list[Candidate], counts: dict[str, int]) -> None:
    raise RuntimeError("write_exec_doc signature updated; call write_exec_doc_v2 instead")


def write_exec_doc_v2(candidates: list[Candidate], counts: dict[str, int], validation_by_code: dict[str, object]) -> None:
    persona_counter = Counter(candidate.persona_tag for candidate in candidates)
    board_counter = Counter(candidate.board_name for candidate in candidates).most_common(20)
    formal_count = sum(1 for candidate in candidates if validation_by_code[candidate.code].is_formal_l5_candidate)
    observation_count = len(candidates) - formal_count
    issue_counter = Counter()
    for candidate in candidates:
        validation = validation_by_code[candidate.code]
        for issue in validation.issues + validation.warnings:
            issue_counter[issue.code] += 1
    lines = [
        "# 静态潜客池-消费品新画像扩池专项-v1",
        "",
        f"- 执行日期：`{TODAY}`",
        f"- 本轮新增 `L5`：`{len(candidates)}`",
        f"- 正式候选：`{formal_count}`",
        f"- 观察/边界对象：`{observation_count}`",
        f"- 当前主表分层：`L1={counts['L1']} / L2={counts['L2']} / L3={counts['L3']} / L4={counts['L4']} / L5={counts['L5']}`",
        f"- 当前总主体：`{counts['total']}`",
        "",
        "## 画像分布",
        "",
        f"- `retail_brand_beauty`：`{persona_counter['retail_brand_beauty']}`",
        f"- `retail_brand_maternal_pet`：`{persona_counter['retail_brand_maternal_pet']}`",
        f"- `retail_fashion_group`：`{persona_counter['retail_fashion_group']}`",
        f"- `retail_high_sku_brand`：`{persona_counter['retail_high_sku_brand']}`",
        "",
        "## 主要来源板块",
        "",
    ]
    for board_name, value in board_counter:
        lines.append(f"- `{board_name}`：`{value}`")
    lines.extend(
        [
            "",
            "## 执行动作",
            "",
            "- 基于东方财富行业板块成分股，围绕新增消费品细分画像和高 SKU 品牌消费品补入新主体。",
            "- 先做 canonical 去重、alias 排除和老客排除；命中风险主体直接跳过，不阻断总目标。",
            "- 已在写主表前统一执行共享校验，程序先分流正式候选和观察/边界对象，再写入主表、档案库与治理队列。",
            "- 当前统一口径为：`L5` 新入池对象先进入最小档案层，模板起草文案不得自动视为公司级事实完成。",
            "",
            "## 主要阻塞/告警分布",
            "",
        ]
    )
    if issue_counter:
        for code, value in issue_counter.most_common():
            lines.append(f"- `{code}`：`{value}`")
    else:
        lines.append("- 本轮未触发额外阻塞或告警。")
    lines.extend(["", "## 抽样校验摘要", ""])
    for candidate in candidates[:10]:
        lines.append(f"### {candidate.name}")
        lines.append(render_validation_summary(validation_by_code[candidate.code]))
        lines.append("")
    EXEC_MD.write_text("\n".join(lines))


def append_memory(candidates: list[Candidate], counts: dict[str, int]) -> None:
    with MEMORY_MD.open("a") as fh:
        fh.write(
            "\n- 完成消费品新画像扩池：基于东方财富行业板块成分股新增 "
            f"`{len(candidates)}` 家 `L5`，当前主表口径更新为 "
            f"`L1={counts['L1']} / L2={counts['L2']} / L3={counts['L3']} / L4={counts['L4']} / L5={counts['L5']}`；"
            "新增对象已同步写入主表、档案库、review_queue 和 evidence_log。\n"
        )


def compute_counts() -> dict[str, int]:
    try:
        wb = load_workbook(MAIN_WB, data_only=True)
        ws = wb[wb.sheetnames[0]]
    except Exception:
        wb = load_workbook(MAIN_SHARED_WB, data_only=True)
        ws = wb["全量主表"]
    rows = list(ws.iter_rows(values_only=True))
    idx = {key: i for i, key in enumerate(rows[0])}
    maturity = Counter(str(row[idx["静态潜客记录成熟度"]]).strip() for row in rows[1:] if row[idx["静态潜客记录成熟度"]])
    total = len(rows) - 1
    return {
        "total": total,
        "L1": maturity["L1"],
        "L2": maturity["L2"],
        "L3": maturity["L3"],
        "L4": maturity["L4"],
        "L5": maturity["L5"],
        "L3_plus": maturity["L1"] + maturity["L2"] + maturity["L3"],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Retail consumer expand provider for L5 candidate discovery.")
    parser.add_argument("--persona-id", action="append", default=[], help="Optional persona_id filter. Can be repeated.")
    parser.add_argument("--limit", type=int, default=TARGET_NEW_L5, help="Maximum number of candidates to select.")
    parser.add_argument("--output-file", help="Optional JSON result package output path.")
    parser.add_argument("--report-only", action="store_true", help="Only build the provider result package without workbook write-back.")
    return parser


def filter_board_persona_pairs(persona_ids: list[str]) -> list[tuple[str, str]]:
    wanted = {item.strip() for item in persona_ids if item.strip()}
    if not wanted:
        return BOARD_PERSONA
    unsupported = sorted(wanted - set(SUPPORTED_PERSONA_IDS))
    if unsupported:
        raise RuntimeError(f"不支持的消费品 persona_id: {', '.join(unsupported)}")
    return [item for item in BOARD_PERSONA if item[1] in wanted]


def run_expand_provider(
    *,
    persona_ids: list[str] | None = None,
    limit: int = TARGET_NEW_L5,
    report_only: bool = False,
    output_file: str | None = None,
) -> dict[str, object]:
    if not output_file:
        output_path = Path(tempfile.gettempdir()) / "codex-static-pool-runs" / "expand_provider_retail_consumer.json"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_file = str(output_path)
    if limit <= 0:
        raise RuntimeError("limit 必须大于 0")

    board_persona_pairs = filter_board_persona_pairs(persona_ids or [])
    board_codes = load_board_codes()
    (
        main_wb,
        main_ws,
        idx,
        current_ids,
        _current_names,
        current_norms,
        _legacy_names,
        legacy_norms,
        _alias_names,
        alias_norms,
        write_target,
    ) = load_existing_sets()

    all_candidates = gather_candidates(board_persona_pairs, board_codes)
    filtered = filter_new_candidates(all_candidates, current_norms, legacy_norms, alias_norms)
    if len(filtered) < limit:
        raise RuntimeError(f"消费品板块净新增候选不足 {limit}，当前仅 {len(filtered)}")

    selected = filtered[:limit]
    new_ids = {f'acc_l5_{candidate.code}' for candidate in selected}
    if new_ids & current_ids:
        raise RuntimeError("发现 account_id 冲突，停止写入")

    payloads: list[dict[str, str]] = []
    payload_map: dict[str, dict[str, str]] = {}
    validation_by_code: dict[str, object] = {}
    for candidate in selected:
        payload = build_main_row(candidate)
        evidence_rows = [
            {
                "source_locator": f"东方财富行业板块：{candidate.board_name}（{candidate.board_code}） / 股票代码 {candidate.code}",
                "source_type": "industry_board_scan",
                "evidence_strength": "B",
                "summary": f"{candidate.name} 已出现在 {candidate.board_name} 板块成分股中。",
            }
        ]
        validation = classify_l5_candidate(payload, evidence_rows)
        payload["review_status"] = validation.review_status
        payload_map[candidate.code] = payload
        validation_by_code[candidate.code] = validation
        payloads.append(payload)

    counts_after_write = None
    if not report_only:
        append_main_rows(main_ws, idx, payloads)
        main_wb.save(MAIN_WB if write_target == "main" else MAIN_SHARED_WB)
        append_profile_rows(selected, payload_map)
        append_gov_rows(selected, validation_by_code)

        counts_after_write = compute_counts()
        for path in [HOME_MD, TOTAL_MD, MILESTONE_MD]:
            replace_counts(path, counts_after_write)
        write_exec_doc_v2(selected, counts_after_write, validation_by_code)
        append_memory(selected, counts_after_write)

    persona_counter = Counter(candidate.persona_tag for candidate in selected)
    candidate_type_counter = Counter(validation_by_code[candidate.code].candidate_type for candidate in selected)
    issue_counter = Counter()
    for candidate in selected:
        validation = validation_by_code[candidate.code]
        for issue in validation.issues + validation.warnings:
            issue_counter[issue.code] += 1

    result = {
        "provider": "retail_consumer_board_scan",
        "track": "零售消费",
        "write_target": write_target,
        "selected_persona_ids": sorted({candidate.persona_tag for candidate in selected}),
        "report_only": report_only,
        "summary": {
            "new_l5": len(selected),
            "persona_counter": dict(persona_counter),
            "candidate_type_counter": dict(candidate_type_counter),
            "issue_counter": dict(issue_counter),
        },
        "counts_after_write": counts_after_write,
        "sample_names": [candidate.name for candidate in selected[:20]],
        "results": [
            {
                "account_id": payload_map[candidate.code]["account_id"],
                "account_canonical_name": candidate.name,
                "persona_tag": candidate.persona_tag,
                "board_name": candidate.board_name,
                "board_code": candidate.board_code,
                "candidate_type": validation_by_code[candidate.code].candidate_type,
                "review_status": validation_by_code[candidate.code].review_status,
                "required_queue_type": validation_by_code[candidate.code].required_queue_type,
                "summary": validation_by_code[candidate.code].summary,
            }
            for candidate in selected
        ],
    }
    if output_file:
        Path(output_file).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> None:
    args = build_parser().parse_args()
    result = run_expand_provider(
        persona_ids=args.persona_id,
        limit=args.limit,
        report_only=args.report_only,
        output_file=args.output_file,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
