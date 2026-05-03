from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import check_workbook_integrity, resolve_static_pool_paths, workbook_write_lock
from shared.static_pool.promotion_writeback import backup_once, build_header_index


DEFAULT_PATCH = "deliveries/archive/milestones/milestone35r_field_governance/milestone35r_safe_profile_patch_v1.json"
DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone35r_field_governance_true_repair"
DEFAULT_REVIEW_MD = "docs/03-执行与校验/Milestone 35R-安全字段真实修复复盘-v1.md"
PROFILE_SHEET = "account_profiles"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Execute M35R safe profile field repair after user confirmation.")
    parser.add_argument("--patch-file", default=DEFAULT_PATCH)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--review-md", default=DEFAULT_REVIEW_MD)
    parser.add_argument("--lock-timeout-seconds", type=float, default=5.0)
    return parser


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(value: object) -> str:
    return str(value or "").strip()


def _write_json(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str | Path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _profile_index(profile_path: Path) -> tuple[dict[str, int], dict[str, int]]:
    wb = load_workbook(profile_path, read_only=True, data_only=True)
    try:
        ws = wb[PROFILE_SHEET]
        headers = [cell.value for cell in ws[1]]
        header_index = {str(value): idx + 1 for idx, value in enumerate(headers) if value}
        account_col = header_index["account_id"]
        row_index: dict[str, int] = {}
        for row in range(2, ws.max_row + 1):
            account_id = _clean(ws.cell(row, account_col).value)
            if account_id:
                row_index[account_id] = row
        return header_index, row_index
    finally:
        wb.close()


def _preflight(profile_path: Path, patch: dict[str, Any]) -> dict[str, Any]:
    header_index, row_index = _profile_index(profile_path)
    missing_accounts: list[str] = []
    missing_fields: set[str] = set()
    non_blank_targets: list[dict[str, str]] = []
    total_fields = 0
    wb = load_workbook(profile_path, read_only=True, data_only=True)
    try:
        ws = wb[PROFILE_SHEET]
        for account in patch.get("accounts") or []:
            account_id = _clean(account.get("account_id"))
            row_num = row_index.get(account_id)
            if not row_num:
                missing_accounts.append(account_id)
                continue
            for field, value in (account.get("updates") or {}).items():
                total_fields += 1
                if field not in header_index:
                    missing_fields.add(field)
                    continue
                current = _clean(ws.cell(row_num, header_index[field]).value)
                if current:
                    non_blank_targets.append({"account_id": account_id, "field": field, "current_value": current, "patch_value": _clean(value)})
    finally:
        wb.close()
    checks = {
        "account_count_is_64": len(patch.get("accounts") or []) == 64,
        "field_count_is_320": total_fields == 320,
        "all_accounts_exist": not missing_accounts,
        "all_fields_exist": not missing_fields,
        "all_targets_blank": not non_blank_targets,
    }
    return {
        "checks": checks,
        "ok": all(checks.values()),
        "account_count": len(patch.get("accounts") or []),
        "field_count": total_fields,
        "missing_accounts": missing_accounts,
        "missing_fields": sorted(missing_fields),
        "non_blank_targets": non_blank_targets[:50],
        "non_blank_target_count": len(non_blank_targets),
    }


def _execute_patch(profile_path: Path, patch: dict[str, Any]) -> dict[str, Any]:
    wb = load_workbook(profile_path)
    ws = wb[PROFILE_SHEET]
    header_index = build_header_index(ws)
    account_col = header_index["account_id"]
    row_index = {
        _clean(ws.cell(row, account_col).value): row
        for row in range(2, ws.max_row + 1)
        if _clean(ws.cell(row, account_col).value)
    }
    updated_accounts: set[str] = set()
    updated_fields = 0
    field_distribution: Counter[str] = Counter()
    samples: list[dict[str, Any]] = []
    for account in patch.get("accounts") or []:
        account_id = _clean(account.get("account_id"))
        row_num = row_index.get(account_id)
        if not row_num:
            continue
        account_updates = 0
        for field, value in (account.get("updates") or {}).items():
            current = _clean(ws.cell(row_num, header_index[field]).value)
            if current:
                continue
            ws.cell(row_num, header_index[field]).value = value
            updated_fields += 1
            account_updates += 1
            field_distribution[field] += 1
        if account_updates:
            updated_accounts.add(account_id)
            if len(samples) < 10:
                samples.append({"account_id": account_id, "account_name": account.get("account_name"), "updated_field_count": account_updates})
    wb.save(profile_path)
    wb.close()
    return {
        "updated_account_count": len(updated_accounts),
        "updated_field_count": updated_fields,
        "field_distribution": dict(field_distribution),
        "samples": samples,
    }


def _render_md(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 35R-安全字段真实修复复盘-v1",
        "",
        "## 摘要",
        "",
        f"- 执行状态：`{payload['status']}`",
        f"- 补写 profile 账户数：`{summary['updated_account_count']}`",
        f"- 补写字段数：`{summary['updated_field_count']}`",
        f"- 工作簿完整性：`{summary['post_workbook_integrity_ok']}`",
        f"- 锁状态：`acquired={summary['workbook_lock_acquired']}`，wait_seconds=`{summary['workbook_lock_wait_seconds']}`",
        "",
        "## 边界",
        "",
        "- 本轮只补 profile 中主表非空、档案为空的安全字段。",
        "- 未处理需人工复核的 profile 差异。",
        "- 未让共享版反向覆盖主表。",
        "- 未写入知识资产注册表或画像注册表。",
    ]
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    pool = resolve_static_pool_paths()
    patch = _load_json(args.patch_file)
    preflight = _preflight(pool["profile"], patch)
    output_dir = Path(args.output_dir)
    if not preflight["ok"]:
        payload = {
            "batch_id": "milestone35r_safe_profile_field_true_repair_v1",
            "generated_at": _now(),
            "status": "blocked_preflight_failed",
            "static_pool_root": str(pool["root"]),
            "preflight": preflight,
        }
        _write_json(output_dir / "milestone35r_safe_profile_field_true_repair_summary_v1.json", payload)
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 1

    with workbook_write_lock(timeout_seconds=args.lock_timeout_seconds) as lock_meta:
        backups = {
            "profile": backup_once(pool["profile"], "m35r_field_repair_backup"),
            "main": backup_once(pool["main"], "m35r_field_repair_reference_backup"),
            "main_shared": backup_once(pool["main_shared"], "m35r_field_repair_reference_backup"),
            "governance": backup_once(pool["governance"], "m35r_field_repair_reference_backup"),
        }
        execution = _execute_patch(pool["profile"], patch)
        post_integrity = check_workbook_integrity([pool["main"], pool["profile"], pool["main_shared"], pool["governance"]], deep_scan=True)

    summary = {
        **execution,
        "post_workbook_integrity_ok": bool(post_integrity.get("ok")),
        "workbook_lock_acquired": bool(lock_meta.get("acquired")),
        "workbook_lock_wait_seconds": lock_meta.get("wait_seconds"),
        "true_field_repair_executed": True,
        "shared_reverse_overwrite_enabled": False,
        "knowledge_asset_write_enabled": False,
        "persona_registry_write_enabled": False,
    }
    payload = {
        "batch_id": "milestone35r_safe_profile_field_true_repair_v1",
        "generated_at": _now(),
        "status": "success",
        "static_pool_root": str(pool["root"]),
        "preflight": preflight,
        "backups": backups,
        "summary": summary,
        "workbook_integrity": post_integrity,
    }
    _write_json(output_dir / "milestone35r_safe_profile_field_true_repair_summary_v1.json", payload)
    _write_text(args.review_md, _render_md(payload))
    print(json.dumps({"output_json": str(output_dir / "milestone35r_safe_profile_field_true_repair_summary_v1.json"), "review_md": args.review_md, "summary": summary}, ensure_ascii=False, indent=2))
    ok = execution["updated_account_count"] == 64 and execution["updated_field_count"] == 320 and summary["post_workbook_integrity_ok"]
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
