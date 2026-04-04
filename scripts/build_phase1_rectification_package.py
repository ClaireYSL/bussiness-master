from __future__ import annotations

import json
from pathlib import Path
import sys

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
ROOT = WORKSPACE.parents[2]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import LEGACY_PERSONA_TAG_MAP, STANDARD_PERSONA_IDS

VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
LEARNING_XLSX = VAULT / "知识资产注册表.xlsx"
DELIVERIES_DIR = WORKSPACE / "deliveries"
OUTPUT_JSON = DELIVERIES_DIR / "phase1_rectification_package_v1.json"

DELEGATE_FILES = (
    DELIVERIES_DIR / "calibration_batch_phase1_retail_v1.delegate.json",
    DELIVERIES_DIR / "calibration_batch_phase1_cbec_v1.delegate.json",
    DELIVERIES_DIR / "calibration_batch_phase1_mfg_v1.delegate.json",
)

PHASE1_DECISIONS = {
    "acc_miniso": {
        "suggested_primary_persona": "retail_multi_store",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留主画像与层级，补 promotion_review 队列和组织结构颗粒度。",
    },
    "acc_threesquirrels": {
        "suggested_primary_persona": "retail_high_sku_brand",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "verification",
        "rectification_action": "保留主画像，重写模板化产品/服务描述并补强来源。",
    },
    "acc_ays": {
        "suggested_primary_persona": "retail_multi_store",
        "secondary_persona_tags": ["retail_brand_maternal_pet"],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留多门店主画像，补次级母婴品牌画像用于辅助解释。",
    },
    "acc_popmart": {
        "suggested_primary_persona": "",
        "secondary_persona_tags": ["retail_fashion_group"],
        "final_candidate_type": "observation",
        "final_review_status": "pending_review",
        "queue_type": "boundary_review",
        "rectification_action": "退出正式候选，先补官方来源和商业模式，再重判主画像。",
    },
    "acc_eastroc": {
        "suggested_primary_persona": "retail_high_sku_brand",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留主画像并接受 L3 建议，作为正向升级样本。",
    },
    "acc_anker": {
        "suggested_primary_persona": "cbec_multi_platform_brand",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留主画像与层级，补平台矩阵与区域颗粒度证据。",
    },
    "acc_songmics": {
        "suggested_primary_persona": "cbec_multi_platform_brand",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留主画像，继续补多平台与区域经营颗粒度。",
    },
    "acc_jihong": {
        "suggested_primary_persona": "cbec_multi_platform_brand",
        "secondary_persona_tags": ["cbec_platform_operator"],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "verification",
        "rectification_action": "清理旧 persona，回到标准画像体系并下调到 L2。",
    },
    "acc_zibuyu": {
        "suggested_primary_persona": "cbec_multi_platform_brand",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "verification",
        "rectification_action": "清理旧 persona，保留跨境品牌出海主画像并下调到 L3。",
    },
    "acc_loctek": {
        "suggested_primary_persona": "",
        "secondary_persona_tags": ["cbec_multi_platform_brand"],
        "final_candidate_type": "observation",
        "final_review_status": "pending_review",
        "queue_type": "boundary_review",
        "rectification_action": "退出正式候选，先补产品/服务、商业模式和官方来源，再重判。",
    },
    "acc_inovance": {
        "suggested_primary_persona": "mfg_multi_factory_group",
        "secondary_persona_tags": ["mfg_rnd_sales_complex"],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留多工厂主画像，并补研产销协同作为次级画像解释。",
    },
    "acc_sany": {
        "suggested_primary_persona": "mfg_multi_factory_group",
        "secondary_persona_tags": [],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "promotion_review",
        "rectification_action": "保留主画像并补多工厂与渠道协同细节。",
    },
    "acc_amec": {
        "suggested_primary_persona": "",
        "secondary_persona_tags": ["mfg_rnd_sales_complex"],
        "final_candidate_type": "observation",
        "final_review_status": "pending_review",
        "queue_type": "boundary_review",
        "rectification_action": "暂不硬贴主画像，先补官方来源与制造协同边界证据。",
    },
    "acc_bozon": {
        "suggested_primary_persona": "mfg_rnd_sales_complex",
        "secondary_persona_tags": ["mfg_multi_factory_group"],
        "final_candidate_type": "formal_candidate",
        "final_review_status": "active",
        "queue_type": "verification",
        "rectification_action": "保留技术型制造主画像，并补多工厂协同作为次级画像说明。",
    },
    "acc_jereh": {
        "suggested_primary_persona": "",
        "secondary_persona_tags": ["mfg_multi_factory_group"],
        "final_candidate_type": "observation",
        "final_review_status": "pending_review",
        "queue_type": "boundary_review",
        "rectification_action": "退出正式候选，先补产品/服务、商业模式与官方来源，再重判。",
    },
}


def load_sheet_rows(path: Path, sheet_name: str) -> list[dict[str, object]]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def load_delegate_samples() -> list[dict[str, object]]:
    samples: list[dict[str, object]] = []
    for path in DELEGATE_FILES:
        payload = json.loads(path.read_text(encoding="utf-8"))
        samples.extend(payload.get("samples", []))
    return samples


def canonical_persona(value: object) -> str:
    text = str(value or "").strip()
    return LEGACY_PERSONA_TAG_MAP.get(text, text)


def maturity_rank(value: object) -> int:
    text = str(value or "").strip().upper()
    if text.startswith("L") and text[1:].isdigit():
        return int(text[1:])
    return 0


def find_matching_learning_queue(track: str, persona_tags: list[str], queue_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    matched: list[dict[str, object]] = []
    track_map = {
        "零售消费": "retail_consumer",
        "跨境电商": "cross_border_ecommerce",
        "先进制造": "advanced_manufacturing",
    }
    queue_track = track_map.get(track, "")
    wanted = set(persona_tags)
    for row in queue_rows:
        if str(row.get("extraction_status") or "") != "queued":
            continue
        trigger_track = str(row.get("trigger_track_id") or "")
        trigger_persona_ids = {
            canonical_persona(item)
            for item in str(row.get("trigger_persona_id") or "").replace("，", ",").split(",")
            if item.strip() and not item.strip().startswith("mgmt_")
        }
        if queue_track and trigger_track == queue_track:
            matched.append(row)
            continue
        if wanted & trigger_persona_ids:
            matched.append(row)
    uniq: list[dict[str, object]] = []
    seen: set[str] = set()
    for row in matched:
        queue_id = str(row.get("queue_id") or "")
        if queue_id and queue_id not in seen:
            seen.add(queue_id)
            uniq.append(row)
    return uniq


def main() -> int:
    profile_rows = load_sheet_rows(PROFILE_XLSX, "account_profiles")
    queue_rows = load_sheet_rows(LEARNING_XLSX, "learning_queue")
    profile_by_id = {str(row.get("account_id") or ""): row for row in profile_rows if row.get("account_id")}
    profile_by_name = {str(row.get("account_canonical_name") or ""): row for row in profile_rows if row.get("account_canonical_name")}

    results: list[dict[str, object]] = []
    summary = {
        "formal_candidate": 0,
        "observation": 0,
        "downgrade": 0,
        "maturity_regraded": 0,
        "persona_adjusted": 0,
        "official_source_repair_needed": 0,
    }

    for sample in load_delegate_samples():
        account_id = str(sample.get("account_id") or "")
        account_name = str(sample.get("account_name") or "")
        profile = profile_by_id.get(account_id) or profile_by_name.get(account_name) or {}
        decision = PHASE1_DECISIONS[account_id]

        current_persona = str(profile.get("persona_tag") or "")
        normalized_current_persona = canonical_persona(current_persona)
        suggested_primary_persona = decision["suggested_primary_persona"]
        secondary_personas = decision["secondary_persona_tags"]
        matched_queue = find_matching_learning_queue(
            str(profile.get("primary_track") or ""),
            [tag for tag in [suggested_primary_persona, normalized_current_persona, *secondary_personas] if tag],
            queue_rows,
        )
        knowledge_refs = [item.strip() for item in str(profile.get("knowledge_asset_refs") or "").split(",") if item.strip()]
        talk_refs = [item.strip() for item in str(profile.get("talk_track_refs") or "").split(",") if item.strip()]
        maturity_status = sample.get("maturity_assessment", {}).get("status")
        suggested_maturity = sample.get("maturity_assessment", {}).get("suggested_maturity")
        if maturity_status == "downgrade":
            summary["downgrade"] += 1
        if maturity_rank(suggested_maturity) and maturity_rank(profile.get("静态潜客记录成熟度")):
            if maturity_rank(suggested_maturity) != maturity_rank(profile.get("静态潜客记录成熟度")):
                summary["maturity_regraded"] += 1
        if sample.get("persona_assessment", {}).get("status") != "keep":
            summary["persona_adjusted"] += 1
        if "official_source_missing" in sample.get("risk_flags", []):
            summary["official_source_repair_needed"] += 1
        summary[decision["final_candidate_type"]] += 1

        result = {
            "account_id": account_id,
            "account_name": account_name,
            "current_state": {
                "primary_track": profile.get("primary_track"),
                "persona_tag": current_persona,
                "normalized_current_persona": normalized_current_persona,
                "静态潜客记录成熟度": profile.get("静态潜客记录成熟度"),
                "review_status": profile.get("archive_status") or "",
                "knowledge_asset_refs": knowledge_refs,
                "talk_track_refs": talk_refs,
            },
            "rectification": {
                "track_decision": sample.get("track_assessment", {}).get("status"),
                "persona_decision": sample.get("persona_assessment", {}).get("status"),
                "minimum_fact_decision": sample.get("minimum_fact_assessment", {}).get("status"),
                "maturity_decision": maturity_status,
                "suggested_maturity": suggested_maturity,
                "suggested_primary_persona": suggested_primary_persona,
                "suggested_secondary_persona_tags": secondary_personas,
                "final_candidate_type": decision["final_candidate_type"],
                "final_review_status": decision["final_review_status"],
                "required_queue_type": decision["queue_type"],
                "rectification_action": decision["rectification_action"],
                "rewrite_suggestion": sample.get("rewrite_suggestion", {}),
                "evidence_gaps": sample.get("evidence_gaps", []),
                "risk_flags": sample.get("risk_flags", []),
            },
            "knowledge_alignment": {
                "strong_related_knowledge_assets": knowledge_refs,
                "strong_related_talk_tracks": talk_refs,
                "matched_learning_queue_items": [
                    {
                        "queue_id": row.get("queue_id"),
                        "material_title": row.get("material_title"),
                        "priority": row.get("priority"),
                        "trigger_persona_id": row.get("trigger_persona_id"),
                    }
                    for row in matched_queue
                ],
            },
        }
        results.append(result)

    output = {
        "milestone": "Milestone 4",
        "title": "治理收口与 Phase 1 全样本纠偏",
        "standard_persona_ids": sorted(STANDARD_PERSONA_IDS),
        "legacy_persona_alias_map": LEGACY_PERSONA_TAG_MAP,
        "summary": summary,
        "results": results,
    }
    OUTPUT_JSON.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output_file": str(OUTPUT_JSON), "result_count": len(results), "summary": summary}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
