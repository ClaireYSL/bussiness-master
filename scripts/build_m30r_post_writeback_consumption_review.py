from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import check_workbook_integrity, resolve_static_pool_paths


DEFAULT_CONSUMPTION = "deliveries/archive/milestones/milestone26r_m25r_trusted_consumption/milestone26r_m25r_trusted_consumption_package_v1.json"
DEFAULT_RUN = "deliveries/archive/milestones/milestone28r_m25r_promote_admission/milestone28r_m25r_promote_admission_run_summary_v1.json"
DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone30r_post_writeback_consumption_review"
DEFAULT_TEMPLATE = "configs/execution_batches/milestone30r_post_writeback_business_review_template_v1.json"
DEFAULT_REVIEW_MD = "docs/03-执行与校验/Milestone 30R-写回后50家可消费池复核与业务抽样评估-v1.md"

SAMPLE_SIZE = 15


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M30R post-writeback consumption review and sampling package.")
    parser.add_argument("--consumption-file", default=DEFAULT_CONSUMPTION)
    parser.add_argument("--run-summary-file", default=DEFAULT_RUN)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--template-json", default=DEFAULT_TEMPLATE)
    parser.add_argument("--review-md", default=DEFAULT_REVIEW_MD)
    parser.add_argument("--sample-size", type=int, default=SAMPLE_SIZE)
    return parser


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(value: object) -> str:
    return str(value or "").strip()


def _load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write_json(path: str | Path, payload: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str | Path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _rows_by(workbook_path: Path, sheet_name: str, key_header: str) -> dict[str, dict[str, Any]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        rows: dict[str, dict[str, Any]] = {}
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            key = _clean(row.get(key_header))
            if key:
                rows[key] = row
        return rows
    finally:
        wb.close()


def _group_rows(workbook_path: Path, sheet_name: str, key_header: str) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            key = _clean(row.get(key_header))
            if key:
                grouped[key].append(row)
        return dict(grouped)
    finally:
        wb.close()


def _card_index(consumption: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        _clean(card.get("account_id")): card
        for card in consumption.get("trusted_prospect_cards") or []
        if isinstance(card, dict) and _clean(card.get("account_id"))
    }


def _promoted_ids(run_summary: dict[str, Any]) -> list[str]:
    samples = (((run_summary.get("promote") or {}).get("write_back") or {}).get("samples") or [])
    return [_clean(item.get("account_id")) for item in samples if isinstance(item, dict) and _clean(item.get("account_id"))]


def _is_l3_ready(main: dict[str, Any], profile: dict[str, Any], shared: dict[str, Any], evidence_rows: list[dict[str, Any]]) -> tuple[bool, list[str]]:
    gaps: list[str] = []
    if _clean(main.get("静态潜客记录成熟度")) != "L3":
        gaps.append("main_not_l3")
    if _clean(profile.get("静态潜客记录成熟度")) != "L3":
        gaps.append("profile_not_l3")
    if shared and _clean(shared.get("静态潜客记录成熟度")) != "L3":
        gaps.append("shared_not_l3")
    if _clean(profile.get("profile_status")) != "standard_ready":
        gaps.append("profile_not_standard_ready")
    for field in ("公司产品与服务概述", "商业模式概述", "admission_reason_summary"):
        if not _clean(main.get(field)):
            gaps.append(f"main_missing_{field}")
    if len(evidence_rows) < 3:
        gaps.append("evidence_less_than_3")
    return not gaps, gaps


def _business_review_item(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "account_id": item["account_id"],
        "company_name": item["company_name"],
        "matched_track": item["matched_track"],
        "matched_persona": item["matched_persona"],
        "sample_reason": item["sample_reason"],
        "why_it_matches": item["match_reason"],
        "core_product_or_service": item["core_product_or_service"],
        "business_model": item["business_model"],
        "primary_evidence_locator": item["primary_evidence_locator"],
        "review_questions": [
            "这家公司是否真的符合当前 ICP/画像？",
            "核心产品/服务和业务模式是否足以支撑首轮判断？",
            "是否值得进入 L2 深研、销售复核、持续监控或淘汰？",
        ],
        "business_feedback_fields": {
            "business_fit_rating": "",
            "worth_following": "",
            "recommended_next_action": "",
            "priority_rank": "",
            "target_scenario": "",
            "disqualify_reason": "",
            "feedback_owner": "",
            "feedback_date": "",
            "feedback_notes": "",
        },
    }


def _select_sample(items: list[dict[str, Any]], sample_size: int) -> list[dict[str, Any]]:
    by_persona: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in sorted(items, key=lambda row: row["account_id"]):
        by_persona[item["matched_persona"]].append(item)

    selected: list[dict[str, Any]] = []
    seen: set[str] = set()
    persona_order = sorted(by_persona, key=lambda persona: (-len(by_persona[persona]), persona))

    # First pass gives every persona at least one seat where possible.
    for persona in persona_order:
        if len(selected) >= sample_size:
            break
        candidate = by_persona[persona][0]
        selected.append({**candidate, "sample_reason": f"画像 `{persona}` 代表样本"})
        seen.add(candidate["account_id"])

    # Round-robin fill keeps large personas represented without losing diversity.
    cursor = 1
    while len(selected) < sample_size:
        added = False
        for persona in persona_order:
            rows = by_persona[persona]
            if cursor >= len(rows):
                continue
            candidate = rows[cursor]
            if candidate["account_id"] in seen:
                continue
            selected.append({**candidate, "sample_reason": f"画像 `{persona}` 分层抽样"})
            seen.add(candidate["account_id"])
            added = True
            if len(selected) >= sample_size:
                break
        if not added:
            break
        cursor += 1
    return selected


def _render_md(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 30R-写回后50家可消费池复核与业务抽样评估-v1",
        "",
        "## 摘要",
        "",
        f"- 写回对象：`{summary['account_count']}`",
        f"- L3 可消费对象：`{summary['l3_share_ready_count']}`",
        f"- 仍需修复对象：`{summary['needs_fix_count']}`",
        f"- 抽样对象：`{summary['sample_count']}`",
        f"- 画像分布：`{summary['persona_counts']}`",
        f"- 工作簿完整性：`{summary['workbook_integrity_ok']}`",
        f"- M28R 写回：`promoted={summary['m28r_promoted']} / skipped={summary['m28r_skipped']}`",
        "",
        "## 结论",
        "",
        "- M30R 不做写回，只做写回后产品化复核和业务抽样准备。",
        "- 这一步的核心问题不是“还能不能跑”，而是“这 50 家是否真能被业务侧读懂、判断、反馈”。",
        "- 潜客观察仍不能直接进入正式知识资产，只能形成业务反馈、source gap 或规则校准建议。",
        "",
        "## 抽样评估清单",
        "",
    ]
    for item in payload["sample_review_items"]:
        lines.extend(
            [
                f"### {item['company_name']}（{item['account_id']}）",
                "",
                f"- 主线/画像：`{item['matched_track']}` / `{item['matched_persona']}`",
                f"- 抽样原因：{item['sample_reason']}",
                f"- 为什么匹配：{item['why_it_matches']}",
                f"- 核心产品/服务：{item['core_product_or_service']}",
                f"- 经营结构：{item['business_model']}",
                f"- 关键 evidence：{item['primary_evidence_locator']}",
                "",
            ]
        )
    lines.extend(
        [
            "## 下一步",
            "",
            "1. 业务/研究侧填写抽样评估模板。",
            "2. 汇总 `worth_following`、`recommended_next_action` 和 `disqualify_reason`。",
            "3. 只把反馈沉淀为候选观察和规则校准建议，不直接改写正式知识资产。",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    pool = resolve_static_pool_paths()
    consumption = _load_json(args.consumption_file)
    run = _load_json(args.run_summary_file)
    cards = _card_index(consumption)
    promoted_ids = _promoted_ids(run)
    main_by_id = _rows_by(pool["main"], "accounts_main", "account_id")
    profile_by_id = _rows_by(pool["profile"], "account_profiles", "account_id")
    shared_by_name = _rows_by(pool["main_shared"], "全量主表", "公司主体")
    evidence_by_id = _group_rows(pool["governance"], "evidence_log", "account_id")
    integrity = check_workbook_integrity([pool["main"], pool["profile"], pool["governance"]], deep_scan=True)

    items: list[dict[str, Any]] = []
    for account_id in promoted_ids:
        card = cards.get(account_id, {})
        main = main_by_id.get(account_id, {})
        profile = profile_by_id.get(account_id, {})
        company_name = _clean(main.get("account_canonical_name") or card.get("company_name"))
        shared = shared_by_name.get(company_name, {})
        evidence_rows = evidence_by_id.get(account_id, [])
        ready, gaps = _is_l3_ready(main, profile, shared, evidence_rows)
        strong = card.get("strong_evidence") if isinstance(card.get("strong_evidence"), list) else []
        primary_evidence = strong[0] if strong else {}
        items.append(
            {
                "account_id": account_id,
                "company_name": company_name,
                "matched_track": _clean(main.get("primary_track") or card.get("matched_track")),
                "matched_persona": _clean(main.get("persona_tag") or card.get("matched_persona")),
                "main_level": _clean(main.get("静态潜客记录成熟度")),
                "profile_level": _clean(profile.get("静态潜客记录成熟度")),
                "shared_level": _clean(shared.get("静态潜客记录成熟度")),
                "profile_status": _clean(profile.get("profile_status")),
                "match_reason": _clean(card.get("match_reason") or main.get("admission_reason_summary")),
                "core_product_or_service": _clean(card.get("core_product_or_service") or main.get("公司产品与服务概述")),
                "business_model": _clean(card.get("business_model") or main.get("商业模式概述")),
                "primary_evidence_locator": _clean(primary_evidence.get("source_locator")),
                "evidence_count": len(evidence_rows),
                "post_writeback_status": "l3_share_ready" if ready else "needs_post_writeback_fix",
                "quality_gaps": gaps,
                "knowledge_asset_boundary_note": "该对象是潜客池消费材料，不是客户案例或正式知识资产来源。",
            }
        )

    sample = _select_sample(items, args.sample_size)
    sample_review_items = [_business_review_item(item) for item in sample]
    promote_writeback = ((run.get("promote") or {}).get("write_back") or {})
    status_counts = Counter(item["post_writeback_status"] for item in items)
    summary = {
        "account_count": len(items),
        "l3_share_ready_count": int(status_counts.get("l3_share_ready") or 0),
        "needs_fix_count": int(status_counts.get("needs_post_writeback_fix") or 0),
        "sample_count": len(sample_review_items),
        "persona_counts": dict(Counter(item["matched_persona"] for item in items)),
        "track_counts": dict(Counter(item["matched_track"] for item in items)),
        "sample_persona_counts": dict(Counter(item["matched_persona"] for item in sample)),
        "workbook_integrity_ok": bool(integrity.get("ok")),
        "m28r_promoted": int(promote_writeback.get("promoted") or 0),
        "m28r_skipped": int(promote_writeback.get("skipped") or 0),
        "true_writeback_enabled": False,
        "formal_knowledge_write_enabled": False,
    }
    payload = {
        "batch_id": "milestone30r_post_writeback_consumption_review_v1",
        "generated_at": _now(),
        "source_files": {
            "consumption_file": args.consumption_file,
            "run_summary_file": args.run_summary_file,
        },
        "static_pool_root": str(pool["root"]),
        "policy": {
            "true_writeback_enabled": False,
            "formal_knowledge_write_enabled": False,
            "not_customer_case": True,
            "not_persona_registry_update": True,
        },
        "summary": summary,
        "items": items,
        "sample_review_items": sample_review_items,
        "workbook_integrity": integrity,
    }
    template = {
        "batch_id": "milestone30r_post_writeback_business_review_template_v1",
        "generated_at": _now(),
        "source_file": str(Path(args.output_dir) / "milestone30r_post_writeback_consumption_review_package_v1.json"),
        "allowed_worth_following_values": ["yes", "no", "unclear"],
        "allowed_recommended_next_actions": [
            "advance_to_l2_research",
            "sales_review",
            "keep_l3_monitoring",
            "disqualify",
            "needs_more_context",
        ],
        "sample_review_items": sample_review_items,
    }
    output_dir = Path(args.output_dir)
    _write_json(output_dir / "milestone30r_post_writeback_consumption_review_package_v1.json", payload)
    _write_json(output_dir / "milestone30r_post_writeback_sample_review_items_v1.json", {"batch_id": "milestone30r_post_writeback_sample_review_items_v1", "items": sample_review_items})
    _write_json(args.template_json, template)
    _write_text(args.review_md, _render_md(payload))
    print(json.dumps({"output_json": str(output_dir / "milestone30r_post_writeback_consumption_review_package_v1.json"), "template_json": args.template_json, "review_md": args.review_md, "summary": summary}, ensure_ascii=False, indent=2))
    ok = (
        summary["account_count"] == 50
        and summary["l3_share_ready_count"] == 50
        and summary["needs_fix_count"] == 0
        and summary["sample_count"] == args.sample_size
        and summary["workbook_integrity_ok"]
        and not summary["true_writeback_enabled"]
    )
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
