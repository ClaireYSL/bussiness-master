from __future__ import annotations

from collections import defaultdict
import os
from pathlib import Path

from openpyxl import load_workbook


VAULT = Path("/Users/clairaipartner/Documents/Obsidian-Codex/潜客池")
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
PERSONA_XLSX = VAULT / "主线与画像注册表.xlsx"
ASSET_XLSX = VAULT / "知识资产注册表.xlsx"
TODAY = "2026-04-02"
TARGET_LEVELS = {x.strip() for x in os.environ.get("ENRICH_LEVELS", "L1").split(",") if x.strip()}
MARKET_PLACEHOLDER = "待补公司级市场参考"
PERSONA_FALLBACKS = {
    "cbec_brand_outbound": "cbec_multi_platform_brand",
}


def normalize_track(track_id: str) -> str:
    return {
        "retail_consumer": "零售消费",
        "cross_border_ecommerce": "跨境电商",
        "advanced_manufacturing": "先进制造",
    }.get(track_id, track_id)


def load_persona_meta():
    wb = load_workbook(PERSONA_XLSX, read_only=True, data_only=True)
    ws = wb["personas"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    meta = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        persona_id = str(row[idx["persona_id"]] or "").strip()
        if not persona_id:
            continue
        meta[persona_id] = {
            "track": normalize_track(str(row[idx["track_id"]] or "").strip()),
            "reference_customers": str(row[idx["reference_customers"]] or "").strip(),
            "reference_cases": str(row[idx["reference_cases"]] or "").strip(),
            "reference_solutions": str(row[idx["reference_solutions"]] or "").strip(),
            "typical_jtbd": str(row[idx["typical_jtbd"]] or "").strip(),
        }
    return meta


def load_assets():
    wb = load_workbook(ASSET_XLSX, read_only=True, data_only=True)
    ws = wb["knowledge_assets"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    by_persona = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        asset_id = str(row[idx["asset_id"]] or "").strip()
        asset_type = str(row[idx["asset_type"]] or "").strip()
        title = str(row[idx["title"]] or "").strip()
        track_ids = [normalize_track(x.strip()) for x in str(row[idx["track_ids"]] or "").split(",") if x.strip()]
        persona_ids = [x.strip() for x in str(row[idx["persona_ids"]] or "").split(",") if x.strip()]
        summary = str(row[idx["summary"]] or "").strip()
        record = {
            "asset_id": asset_id,
            "asset_type": asset_type,
            "title": title,
            "track_ids": track_ids,
            "persona_ids": persona_ids,
            "summary": summary,
        }
        for persona_id in persona_ids:
            by_persona[persona_id].append(record)
    return by_persona


def choose_assets(records: list[dict], types: tuple[str, ...], limit: int = 3) -> list[str]:
    chosen = []
    for record in records:
        if record["asset_type"] in types and record["asset_id"] not in chosen:
            chosen.append(record["asset_id"])
        if len(chosen) >= limit:
            break
    return chosen


def persona_with_fallback(persona_meta: dict, assets_by_persona: dict, persona: str):
    meta = persona_meta.get(persona)
    records = assets_by_persona.get(persona, [])
    if meta or records:
        return meta, records
    fallback = PERSONA_FALLBACKS.get(persona)
    if fallback:
        return persona_meta.get(fallback), assets_by_persona.get(fallback, [])
    return None, []


def infer_mgmt_tags(track: str, persona: str) -> str:
    mapping = {
        "retail_multi_store": "mgmt_hq_operating_visibility,mgmt_frontline_action_loop",
        "retail_multi_store_chain": "mgmt_hq_operating_visibility,mgmt_frontline_action_loop",
        "retail_chain_fnb": "mgmt_hq_operating_visibility,mgmt_frontline_action_loop",
        "fnb_chain_standardized": "mgmt_hq_operating_visibility,mgmt_frontline_action_loop",
        "fnb_chain_beverage_coffee": "mgmt_hq_operating_visibility,mgmt_frontline_action_loop",
        "retail_high_sku_brand": "mgmt_inventory_supply_coordination,mgmt_profit_improvement",
        "retail_brand_beauty": "mgmt_inventory_supply_coordination,mgmt_profit_improvement",
        "retail_brand_maternal_pet": "mgmt_inventory_supply_coordination,mgmt_profit_improvement",
        "retail_fashion_group": "mgmt_hq_operating_visibility,mgmt_inventory_supply_coordination",
        "cbec_multi_platform_brand": "mgmt_profit_improvement,mgmt_group_coordination",
        "cbec_brand_outbound": "mgmt_profit_improvement,mgmt_group_coordination",
        "cbec_supply_chain_complex": "mgmt_inventory_supply_coordination,mgmt_profit_improvement",
        "cbec_platform_operator": "mgmt_profit_improvement,mgmt_hq_operating_visibility",
        "mfg_multi_factory_group": "mgmt_group_coordination",
        "mfg_rnd_sales_complex": "mgmt_group_coordination,mgmt_profit_improvement",
    }
    return mapping.get(persona, "")


def build_market_text(track: str, persona: str, reference_customers: str) -> tuple[str, str]:
    refs = [x.strip() for x in reference_customers.split(",") if x.strip()]
    similar = "、".join(refs[:3]) if len(refs) >= 2 else MARKET_PLACEHOLDER
    if len(refs) >= 2:
        competitor = f"可先参考{ '、'.join(refs[:3]) }等相近样本。"
    else:
        competitor = MARKET_PLACEHOLDER
    return similar or MARKET_PLACEHOLDER, competitor


def update_sheet(ws, persona_meta: dict, assets_by_persona: dict, target_levels: set[str]) -> int:
    headers = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}
    updated = 0
    for row in range(2, ws.max_row + 1):
        level = str(ws.cell(row, idx["静态潜客记录成熟度"]).value or "").strip()
        if level not in target_levels:
            continue
        persona = str(ws.cell(row, idx["persona_tag"]).value or "").strip()
        track = str(ws.cell(row, idx["primary_track"]).value or "").strip()
        meta, records = persona_with_fallback(persona_meta, assets_by_persona, persona)
        if not meta:
            continue

        if idx.get("management_persona_tags"):
            current = str(ws.cell(row, idx["management_persona_tags"]).value or "").strip()
            inferred = infer_mgmt_tags(track, persona)
            if not current and inferred:
                ws.cell(row, idx["management_persona_tags"]).value = inferred
                updated += 1

        if idx.get("knowledge_asset_refs"):
            current = str(ws.cell(row, idx["knowledge_asset_refs"]).value or "").strip()
            if not current:
                chosen = choose_assets(records, ("customer_case", "industry_insight", "scenario_pack", "solution_playbook"), limit=3)
                if chosen:
                    ws.cell(row, idx["knowledge_asset_refs"]).value = ",".join(chosen)
                    updated += 1
        if idx.get("talk_track_refs"):
            current = str(ws.cell(row, idx["talk_track_refs"]).value or "").strip()
            if not current:
                chosen = choose_assets(records, ("sales_narrative",), limit=2)
                if chosen:
                    ws.cell(row, idx["talk_track_refs"]).value = ",".join(chosen)
                    updated += 1

        similar, competitor = build_market_text(track, persona, meta["reference_customers"])
        if idx.get("相似客户线索"):
            current = str(ws.cell(row, idx["相似客户线索"]).value or "").strip()
            if current in {"", MARKET_PLACEHOLDER, "待补充"} and similar != MARKET_PLACEHOLDER:
                ws.cell(row, idx["相似客户线索"]).value = similar
                updated += 1
        if idx.get("主要竞品概述"):
            current = str(ws.cell(row, idx["主要竞品概述"]).value or "").strip()
            if current in {"", MARKET_PLACEHOLDER, "待补充"} and competitor != MARKET_PLACEHOLDER:
                ws.cell(row, idx["主要竞品概述"]).value = competitor
                updated += 1
    return updated


def main():
    persona_meta = load_persona_meta()
    assets_by_persona = load_assets()

    profile_wb = load_workbook(PROFILE_XLSX)
    profile_updated = update_sheet(profile_wb["account_profiles"], persona_meta, assets_by_persona, TARGET_LEVELS)
    profile_wb.save(PROFILE_XLSX)

    main_wb = load_workbook(MAIN_XLSX)
    main_updated = update_sheet(main_wb["accounts_main"], persona_meta, assets_by_persona, TARGET_LEVELS)
    main_wb.save(MAIN_XLSX)

    print({
        "profile_updated": profile_updated,
        "main_updated": main_updated,
        "target_levels": sorted(TARGET_LEVELS),
    })


if __name__ == "__main__":
    main()
