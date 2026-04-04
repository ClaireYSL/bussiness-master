from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


TODAY = "2026-03-31"
ROOT = Path("/Users/clairaipartner")
WORKSPACE = ROOT / ".openclaw/workspace-main/bussiness-master"
VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"

MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
GOV_XLSX = VAULT / "治理与证据.xlsx"
DOC_PATH = WORKSPACE / "docs/03-执行与校验/静态潜客池-L3到L2批量上移专项-v5.md"
MEMORY_PATH = WORKSPACE / "memory/2026-03-31.md"

CONSUMER_PERSONAS = {
    "retail_brand_beauty",
    "retail_brand_maternal_pet",
    "retail_fashion_group",
    "retail_multi_store",
    "retail_high_sku_brand",
}
BLOCKED_IDS = {"acc_dreame", "acc_tomtop"}


def main() -> None:
    main_wb = load_workbook(MAIN_XLSX)
    main_ws = main_wb["accounts_main"]
    mh = [c.value for c in main_ws[1]]
    mi = {h: i + 1 for i, h in enumerate(mh)}

    profile_wb = load_workbook(PROFILE_XLSX)
    profile_ws = profile_wb["account_profiles"]
    coverage_ws = profile_wb["profile_coverage"]
    ph = [c.value for c in profile_ws[1]]
    pi = {h: i + 1 for i, h in enumerate(ph)}
    ch = [c.value for c in coverage_ws[1]]
    ci = {h: i + 1 for i, h in enumerate(ch)}

    gov_wb = load_workbook(GOV_XLSX)
    evidence_ws = gov_wb["evidence_log"]
    review_ws = gov_wb["review_queue"]
    eh = [c.value for c in evidence_ws[1]]
    ei = {h: i + 1 for i, h in enumerate(eh)}
    qh = [c.value for c in review_ws[1]]
    qi = {h: i + 1 for i, h in enumerate(qh)}

    main_rows = {}
    for r in range(2, main_ws.max_row + 1):
        aid = str(main_ws.cell(r, mi["account_id"]).value or "")
        if aid:
            main_rows[aid] = r

    coverage_rows = {}
    for r in range(2, coverage_ws.max_row + 1):
        aid = str(coverage_ws.cell(r, ci["account_id"]).value or "")
        if aid:
            coverage_rows[aid] = r

    review_rows = []
    for r in range(2, review_ws.max_row + 1):
        review_rows.append(
            {
                "row": r,
                "queue_type": str(review_ws.cell(r, qi["queue_type"]).value or ""),
                "account_id": str(review_ws.cell(r, qi["account_id"]).value or ""),
                "status": str(review_ws.cell(r, qi["status"]).value or ""),
            }
        )

    promoted = []
    by_reason = Counter()

    for r in range(2, profile_ws.max_row + 1):
        aid = str(profile_ws.cell(r, pi["account_id"]).value or "")
        if aid in BLOCKED_IDS:
            continue
        level = str(profile_ws.cell(r, pi["静态潜客记录成熟度"]).value or "")
        if level != "L3":
            continue

        info = str(profile_ws.cell(r, pi["信息扎实度"]).value or "")
        icp = str(profile_ws.cell(r, pi["ICP匹配概率"]).value or "")
        official = int(profile_ws.cell(r, pi["official_source_count"]).value or 0)
        high_conf = int(profile_ws.cell(r, pi["high_confidence_source_count"]).value or 0)
        profile_status = str(profile_ws.cell(r, pi["profile_status"]).value or "")
        track = str(profile_ws.cell(r, pi["primary_track"]).value or "")
        persona = str(profile_ws.cell(r, pi["persona_tag"]).value or "")
        static_priority = str(profile_ws.cell(r, pi["static_priority"]).value or "")

        if info not in ("中高", "高"):
            continue
        if official < 1 or high_conf < 2:
            continue
        if profile_status != "standard_ready":
            continue

        reason = None
        if icp in ("中高", "高"):
            reason = "already_qualified"
        elif (
            icp == "中"
            and track == "零售消费"
            and persona in CONSUMER_PERSONAS
            and static_priority in ("A", "B")
        ):
            reason = "consumer_reassessed"
        else:
            continue

        mr = main_rows.get(aid)
        if not mr:
            continue

        if reason == "consumer_reassessed":
            profile_ws.cell(r, pi["ICP匹配概率"]).value = "中高"
            main_ws.cell(mr, mi["ICP匹配概率"]).value = "中高"

        profile_ws.cell(r, pi["静态潜客记录成熟度"]).value = "L2"
        profile_ws.cell(r, pi["static_maturity_level"]).value = "L2"
        profile_ws.cell(r, pi["profile_status"]).value = "high_quality_ready"
        profile_ws.cell(r, pi["last_profiled_at"]).value = TODAY
        validation_gap = str(profile_ws.cell(r, pi["validation_gap"]).value or "")
        if "已进入L2" not in validation_gap:
            prefix = "已进入L2；后续重点继续补官网、年报、IR和财报口径字段。"
            profile_ws.cell(r, pi["validation_gap"]).value = f"{prefix}{validation_gap}" if validation_gap else prefix

        main_ws.cell(mr, mi["静态潜客记录成熟度"]).value = "L2"
        main_ws.cell(mr, mi["信息扎实度"]).value = info
        src = str(main_ws.cell(mr, mi["source_note"]).value or "")
        if "2026-03-31 L3到L2批量上移" not in src:
            main_ws.cell(mr, mi["source_note"]).value = f"{src}；2026-03-31 L3到L2批量上移".strip("；")
        gap_main = str(main_ws.cell(mr, mi["validation_gap"]).value or "")
        if "已进入L2" not in gap_main:
            prefix = "已进入L2；后续重点继续补官网、年报、IR和财报口径字段。"
            main_ws.cell(mr, mi["validation_gap"]).value = f"{prefix}{gap_main}" if gap_main else prefix

        cr = coverage_rows.get(aid)
        if cr:
            coverage_ws.cell(cr, ci["target_scope"]).value = "L1-L2"
            coverage_ws.cell(cr, ci["official_source_ready"]).value = "yes"
            coverage_ws.cell(cr, ci["high_confidence_ready"]).value = "yes"
            coverage_ws.cell(cr, ci["profile_complete_status"]).value = "high_quality_ready"
            coverage_ws.cell(cr, ci["next_action"]).value = "继续补官网、年报、IR 和财报口径字段；必要时再评估是否进入 L1。"
            coverage_ws.cell(cr, ci["missing_core_fields"]).value = "收入规模、利润状态、营收增长仍需回到财报/年报/IR 口径继续补齐。"

        evidence_id = f"ev_{aid}_promote_l2_20260331"
        exists = False
        for er in range(2, evidence_ws.max_row + 1):
            if str(evidence_ws.cell(er, ei["evidence_id"]).value or "") == evidence_id:
                exists = True
                break
        if not exists:
            evidence_ws.append(
                [
                    evidence_id,
                    aid,
                    "promotion_assessment",
                    "internal://promotion/l3_to_l2/2026-03-31",
                    "B",
                    "persona_support,share_support",
                    "已满足 L2 门槛：信息扎实度、官方源、高可信源和档案完整度均已达标；本轮批量上移至 L2。",
                    "codex_llm",
                    TODAY,
                    "",
                    "静态潜客记录成熟度",
                    "L2",
                ]
            )

        for q in review_rows:
            if q["account_id"] == aid and q["queue_type"] == "promotion_review" and q["status"] == "open":
                review_ws.cell(q["row"], qi["status"]).value = "resolved"
                review_ws.cell(q["row"], qi["resolved_at"]).value = TODAY

        promoted.append((aid, str(profile_ws.cell(r, pi["account_canonical_name"]).value or ""), reason))
        by_reason[reason] += 1

    main_wb.save(MAIN_XLSX)
    profile_wb.save(PROFILE_XLSX)
    gov_wb.save(GOV_XLSX)

    l2_total = 0
    l3_total = 0
    for r in range(2, main_ws.max_row + 1):
        lvl = str(main_ws.cell(r, mi["静态潜客记录成熟度"]).value or "")
        if lvl == "L2":
            l2_total += 1
        elif lvl == "L3":
            l3_total += 1

    DOC_PATH.write_text(
        "\n".join(
            [
                "# 静态潜客池-L3到L2批量上移专项-v5",
                "",
                f"- 执行日期：`{TODAY}`",
                f"- 本轮上移：`{len(promoted)}` 家 `L3 -> L2`。",
                f"- 其中原本已满足门槛：`{by_reason['already_qualified']}` 家。",
                f"- 其中基于消费品新画像与官方源补强重判 ICP 后上移：`{by_reason['consumer_reassessed']}` 家。",
                f"- 当前主表口径：`L2={l2_total} / L3={l3_total}`。",
                "",
                "## 本轮上移门槛",
                "",
                "- `信息扎实度 ∈ {中高, 高}`",
                "- `official_source_count >= 1`",
                "- `high_confidence_source_count >= 2`",
                "- `profile_status = standard_ready`",
                "- 原本已满足 `ICP匹配概率 ∈ {中高, 高}`，或属于本轮完成官方源补强的消费品 active 画像并完成重判。",
                "",
                "## 说明",
                "",
                "- 本轮不扩 `L1`，继续保持 `L1` 为锚点层。",
                "- 剩余未上移对象应继续留在 `L3`，待补更强主体与官方源后再评估。",
            ]
        ),
        encoding="utf-8",
    )

    with MEMORY_PATH.open("a", encoding="utf-8") as f:
        f.write(
            f"\n- 完成一轮 `L3 -> L2` 批量上移，实际新增上移 `{len(promoted)}` 家；其中 `{by_reason['consumer_reassessed']}` 家基于消费品新画像与官方源补强重判后进入 `L2`。\n"
        )

    print(
        {
            "promoted": len(promoted),
            "by_reason": dict(by_reason),
            "sample": promoted[:40],
        }
    )


if __name__ == "__main__":
    main()
