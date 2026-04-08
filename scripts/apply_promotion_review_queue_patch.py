from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

VAULT = Path.home() / "Documents/Obsidian-Codex/潜客池"
GOV_XLSX = VAULT / "治理与证据.xlsx"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Open promotion_review queue items from a JSON patch.")
    parser.add_argument("--queue-patch-file", required=True, help="Queue patch JSON payload.")
    parser.add_argument("--output-file", help="Optional JSON summary output.")
    return parser


def _clean(value: object) -> str:
    return str(value or "").strip()


def build_header_index(ws) -> dict[str, int]:
    headers = [cell.value for cell in ws[1]]
    return {str(value): idx + 1 for idx, value in enumerate(headers) if value}


def main() -> int:
    args = build_parser().parse_args()
    payload = json.loads(Path(args.queue_patch_file).read_text(encoding="utf-8"))

    wb = load_workbook(GOV_XLSX)
    ws = wb["review_queue"]
    header_index = build_header_index(ws)
    created = 0
    reopened = 0

    for item in payload.get("accounts") or []:
        account_id = _clean(item.get("account_id"))
        queue_type = _clean(item.get("queue_type"))
        target_row = None
        for row in range(2, ws.max_row + 1):
            if _clean(ws.cell(row, header_index["account_id"]).value) != account_id:
                continue
            if _clean(ws.cell(row, header_index["queue_type"]).value) != queue_type:
                continue
            target_row = row
            break

        if target_row is None:
            target_row = ws.max_row + 1
            ws.cell(target_row, header_index["queue_item_id"]).value = f"{payload['batch_id']}_{account_id}_{queue_type}"
            created += 1
        else:
            reopened += 1

        ws.cell(target_row, header_index["queue_type"]).value = queue_type
        ws.cell(target_row, header_index["account_id"]).value = account_id
        ws.cell(target_row, header_index["priority"]).value = item.get("priority") or "P1"
        ws.cell(target_row, header_index["status"]).value = item.get("status") or "open"
        ws.cell(target_row, header_index["owner"]).value = item.get("owner") or "codex"
        ws.cell(target_row, header_index["note"]).value = item.get("note") or ""
        ws.cell(target_row, header_index["created_at"]).value = datetime.now().strftime("%Y-%m-%d")
        if "resolved_at" in header_index:
            ws.cell(target_row, header_index["resolved_at"]).value = ""

    wb.save(GOV_XLSX)
    result = {
        "batch_id": payload.get("batch_id"),
        "created": created,
        "reopened_or_updated": reopened,
        "account_count": len(payload.get("accounts") or []),
    }
    if args.output_file:
        Path(args.output_file).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
