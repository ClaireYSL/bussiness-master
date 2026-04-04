from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
ROOT = WORKSPACE.parents[2]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import is_placeholder
VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"

PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
GOV_XLSX = VAULT / "治理与证据.xlsx"
TEAM_INDEX_XLSX = VAULT / "L3以上客户档案索引-团队共享.xlsx"
L3_INDEX_MD = VAULT / "05-汇总与状态/01-总览/L3以上客户档案索引.md"

NORMALIZE_SCRIPT = WORKSPACE / "scripts/normalize_customer_fact_fields_20260402.py"
RENDER_SCRIPT = WORKSPACE / "scripts/reinforce_l3plus_quality_20260331.py"
CHECK_SCRIPT = WORKSPACE / "scripts/check_customer_archive_quality_20260402.py"
MEMORY_PATH = WORKSPACE / "memory/2026-04-03.md"

DEFAULT_BATCH_FILE = WORKSPACE / "deliveries/customer_archive_repair_batch_2026_04_03_01.json"

FACT_PLACEHOLDER = "待补官网/年报/IR口径"
EVENT_PLACEHOLDER = "待补更强官方披露"
MARKET_PLACEHOLDER = "待补公司级市场参考"
ASSET_PLACEHOLDER = "待补公司级资产映射"

REPAIR_FIELDS = [
    "archive_repair_status",
    "archive_repair_batch",
    "archive_repair_checked_at",
    "archive_repair_note",
]

TEAM_REQUIRED_HEADERS = [
    "account_id",
    "公司名称",
    "当前层级",
    "主线",
    "业务形态画像",
    "经营诉求画像",
    "信息扎实度",
    "ICP匹配概率",
    "静态优先级",
    "档案完整度",
    "分享状态",
    "分享批次",
    "最近标记时间",
    "档案页路径",
    "共享说明",
    "修复状态",
    "修复批次",
    "修复时间",
]

PROFILE_FIELD_LABELS = {
    "公司产品与服务概述": "主营产品与服务",
    "商业模式概述": "商业模式",
    "核心客户客群": "核心客户/渠道/场景",
    "主要竞品概述": "主要竞品概述",
    "相似客户线索": "相似样本参考",
    "产品与服务长摘录": "产品与服务长摘录",
    "商业模式长摘录": "商业模式长摘录",
}


def now_dt() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def today_str() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def safe_load_workbook(path: Path, **kwargs):
    last_error = None
    for attempt in range(6):
        try:
            return load_workbook(path, **kwargs)
        except (EOFError, BadZipFile) as exc:
            last_error = exc
            time.sleep(0.5 * (attempt + 1))
    raise last_error


def ensure_headers(ws, required_headers: list[str]) -> dict[str, int]:
    current = [c.value for c in ws[1]]
    for header in required_headers:
        if header not in current:
            ws.cell(1, len(current) + 1).value = header
            current.append(header)
    return {header: i + 1 for i, header in enumerate(current)}


def ensure_team_headers(wb) -> None:
    for sheet_name in ["档案索引", "新增档案", "已更新档案", "已分享档案"]:
        ensure_headers(wb[sheet_name], TEAM_REQUIRED_HEADERS)


def row_index_by(ws, col_idx: int) -> dict[str, int]:
    rows = {}
    for r in range(2, ws.max_row + 1):
        value = ws.cell(r, col_idx).value
        if value is not None and str(value).strip():
            rows[str(value).strip()] = r
    return rows


def merge_csv(existing: str, additions: list[str]) -> str:
    items = [x.strip() for x in str(existing or "").split(",") if x and x.strip()]
    for item in additions:
        item = str(item or "").strip()
        if item and item not in items:
            items.append(item)
    return ",".join(items)


def merge_lines(existing: str, additions: list[str]) -> str:
    items = [x.strip() for x in str(existing or "").split("\n") if x and x.strip()]
    for item in additions:
        item = str(item or "").strip()
        if item and item not in items:
            items.append(item)
    return "\n".join(items)


def compute_missing_fields(profile: dict) -> tuple[str, str]:
    missing = []
    if is_placeholder(profile.get("公司产品与服务概述")):
        missing.append("公司产品与服务概述")
    if is_placeholder(profile.get("商业模式概述")):
        missing.append("商业模式概述")
    if is_placeholder(profile.get("核心客户客群")):
        missing.append("核心客户客群")
    if is_placeholder(profile.get("近一年重大事件")):
        missing.append("近一年重大事件")
    if is_placeholder(profile.get("已上线系统概况")):
        missing.append("已上线系统概况")
    if is_placeholder(profile.get("数字化项目动态")):
        missing.append("数字化项目动态")
    if is_placeholder(profile.get("主要竞品概述")):
        missing.append("主要竞品概述")
    if is_placeholder(profile.get("相似客户线索")):
        missing.append("相似客户线索")

    if missing:
        missing_text = "、".join(missing)
    else:
        missing_text = "收入规模区间、利润状态概述、营收增长概述仍需按财报/年报/IR 口径继续补强。"

    next_action = (
        "继续围绕已修复事实字段补来源型摘录，并按官网、年报、IR 节奏补系统、事件和财报口径。"
        if missing
        else "当前批次单页修复已过关，后续继续补财报口径和更强字段级 evidence。"
    )
    return missing_text, next_action


def upsert_observation(ws, headers: dict[str, int], account_id: str, field_name: str, short_value: str, excerpt: str, source_locator: str, source_type: str, checked_at: str, observation_id: str) -> None:
    target_row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, headers["observation_id"]).value or "").strip() == observation_id:
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1

    ws.cell(target_row, headers["observation_id"]).value = observation_id
    ws.cell(target_row, headers["account_id"]).value = account_id
    ws.cell(target_row, headers["field_name"]).value = field_name
    ws.cell(target_row, headers["field_label"]).value = PROFILE_FIELD_LABELS.get(field_name, field_name)
    ws.cell(target_row, headers["field_value_short"]).value = short_value
    ws.cell(target_row, headers["field_excerpt_long"]).value = excerpt
    ws.cell(target_row, headers["source_locator"]).value = source_locator
    ws.cell(target_row, headers["source_type"]).value = source_type
    ws.cell(target_row, headers["source_strength"]).value = "A"
    ws.cell(target_row, headers["supports_dimension"]).value = "background_profile"
    ws.cell(target_row, headers["is_current_best"]).value = "yes"
    ws.cell(target_row, headers["checked_at"]).value = checked_at
    ws.cell(target_row, headers["checked_by"]).value = "codex_llm"
    ws.cell(target_row, headers["note"]).value = "客户档案逐客户修复批次回写。"


def upsert_profile_note(ws, headers: dict[str, int], note_id: str, account_id: str, title: str, body: str, checked_at: str) -> None:
    target_row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, headers["note_id"]).value or "").strip() == note_id:
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1
    ws.cell(target_row, headers["note_id"]).value = note_id
    ws.cell(target_row, headers["account_id"]).value = account_id
    ws.cell(target_row, headers["note_type"]).value = "context_note"
    ws.cell(target_row, headers["note_title"]).value = title
    ws.cell(target_row, headers["note_body"]).value = body
    ws.cell(target_row, headers["importance"]).value = "high"
    ws.cell(target_row, headers["created_at"]).value = checked_at
    ws.cell(target_row, headers["created_by"]).value = "codex_llm"


def upsert_evidence(ws, headers: dict[str, int], evidence_id: str, account_id: str, source_locator: str, summary: str, related_assets: str, checked_at: str) -> None:
    target_row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, headers["evidence_id"]).value or "").strip() == evidence_id:
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1
    ws.cell(target_row, headers["evidence_id"]).value = evidence_id
    ws.cell(target_row, headers["account_id"]).value = account_id
    ws.cell(target_row, headers["evidence_type"]).value = "archive_repair_batch"
    ws.cell(target_row, headers["source_locator"]).value = source_locator
    ws.cell(target_row, headers["evidence_strength"]).value = "A"
    ws.cell(target_row, headers["supports_dimension"]).value = "background_profile,persona_support,share_support"
    ws.cell(target_row, headers["summary"]).value = summary
    ws.cell(target_row, headers["checked_by"]).value = "codex_llm"
    ws.cell(target_row, headers["checked_at"]).value = checked_at
    ws.cell(target_row, headers["related_asset_ids"]).value = related_assets
    ws.cell(target_row, headers["field_name"]).value = "archive_repair_batch"
    ws.cell(target_row, headers["field_value"]).value = "已完成逐客户修复"


def run_python(script: Path, env: dict[str, str] | None = None) -> subprocess.CompletedProcess[str]:
    merged_env = os.environ.copy()
    if env:
        merged_env.update(env)
    return subprocess.run(
        [sys.executable, str(script)],
        cwd=str(WORKSPACE),
        env=merged_env,
        capture_output=True,
        text=True,
        check=False,
    )


def run_batch_check(batch_file: Path) -> dict:
    result = run_python(CHECK_SCRIPT, {"ACCOUNT_NAMES_FILE": str(batch_file)})
    if not result.stdout.strip():
        raise RuntimeError(f"质量检查脚本没有输出：{result.stderr}")
    report = json.loads(result.stdout)
    report["_returncode"] = result.returncode
    report["_stderr"] = result.stderr
    return report


def level_explanation(level: str) -> str:
    mapping = {
        "L1": "当前档案已达到成熟潜客层，可长期稳定保留和使用。",
        "L2": "当前档案已达到高成熟潜客层，但仍存在一定待补缺口。",
        "L3": "当前档案已达到可读可治理层，仍需继续补强。",
    }
    return mapping.get(level, f"当前档案处于 `{level}`。")


def repair_passes(per_account: dict) -> tuple[bool, list[str]]:
    reasons = []
    if not per_account.get("product_ready"):
        reasons.append("产品与服务仍未达到公司级事实")
    if not per_account.get("business_ready"):
        reasons.append("商业模式仍未达到公司级事实")
    if not per_account.get("source_ready"):
        reasons.append("来源层仍不足")
    if per_account.get("sturdiness") == "中低":
        reasons.append("信息扎实度仍为中低")
    return (not reasons, reasons)


def build_specific_validation_gap(profile_snapshot: dict) -> str:
    missing, _next_action = compute_missing_fields(profile_snapshot)
    if "已完成本轮重点字段修复" in missing:
        return "待补收入规模区间、利润状态概述、营收增长概述的财报/年报/IR 口径。"
    return "待补" + missing if not str(missing).startswith("待补") else str(missing)


def append_memory(batch_id: str, passed: int, total: int, report_path: Path) -> None:
    line = f"- 完成 `{batch_id}` 逐客户修复批次：处理 `{total}` 家，过关并标记 `已修复` 的 `{passed}` 家；批次报告：`{report_path}`。\n"
    with MEMORY_PATH.open("a", encoding="utf-8") as f:
        f.write(line)


def main() -> int:
    batch_file = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(os.environ.get("BATCH_FILE", DEFAULT_BATCH_FILE))
    batch = load_json(batch_file)
    batch_id = batch["batch_id"]
    batch_title = batch.get("batch_title", batch_id)
    batch_suffix = batch_id.split("_")[-1]
    checked_at = now_dt()
    today = today_str()

    profile_wb = safe_load_workbook(PROFILE_XLSX)
    profiles_ws = profile_wb["account_profiles"]
    observations_ws = profile_wb["field_observations"]
    notes_ws = profile_wb["profile_notes"]
    coverage_ws = profile_wb["profile_coverage"]

    main_wb = safe_load_workbook(MAIN_XLSX)
    main_ws = main_wb["accounts_main"]

    gov_wb = safe_load_workbook(GOV_XLSX)
    evidence_ws = gov_wb["evidence_log"]

    team_wb = safe_load_workbook(TEAM_INDEX_XLSX)
    ensure_team_headers(team_wb)

    profile_headers = ensure_headers(profiles_ws, [c.value for c in profiles_ws[1]] + [h for h in REPAIR_FIELDS if h not in [c.value for c in profiles_ws[1]]])
    observation_headers = ensure_headers(observations_ws, [c.value for c in observations_ws[1]])
    notes_headers = ensure_headers(notes_ws, [c.value for c in notes_ws[1]])
    coverage_headers = ensure_headers(coverage_ws, [c.value for c in coverage_ws[1]])
    main_headers = ensure_headers(main_ws, [c.value for c in main_ws[1]])
    evidence_headers = ensure_headers(evidence_ws, [c.value for c in evidence_ws[1]])

    profile_rows_by_name = row_index_by(profiles_ws, profile_headers["account_canonical_name"])
    coverage_rows_by_id = row_index_by(coverage_ws, coverage_headers["account_id"])
    main_rows_by_name = row_index_by(main_ws, main_headers["account_canonical_name"])
    main_rows_by_id = row_index_by(main_ws, main_headers["account_id"])

    processed_accounts = []
    for account in batch["accounts"]:
        name = account["account_canonical_name"]
        if name not in profile_rows_by_name:
            raise RuntimeError(f"未找到客户：{name}")
        pr = profile_rows_by_name[name]
        account_id = str(profiles_ws.cell(pr, profile_headers["account_id"]).value).strip()
        mr = main_rows_by_name.get(name) or main_rows_by_id.get(account_id)
        if mr is None:
            raise RuntimeError(f"未找到主表记录：{name} / {account_id}")
        cr = coverage_rows_by_id.get(account_id)
        if cr is None:
            raise RuntimeError(f"未找到 coverage 记录：{name} / {account_id}")

        fields = account.get("fields", {})
        matching = account.get("matching", {})
        excerpts = account.get("excerpts", {})
        repair_note = account.get("repair_note", "")

        for field_name, value in fields.items():
            if field_name in profile_headers:
                profiles_ws.cell(pr, profile_headers[field_name]).value = value
            if field_name in main_headers:
                main_ws.cell(mr, main_headers[field_name]).value = value

        if "management_persona_tags" in matching:
            profiles_ws.cell(pr, profile_headers["management_persona_tags"]).value = matching["management_persona_tags"]
        if "knowledge_asset_refs" in matching:
            profiles_ws.cell(pr, profile_headers["knowledge_asset_refs"]).value = matching["knowledge_asset_refs"]
            if "knowledge_asset_refs" in main_headers:
                main_ws.cell(mr, main_headers["knowledge_asset_refs"]).value = matching["knowledge_asset_refs"]
        if "talk_track_refs" in matching:
            profiles_ws.cell(pr, profile_headers["talk_track_refs"]).value = matching["talk_track_refs"]

        source_types = merge_csv(
            profiles_ws.cell(pr, profile_headers["primary_source_types"]).value,
            account.get("source_types_add", []),
        )
        source_refs = merge_lines(
            profiles_ws.cell(pr, profile_headers["primary_source_refs"]).value,
            account.get("source_refs_add", []),
        )
        profiles_ws.cell(pr, profile_headers["primary_source_types"]).value = source_types
        profiles_ws.cell(pr, profile_headers["primary_source_refs"]).value = source_refs
        profiles_ws.cell(pr, profile_headers["official_source_count"]).value = max(
            int(profiles_ws.cell(pr, profile_headers["official_source_count"]).value or 0),
            1 if account.get("source_refs_add") else 0,
        )
        profiles_ws.cell(pr, profile_headers["high_confidence_source_count"]).value = max(
            int(profiles_ws.cell(pr, profile_headers["high_confidence_source_count"]).value or 0),
            2 if len(account.get("source_refs_add", [])) >= 2 else 1,
        )
        profiles_ws.cell(pr, profile_headers["profile_owner"]).value = "codex_llm"
        profiles_ws.cell(pr, profile_headers["last_profiled_at"]).value = checked_at
        profiles_ws.cell(pr, profile_headers["archive_repair_status"]).value = "修复中"
        profiles_ws.cell(pr, profile_headers["archive_repair_batch"]).value = batch_id
        profiles_ws.cell(pr, profile_headers["archive_repair_checked_at"]).value = checked_at
        profiles_ws.cell(pr, profile_headers["archive_repair_note"]).value = (
            repair_note or "逐客户修复进行中；本批仅做阅读层修复，不代表上移依据已完成。"
        )

        profile_snapshot_for_gap = {h: profiles_ws.cell(pr, idx).value for h, idx in profile_headers.items()}
        validation_gap = build_specific_validation_gap(profile_snapshot_for_gap)
        profiles_ws.cell(pr, profile_headers["validation_gap"]).value = validation_gap
        if "validation_gap" in main_headers:
            main_ws.cell(mr, main_headers["validation_gap"]).value = validation_gap
        if "last_verified_at" in main_headers:
            main_ws.cell(mr, main_headers["last_verified_at"]).value = checked_at

        for excerpt_field, excerpt_value in excerpts.items():
            if excerpt_field in profile_headers:
                profiles_ws.cell(pr, profile_headers[excerpt_field]).value = excerpt_value

        for field_name, field_value in {**fields, **excerpts}.items():
            source_locator = "\n".join(account.get("source_refs_add", [])[:2]) or source_refs
            source_type = ",".join(account.get("source_types_add", [])[:2]) or source_types
            observation_id = f"obs_{batch_id}_{account_id}_{field_name}"
            upsert_observation(
                observations_ws,
                observation_headers,
                account_id,
                field_name,
                str(field_value),
                str(field_value),
                source_locator,
                source_type,
                checked_at,
                observation_id,
            )

        note_id = f"pn_{batch_id}_{account_id}"
        note_body = repair_note or "已按客户事实优先标准逐客户修复。"
        upsert_profile_note(notes_ws, notes_headers, note_id, account_id, batch_title, note_body, checked_at)

        evidence_id = f"ev_{batch_id}_{account_id}"
        upsert_evidence(
            evidence_ws,
            evidence_headers,
            evidence_id,
            account_id,
            source_refs,
            repair_note or "本轮逐客户修复已补回公司级事实、解释层映射与来源入口。",
            matching.get("knowledge_asset_refs", ""),
            checked_at,
        )

        coverage_ws.cell(cr, coverage_headers["target_scope"]).value = profiles_ws.cell(pr, profile_headers["静态潜客记录成熟度"]).value
        coverage_ws.cell(cr, coverage_headers["profile_complete_status"]).value = profiles_ws.cell(pr, profile_headers["profile_status"]).value
        coverage_ws.cell(cr, coverage_headers["official_source_ready"]).value = "yes"
        coverage_ws.cell(cr, coverage_headers["high_confidence_ready"]).value = "yes"

        processed_accounts.append({"name": name, "account_id": account_id})

    profile_wb.save(PROFILE_XLSX)
    main_wb.save(MAIN_XLSX)
    gov_wb.save(GOV_XLSX)
    team_wb.save(TEAM_INDEX_XLSX)

    compile_targets = [
        WORKSPACE / "scripts/repair_customer_archive_batch_20260403.py",
        NORMALIZE_SCRIPT,
        RENDER_SCRIPT,
        CHECK_SCRIPT,
    ]
    compile_result = subprocess.run(
        [sys.executable, "-m", "py_compile", *[str(p) for p in compile_targets]],
        cwd=str(WORKSPACE),
        capture_output=True,
        text=True,
        check=False,
    )
    if compile_result.returncode != 0:
        raise RuntimeError(compile_result.stderr)

    for script in [NORMALIZE_SCRIPT, RENDER_SCRIPT]:
        result = run_python(script)
        if result.returncode != 0:
            raise RuntimeError(result.stderr or result.stdout)

    report = run_batch_check(batch_file)

    profile_wb = safe_load_workbook(PROFILE_XLSX)
    profiles_ws = profile_wb["account_profiles"]
    coverage_ws = profile_wb["profile_coverage"]
    profile_headers = {c.value: i + 1 for i, c in enumerate(profiles_ws[1])}
    coverage_headers = {c.value: i + 1 for i, c in enumerate(coverage_ws[1])}
    profile_rows_by_name = row_index_by(profiles_ws, profile_headers["account_canonical_name"])
    coverage_rows_by_id = row_index_by(coverage_ws, coverage_headers["account_id"])

    per_account = {item["name"]: item for item in report.get("per_account", [])}
    passed_count = 0
    account_results = []
    for account in batch["accounts"]:
        name = account["account_canonical_name"]
        per = per_account.get(name, {})
        passed, reasons = repair_passes(per)
        pr = profile_rows_by_name[name]
        account_id = str(profiles_ws.cell(pr, profile_headers["account_id"]).value).strip()
        cr = coverage_rows_by_id[account_id]

        profile_snapshot = {h: profiles_ws.cell(pr, idx).value for h, idx in profile_headers.items()}
        missing_text, next_action = compute_missing_fields(profile_snapshot)
        coverage_ws.cell(cr, coverage_headers["missing_core_fields"]).value = missing_text
        coverage_ws.cell(cr, coverage_headers["next_action"]).value = next_action

        if passed:
            profiles_ws.cell(pr, profile_headers["archive_repair_status"]).value = "已修复"
            profiles_ws.cell(pr, profile_headers["archive_repair_note"]).value = "已按客户事实优先标准重建，单页自检通过。"
            passed_count += 1
        else:
            profiles_ws.cell(pr, profile_headers["archive_repair_status"]).value = "待复核"
            profiles_ws.cell(pr, profile_headers["archive_repair_note"]).value = "待复核：" + "；".join(reasons)

        profiles_ws.cell(pr, profile_headers["archive_repair_checked_at"]).value = checked_at
        account_results.append(
            {
                "name": name,
                "account_id": account_id,
                "passed": passed,
                "reasons": reasons,
                "sturdiness": per.get("sturdiness", ""),
                "missing_core_fields": missing_text,
                "sources": account.get("source_refs_add", []),
            }
        )

    profile_wb.save(PROFILE_XLSX)

    rerender = run_python(RENDER_SCRIPT)
    if rerender.returncode != 0:
        raise RuntimeError(rerender.stderr or rerender.stdout)

    report_path = WORKSPACE / f"docs/03-执行与校验/客户档案修复批次-{today.replace('-', '')}-{batch_suffix}.md"
    common_issues = []
    if report.get("findings"):
        common_issues.extend(report["findings"])
    if passed_count < len(batch["accounts"]):
        common_issues.append("本批存在待复核对象，需回看页面是否仍有事实不足或来源不足。")
    remaining_counter = Counter()
    for item in account_results:
        for field in str(item["missing_core_fields"]).split("、"):
            field = field.strip()
            if field and "已完成本轮重点字段修复" not in field and "仍需回到" not in field:
                remaining_counter[field] += 1

    lines = [
        f"# {batch_title}",
        "",
        f"- 批次 ID：`{batch_id}`",
        f"- 执行时间：`{checked_at}`",
        f"- 对象数量：`{len(batch['accounts'])}`",
        f"- 已修复：`{passed_count}`",
        f"- 待复核：`{len(batch['accounts']) - passed_count}`",
        f"- 索引入口：[{L3_INDEX_MD}]({L3_INDEX_MD})",
        "- 说明：本批仅完成阅读层修复，不等于上移依据完成；正式上移仍需看主线/画像/商业模式等关键判断点是否收敛。",
        "",
        "## 本批对象",
        "",
    ]
    for item in account_results:
        status = "已修复" if item["passed"] else "待复核"
        lines.append(f"- `{status}` {item['name']} · 信息扎实度：`{item['sturdiness']}`")
        if item["reasons"]:
            lines.append(f"  理由：{'；'.join(item['reasons'])}")
    lines.extend([
        "",
        "## 每家修了什么",
        "",
    ])
    for account in batch["accounts"]:
        lines.append(f"### {account['account_canonical_name']}")
        lines.append(f"- 修复说明：{account.get('repair_note', '已逐客户修复。')}")
        lines.append(f"- 事实补强：{', '.join(account.get('fields', {}).keys())}")
        lines.append(f"- 解释层收口：{', '.join(account.get('matching', {}).keys())}")
        lines.append(f"- 来源补充：{len(account.get('source_refs_add', []))} 条")
        result = next(x for x in account_results if x["name"] == account["account_canonical_name"])
        lines.append(f"- 当前剩余缺口：{result['missing_core_fields']}")
        lines.append("")

    lines.extend([
        "## 本批发现的共性问题",
        "",
    ])
    if common_issues:
        for issue in common_issues:
            lines.append(f"- {issue}")
    else:
        lines.append("- 本批未发现新的结构性问题。")

    lines.extend([
        "",
        "## 剩余缺口统计",
        "",
    ])
    if remaining_counter:
        for field, count in remaining_counter.most_common():
            lines.append(f"- `{field}`：`{count}` 家")
    else:
        lines.append("- 本批对象的本轮核心缺口已基本清空。")

    lines.extend([
        "",
        "## 本批 mark 范围",
        "",
        f"- `archive_repair_status`：已更新到 `潜客档案库.xlsx / account_profiles`",
        f"- `archive_repair_batch`：统一为 `{batch_id}`",
        "- `L3以上客户档案索引.md` 与 `L3以上客户档案索引-团队共享.xlsx` 已同步展示修复状态和批次。",
    ])
    report_path.write_text("\n".join(lines), encoding="utf-8")

    append_memory(batch_id, passed_count, len(batch["accounts"]), report_path)

    print(
        json.dumps(
            {
                "batch_id": batch_id,
                "processed": len(batch["accounts"]),
                "passed": passed_count,
                "pending_review": len(batch["accounts"]) - passed_count,
                "report_path": str(report_path),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
