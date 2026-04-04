from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


VAULT = Path("/Users/clairaipartner/Documents/Obsidian-Codex/潜客池")
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MEMORY_PATH = Path("/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/memory/2026-04-02.md")

SAMPLES = {
    "acc_uih": {
        "公司产品与服务概述": "主营医学影像设备、放射治疗产品和生命科学仪器，并提供医疗云与综合解决方案。",
        "商业模式概述": "以高端医疗装备研发制造为核心，覆盖整机系统、核心部件、软件平台和医疗云服务，向医院与科研机构提供设备与整体解决方案。",
        "核心客户客群": "各级医疗机构、医学中心、科研机构及其他医疗服务机构。",
        "已上线系统概况": "年报已明确公司基于联影云系统架构提供联影医疗云服务，实现设备与应用云端协同及医疗资源共享。",
        "数字化项目动态": "年报明确提到 AI 大模型与影像、治疗、介入等关键场景融合，并以硬软结合、院内院外联动的整体解决方案推进产品与服务落地。",
        "主要竞品概述": "年报可见可比对象包括迈瑞医疗、万东医疗、东软医疗以及国际影像和放疗设备厂商。",
        "招聘代表岗位": "研发管理、营销运营、计划协同、数据分析。",
        "近一年重大事件": "2024 年年报显示，公司在全球贸易环境复杂和国内设备更新节奏变化背景下，继续推进 AI 大模型、医学影像和放疗等能力融合。",
        "产品与服务长摘录": "2024 年年报披露，公司已形成医学影像设备、放射治疗产品和生命科学仪器的完整产品线，截至报告期末累计推出 140 多款产品。",
        "商业模式长摘录": "年报披露，公司围绕“整机系统-核心部件-底层元器件”构建技术攻坚体系，并基于联影云系统架构提供设备与应用云端协同及医疗资源共享。",
        "客户客群长摘录": "年报中的产品用途说明显示，其磁共振、CT、XR、分子影像、放疗等产品广泛面向各级医疗机构临床诊断、治疗及科研场景。",
        "系统与数字化长摘录": "董事长致辞明确提到 AI 已成为公司设备、临床、科研与解决方案的技术底座，并与影像、放疗、信息化系统做有机融合。",
        "重大事件长摘录": "2024 年年报将全球供应链波动、国内高端医疗装备国产化、AI 大模型加速发展列为公司当期经营与战略推进的重要背景。",
        "primary_source_types": "annual_report",
        "primary_source_refs": "https://global.united-imaging.com/-/media/uih/pdf/investor/20250430-cn/united-imaging-healthcare-annual-report-for-2024.pdf",
        "official_source_count": 1,
        "high_confidence_source_count": 1,
        "validation_gap": "收入、利润与增长字段仍需继续回到财报口径细化；渠道和区域经营颗粒度也可继续补。",
        "notes": [
            ("obs_acc_uih_fact_product_20260402", "公司产品与服务概述", "公司产品与服务概述", "主营医学影像设备、放射治疗产品和生命科学仪器。", "年报披露，公司已形成医学影像设备、放射治疗产品和生命科学仪器的完整产品线，并累计推出 140 多款产品。", "https://global.united-imaging.com/-/media/uih/pdf/investor/20250430-cn/united-imaging-healthcare-annual-report-for-2024.pdf", "annual_report", "S", "background_profile"),
            ("obs_acc_uih_fact_model_20260402", "商业模式概述", "商业模式概述", "高端医疗装备研发制造与整体解决方案。", "年报披露，公司围绕整机系统、核心部件和底层元器件构建技术体系，并提供医疗云与综合解决方案。", "https://global.united-imaging.com/-/media/uih/pdf/investor/20250430-cn/united-imaging-healthcare-annual-report-for-2024.pdf", "annual_report", "S", "business_model"),
            ("obs_acc_uih_fact_system_20260402", "已上线系统概况", "已上线系统概况", "已提供联影医疗云服务。", "年报披露，公司基于联影云系统架构实现设备与应用云端协同及医疗资源共享。", "https://global.united-imaging.com/-/media/uih/pdf/investor/20250430-cn/united-imaging-healthcare-annual-report-for-2024.pdf", "annual_report", "S", "background_profile"),
        ],
    },
    "acc_cn_603288": {
        "公司产品与服务概述": "主营酱油、蚝油、调味酱、醋、料酒等调味品，属于典型的品牌调味品消费品公司。",
        "商业模式概述": "以品牌调味品研发、生产和全国渠道销售为核心，通过经销、零售终端和线上线下渠道覆盖消费市场。",
        "核心客户客群": "家庭消费者、餐饮客户、经销商与零售渠道。",
        "已上线系统概况": "目前已能确认公司具备支撑全国品牌消费品经营的供应链、渠道和财务管理基础系统，但系统名称与项目口径仍待年报/IR 补强。",
        "数字化项目动态": "当前更能稳定确认的是其品牌消费品经营复杂度和渠道广度；具体数字化项目口径仍待官网、年报或 IR 材料补充。",
        "主要竞品概述": "调味品与食品饮料领域的全国化品牌企业，重点看多品类、多渠道、供应链与利润协同。",
        "招聘代表岗位": "商品企划、渠道分析、供应链计划、财务分析、BI/数据岗位。",
        "近一年重大事件": "官网持续展示全系产品、子品牌和面向家庭及餐饮场景的产品矩阵；更细的经营事件仍待年报与 IR 披露补强。",
        "产品与服务长摘录": "海天官网首页可见全系产品橱窗，并直接展示有机酱油、海天上等蚝油、有机五谷醋等产品系列，说明其品类矩阵覆盖酱油、蚝油、醋等核心调味品。",
        "商业模式长摘录": "海天官网以品牌官网形式展示产品家族、子品牌和全国化消费场景，结合上市主体公开披露，可以稳定判断其属于多品类、多渠道经营的品牌调味品公司。",
        "客户客群长摘录": "官网产品展示同时覆盖家庭烹饪和餐饮调味场景，现阶段可先确认其客户链路至少覆盖终端消费者、餐饮场景和零售渠道。",
        "系统与数字化长摘录": "当前对海天的系统与数字化判断仍应克制：可以确认其经营规模和渠道复杂度较高，但具体系统名称、项目范围和数字化组织口径仍待年报/IR 继续补。",
        "重大事件长摘录": "海天官网目前能直接支撑的是产品家族和品牌经营范围；若要描述近一年经营动作，仍需回到年报、公告或 IR 材料，而不应继续用市值或泛化句子替代。",
        "primary_source_types": "official_website,行业分类,上市公司公开披露",
        "primary_source_refs": "https://www.haday.com/\nhttps://quote.eastmoney.com/concept/sh603288.html",
        "official_source_count": 1,
        "high_confidence_source_count": 2,
        "validation_gap": "收入、利润和增长字段仍需回到年报/IR；具体系统和组织口径也需要官方材料继续补。",
        "notes": [
            ("obs_acc_cn_603288_fact_product_20260402", "公司产品与服务概述", "公司产品与服务概述", "主营酱油、蚝油、调味酱、醋、料酒等调味品。", "海天官网首页产品橱窗直接展示有机酱油、海天上等蚝油、有机五谷醋等产品系列，可支撑其调味品品牌公司定位。", "https://www.haday.com/", "official_website", "A", "background_profile"),
            ("obs_acc_cn_603288_fact_model_20260402", "商业模式概述", "商业模式概述", "品牌调味品研发、生产和全国渠道销售。", "官网以品牌官网形式展示全系产品和子品牌，结合上市公司公开披露，可以稳定判断其经营核心是品牌调味品生产与全国渠道销售。", "https://www.haday.com/", "official_website", "A", "business_model"),
            ("obs_acc_cn_603288_fact_customer_20260402", "核心客户客群", "核心客户客群", "家庭消费者、餐饮客户、经销商与零售渠道。", "官网产品展示覆盖家庭调味和餐饮场景，现阶段可先确认其客户链路至少包含终端消费者、餐饮场景和零售渠道。", "https://www.haday.com/", "official_website", "A", "background_profile"),
        ],
    },
    "acc_anker": {
        "公司产品与服务概述": "主营智能充电、智能家居、智能影音等消费电子产品，旗下包括 Anker、Eufy、Nebula 等品牌。",
        "商业模式概述": "以自有品牌出海为核心，通过多平台、多国家和多品类经营触达全球消费者。",
        "核心客户客群": "海外消费电子用户及全球主流电商平台消费者。",
        "已上线系统概况": "官网材料可支撑其已形成跨品牌、跨品类和跨区域的全球经营基础设施；更细的系统名称和财务口径仍待年报/IR 补充。",
        "数字化项目动态": "当前可稳定确认的是其全球经营复杂度、品牌矩阵和多平台特征；具体经营分析项目仍待年报、案例和 IR 口径补强。",
        "主要竞品概述": "消费电子品牌出海企业，重点对标多品牌、多平台和全球经营能力。",
        "招聘代表岗位": "跨境运营、财务分析、供应链计划、BI/数据岗位。",
        "近一年重大事件": "官网仍持续以全球化品牌矩阵方式对外展示 Anker、Eufy、Nebula 等品牌；更细的区域经营动作仍待年报或 IR 披露。",
        "产品与服务长摘录": "安克创新官网首页描述，公司在全球市场塑造中国消费电子品牌，并已打造 Anker，同时推出 Eufy、Nebula 等智能硬件品牌。",
        "商业模式长摘录": "官网对外呈现的是多品牌、多品类的全球市场经营，而不是单一产品出口，这可以直接支撑其“多平台品牌出海型”画像。",
        "客户客群长摘录": "官网描述聚焦全球消费电子市场，当前可稳定确认其核心服务对象是海外终端消费者和全球主流电商平台用户。",
        "系统与数字化长摘录": "安克当前最能被稳定确认的是全球经营和品牌矩阵复杂度；具体站点经营、利润分析与供应链系统口径仍需继续回到年报、案例和 evidence。",
        "重大事件长摘录": "现阶段可直接引用的官方信息主要来自官网品牌矩阵与市场描述；若要继续描述年度经营动作，仍需补年报或 IR 材料。",
        "primary_source_types": "official_website",
        "primary_source_refs": "https://www.anker.com.cn/",
        "official_source_count": 1,
        "high_confidence_source_count": 1,
        "validation_gap": "平台矩阵、区域经营和财报口径仍需继续回到年报、IR 和字段级 evidence。",
        "notes": [
            ("obs_acc_anker_fact_product_20260402", "公司产品与服务概述", "公司产品与服务概述", "主营智能充电、智能家居、智能影音等消费电子产品。", "安克官网首页描述，公司在全球市场塑造中国消费电子品牌，并打造了 Anker、Eufy、Nebula 等品牌。", "https://www.anker.com.cn/", "official_website", "A", "background_profile"),
            ("obs_acc_anker_fact_model_20260402", "商业模式概述", "商业模式概述", "多品牌、多平台、全球化品牌出海。", "官网对外呈现为多品牌智能硬件企业，业务覆盖智能充电、智能家居、智能影音等多个产品方向。", "https://www.anker.com.cn/", "official_website", "A", "business_model"),
            ("obs_acc_anker_fact_customer_20260402", "核心客户客群", "核心客户客群", "海外消费电子用户及全球主流电商平台消费者。", "官网描述聚焦全球消费电子市场，可支撑其面向海外终端消费者和全球平台市场的客户定位。", "https://www.anker.com.cn/", "official_website", "A", "background_profile"),
        ],
    },
    "acc_catl": {
        "公司产品与服务概述": "主营动力电池系统、储能系统及相关新能源解决方案，属于全球化新能源制造企业。",
        "商业模式概述": "以动力电池和储能系统研发、制造与销售为核心，通过多基地制造与集团协同服务全球新能源应用市场。",
        "核心客户客群": "新能源汽车产业链客户、储能场景客户及全球新能源应用市场。",
        "已上线系统概况": "官网可支撑其属于全球新能源创新科技公司；更细的集团系统、制造执行与财务协同口径仍待年报或 IR 继续补。",
        "数字化项目动态": "当前更能稳定确认的是多基地协同、全球制造和经营驾驶舱类需求；具体数字化项目名称仍需依赖年报、IR 或案例补充。",
        "主要竞品概述": "动力电池、储能系统及多基地制造协同的新能源制造企业。",
        "招聘代表岗位": "计划、供应链、财务分析、数字化岗位。",
        "近一年重大事件": "官网持续将公司定位为全球领先的新能源创新科技公司，强调动力电池、储能系统与全球新能源应用解决方案。",
        "产品与服务长摘录": "宁德时代官网首页描述，公司是全球领先的新能源创新科技公司，致力于为全球新能源应用提供一流解决方案和服务。",
        "商业模式长摘录": "官网 meta 描述明确提到公司专注于新能源汽车动力电池系统、储能系统的研发、生产和销售，可直接支撑其新能源制造和多基地协同属性。",
        "客户客群长摘录": "基于官网表述，当前可先确认其服务对象覆盖全球新能源应用市场，具体客户结构仍待年报、IR 或一手案例材料继续补齐。",
        "系统与数字化长摘录": "宁德时代当前最能被稳定确认的是全球化新能源制造与多基地协同特征；具体系统名称、数据平台和经营驾驶舱口径仍待官方材料补充。",
        "重大事件长摘录": "当前能直接引用的官方材料主要来自官网定位和历史强核验材料；若要描述具体年度经营动作，仍需回到年报和 IR。",
        "primary_source_types": "official_website,archive,manual_note",
        "primary_source_refs": "https://www.catl.com/\n/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/docs/外部目标客户池-v1.0-第四至第六批关键账户强核验-v0.1.md",
        "official_source_count": 1,
        "high_confidence_source_count": 2,
        "validation_gap": "收入、利润、增长和更细的基地/事业部颗粒度仍需回到年报、IR 与字段级 evidence。",
        "notes": [
            ("obs_acc_catl_fact_product_20260402", "公司产品与服务概述", "公司产品与服务概述", "主营动力电池系统、储能系统及相关新能源解决方案。", "宁德时代官网描述，公司专注于新能源汽车动力电池系统、储能系统的研发、生产和销售。", "https://www.catl.com/", "official_website", "A", "background_profile"),
            ("obs_acc_catl_fact_model_20260402", "商业模式概述", "商业模式概述", "动力电池和储能系统研发制造与全球销售。", "官网将公司定位为全球领先的新能源创新科技公司，面向全球新能源应用提供解决方案和服务。", "https://www.catl.com/", "official_website", "A", "business_model"),
            ("obs_acc_catl_fact_customer_20260402", "核心客户客群", "核心客户客群", "新能源汽车产业链客户、储能场景客户。", "基于官网表述，可先确认其服务对象覆盖全球新能源应用市场，具体客户名单仍待更强官方披露继续补。", "https://www.catl.com/", "official_website", "A", "background_profile"),
        ],
    },
    "acc_l5_920123": {
        "公司产品与服务概述": "主营化妆品研发、生产、销售及检测服务，属于美妆个护产业链中的品牌与代工能力主体。",
        "商业模式概述": "以化妆品代加工和品牌产品经营为核心，兼具研发、生产和服务能力。",
        "核心客户客群": "化妆品品牌方、渠道品牌、新锐品牌及终端消费品经营客户。",
        "已上线系统概况": "官网可支撑其属于化妆品加工与品牌产品并行的经营主体；更细的 ERP、会员或财务系统名称仍待年报或 IR 补充。",
        "数字化项目动态": "当前更能稳定确认的是产品研发、代工生产和品牌经营并行；具体数字化项目仍待官网深页、年报或 IR 继续补。",
        "主要竞品概述": "美妆个护品牌、化妆品代工与高 SKU 消费品企业。",
        "招聘代表岗位": "商品企划、会员运营、渠道分析、分货补货、BI/数据岗位。",
        "近一年重大事件": "官网持续强调其化妆品加工、代加工和生产基地定位；更细的年度经营动作仍待后续补官方披露。",
        "产品与服务长摘录": "芭薇官网 meta 描述显示，公司成立于 2005 年，专注于化妆品代加工领域，并为众多知名品牌提供加工服务。",
        "商业模式长摘录": "官网首页同时展示品牌产品与化妆品加工业务，当前可稳定判断其兼具研发、生产和服务能力，而不是单一渠道品牌。",
        "客户客群长摘录": "基于官网描述，当前至少可确认其客户包括化妆品品牌方和代工需求方；更细的渠道与终端结构仍待继续补。",
        "系统与数字化长摘录": "芭薇当前最能被稳定确认的是研发、生产、销售与检测并行的经营结构；具体数据平台、会员系统和财务系统名称仍待官方材料补强。",
        "重大事件长摘录": "官网当前最直接支撑的是“成立于 2005 年、专注化妆品代加工”的公司定位；若要描述近一年经营动作，仍需继续补公告、年报或 IR。",
        "primary_source_types": "official_website,上市公司基础资料,上市公司公开披露",
        "primary_source_refs": "https://www.gzbawei.com\nhttps://quote.eastmoney.com/bj920123.html\ncninfo:920123",
        "official_source_count": 1,
        "high_confidence_source_count": 2,
        "validation_gap": "收入、利润、增长和更细的客户/渠道结构仍需回到年报、IR 和字段级 evidence。",
        "notes": [
            ("obs_acc_l5_920123_fact_product_20260402", "公司产品与服务概述", "公司产品与服务概述", "主营化妆品研发、生产、销售及检测服务。", "芭薇官网 meta 描述显示，公司专注于化妆品代加工领域，并为众多知名品牌提供加工服务。", "https://www.gzbawei.com", "official_website", "A", "background_profile"),
            ("obs_acc_l5_920123_fact_model_20260402", "商业模式概述", "商业模式概述", "化妆品代加工与品牌产品经营并行。", "官网首页同时展示品牌产品与加工能力，说明其业务不只是品牌销售，还覆盖研发、生产和服务。", "https://www.gzbawei.com", "official_website", "A", "business_model"),
            ("obs_acc_l5_920123_fact_customer_20260402", "核心客户客群", "核心客户客群", "化妆品品牌方、渠道品牌和终端消费品经营客户。", "官网描述“为众多知名品牌提供加工服务”，可先支撑其客户结构中存在品牌方与代工需求方。", "https://www.gzbawei.com", "official_website", "A", "background_profile"),
        ],
    },
    "acc_tomtop": {
        "公司产品与服务概述": "目前可稳定确认其属于跨境电商经营主体，但主营品牌、产品矩阵和自有品牌结构仍待一手材料继续补。",
        "商业模式概述": "历史候选池和补强记录只能支撑其具备多平台、多地区跨境经营属性；更细的品牌化、平台结构和履约模式仍待补充。",
        "核心客户客群": "当前只能保守确认其面向海外终端消费者和跨境平台链路，具体客群结构仍待官网或一手案例材料补充。",
        "已上线系统概况": "待补官网口径",
        "数字化项目动态": "待补官网口径",
        "主要竞品概述": "与品牌出海和多平台跨境经营主体相邻，但当前不宜再用模板化竞品描述替代公司事实。",
        "招聘代表岗位": "待补官网口径",
        "近一年重大事件": "待补官网口径",
        "产品与服务长摘录": "第十五批放量候选池仅能支撑“跨境电商经营属性明确、适合作为品牌出海相邻候选”这一层级，无法直接展开其具体产品与品牌矩阵。",
        "商业模式长摘录": "第三批放量候选池曾将其描述为多平台、多地区跨境经营画像相邻，说明其经营模式至少涉及多平台跨境链路，但仍不足以细化为具体品牌和渠道结构。",
        "客户客群长摘录": "当前关于通拓的客户事实仍然偏薄，只能保守确认其服务于跨境平台和海外消费者链路；更细的国家、站点和品牌颗粒度仍待一手材料。",
        "系统与数字化长摘录": "对于通拓，当前不应再用模板化“已具备基础系统环境”替代公司事实；目前更诚实的做法是直接承认系统与数字化口径待补官网或一手案例材料。",
        "重大事件长摘录": "现有公开材料不足以支持对其近一年经营事件做出具体描述，因此本项保留为待补，不再用泛化句子填充。",
        "primary_source_types": "archive,high_confidence_public",
        "primary_source_refs": "/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/docs/archive/外部目标客户池-v1.0/外部目标客户池-v1.0-第三批放量候选池-v0.1.md\n/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/docs/archive/外部目标客户池-v1.0/外部目标客户池-v1.0-第十五批放量候选池-v0.1.md",
        "official_source_count": 0,
        "high_confidence_source_count": 1,
        "validation_gap": "当前仍缺稳定官网、官方介绍页或一手案例材料；在补到这些来源前，保持 L3 且不继续上移。",
        "notes": [
            ("obs_acc_tomtop_fact_product_20260402", "公司产品与服务概述", "公司产品与服务概述", "跨境电商经营主体，具体产品矩阵待补。", "第十五批放量候选池只能支撑其跨境电商经营属性明确，尚不足以展开具体品牌和产品矩阵。", "/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/docs/archive/外部目标客户池-v1.0/外部目标客户池-v1.0-第十五批放量候选池-v0.1.md", "archive", "B", "background_profile"),
            ("obs_acc_tomtop_fact_model_20260402", "商业模式概述", "商业模式概述", "多平台、多地区跨境经营属性相邻。", "第三批放量候选池曾将其标为多平台、多地区跨境经营画像相邻，说明其经营模式至少涉及多平台跨境链路。", "/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/docs/archive/外部目标客户池-v1.0/外部目标客户池-v1.0-第三批放量候选池-v0.1.md", "archive", "B", "business_model"),
            ("obs_acc_tomtop_fact_gap_20260402", "已上线系统概况", "已上线系统概况", "待补官网口径。", "现有档案来源仍不足以支撑其系统与数字化事实描述，因此本字段保持待补，而不再用模板化推断填满。", "/Users/clairaipartner/.openclaw/workspace-main/bussiness-master/docs/03-执行与校验/静态潜客池-L3全量补强专项-v1.md", "high_confidence_public", "B", "background_profile"),
        ],
    },
}


PROFILE_FIELDS = [
    "公司产品与服务概述",
    "商业模式概述",
    "核心客户客群",
    "已上线系统概况",
    "数字化项目动态",
    "主要竞品概述",
    "招聘代表岗位",
    "近一年重大事件",
    "产品与服务长摘录",
    "商业模式长摘录",
    "客户客群长摘录",
    "系统与数字化长摘录",
    "重大事件长摘录",
    "primary_source_types",
    "primary_source_refs",
    "official_source_count",
    "high_confidence_source_count",
    "validation_gap",
]


def build_index(ws):
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {h: i + 1 for i, h in enumerate(headers)}


def find_row(ws, idx, account_id):
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, idx["account_id"]).value or "") == account_id:
            return r
    raise KeyError(account_id)


def upsert_observation(ws, idx, row_data):
    obs_id = row_data[0]
    target_row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, idx["observation_id"]).value or "") == obs_id:
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1
    for col, value in zip(range(1, len(row_data) + 1), row_data):
        ws.cell(target_row, col).value = value


def main():
    main_wb = load_workbook(MAIN_XLSX)
    main_ws = main_wb["accounts_main"]
    mi = build_index(main_ws)

    profile_wb = load_workbook(PROFILE_XLSX)
    profiles_ws = profile_wb["account_profiles"]
    observations_ws = profile_wb["field_observations"]
    notes_ws = profile_wb["profile_notes"]
    pi = build_index(profiles_ws)
    oi = build_index(observations_ws)
    ni = build_index(notes_ws)

    for account_id, payload in SAMPLES.items():
        main_row = find_row(main_ws, mi, account_id)
        profile_row = find_row(profiles_ws, pi, account_id)

        for field in PROFILE_FIELDS:
            value = payload[field]
            if field in mi:
                main_ws.cell(main_row, mi[field]).value = value
            profiles_ws.cell(profile_row, pi[field]).value = value
        main_ws.cell(main_row, mi["last_verified_at"]).value = "2026-04-02"
        profiles_ws.cell(profile_row, pi["last_profiled_at"]).value = "2026-04-02"
        profiles_ws.cell(profile_row, pi["profile_owner"]).value = "Codex 结构化沉淀"

        for r in range(2, observations_ws.max_row + 1):
            if str(observations_ws.cell(r, oi["account_id"]).value or "") != account_id:
                continue
            field_name = str(observations_ws.cell(r, oi["field_name"]).value or "")
            if field_name in {
                "公司产品与服务概述",
                "商业模式概述",
                "核心客户客群",
                "已上线系统概况",
            }:
                observations_ws.cell(r, oi["is_current_best"]).value = "no"

        for obs_id, field_name, field_label, short_value, excerpt, source_locator, source_type, source_strength, supports_dimension in payload["notes"]:
            row_data = [
                obs_id,
                account_id,
                field_name,
                field_label,
                short_value,
                excerpt,
                source_locator,
                source_type,
                source_strength,
                supports_dimension,
                "yes",
                "2026-04-02",
                "codex_llm",
                "客户事实优先样板：改用来源型摘录，不再用模板解释句。",
            ]
            upsert_observation(observations_ws, oi, row_data)

        note_id = f"note_{account_id}_fact_first_20260402"
        note_row = None
        for r in range(2, notes_ws.max_row + 1):
            if str(notes_ws.cell(r, ni["note_id"]).value or "") == note_id:
                note_row = r
                break
        if note_row is None:
            note_row = notes_ws.max_row + 1
        notes_ws.cell(note_row, ni["note_id"]).value = note_id
        notes_ws.cell(note_row, ni["account_id"]).value = account_id
        notes_ws.cell(note_row, ni["note_type"]).value = "context_note"
        notes_ws.cell(note_row, ni["note_title"]).value = "客户事实优先样板"
        notes_ws.cell(note_row, ni["note_body"]).value = "本档案页已按“客户事实优先”重写：先描述公司事实，再展示静态池判断；缺口保持待补，不再用模板化解释补满。"
        notes_ws.cell(note_row, ni["importance"]).value = "high"
        notes_ws.cell(note_row, ni["created_at"]).value = "2026-04-02"
        notes_ws.cell(note_row, ni["created_by"]).value = "codex_llm"

    main_wb.save(MAIN_XLSX)
    profile_wb.save(PROFILE_XLSX)

    with MEMORY_PATH.open("a") as f:
        f.write("\n- 已对 6 个客户档案样板对象执行“客户事实优先”重写，回写主表、档案库、field_observations 和 profile_notes。\n")


if __name__ == "__main__":
    main()
