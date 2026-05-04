from __future__ import annotations

import json
import os
import re
from pathlib import Path

import requests
from openpyxl import load_workbook


VAULT = Path("/Users/clairaipartner/Documents/Obsidian-Codex/潜客池")
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
TODAY = "2026-04-02"
FACT_PLACEHOLDER = "待补官网/年报/IR口径"
DEFAULT_BATCH_LIMIT = int(os.environ.get("ENRICH_BATCH_LIMIT", "30"))
TARGET_LEVELS = {x.strip() for x in os.environ.get("ENRICH_LEVELS", "L1").split(",") if x.strip()}
ENRICH_TARGET = os.environ.get("ENRICH_TARGET", "all").strip().lower()
GENERIC_CUSTOMER_VALUES = {
    FACT_PLACEHOLDER,
    "工业企业、行业客户及项目型客户。",
    "终端消费者、经销/零售渠道及线上线下平台。",
    "医院、医疗机构、科研机构及专业客户。",
    "工程施工、基础设施建设及工业客户。",
}
GENERIC_PRODUCT_VALUES = {
    FACT_PLACEHOLDER,
}
GENERIC_BUSINESS_VALUES = {
    FACT_PLACEHOLDER,
    "以品牌经营和渠道销售为核心，通过线上线下渠道覆盖目标市场。",
    "以多平台跨境销售为核心，覆盖线上平台与海外渠道经营。",
    "品牌零售或连锁经营，依赖总部、区域与门店协同管理。",
    "品牌消费品全渠道经营，依赖商品、供应链和渠道协同。",
    "技术驱动型制造，覆盖研发、制造、销售与项目交付协同。",
}

USER_AGENT = {"User-Agent": "Mozilla/5.0"}
EASTMONEY_SEARCH_TOKEN = os.environ.get("EASTMONEY_SEARCH_TOKEN", "").strip()
SESSION = requests.Session()
SESSION.headers.update(USER_AGENT)


def clean_text(value: str) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", "", text)
    return text.replace("\u3000", "")


def legal_short_name(name: str) -> str:
    text = clean_text(name)
    for suffix in [
        "股份有限公司",
        "集团股份公司",
        "集团股份有限公司",
        "股份公司",
        "有限责任公司",
        "有限公司",
    ]:
        if text.endswith(suffix):
            return text[: -len(suffix)]
    return text


def quote_id_from_refs(refs: str) -> str | None:
    text = str(refs or "")
    m = re.search(r"quote\.eastmoney\.com/(?:concept/)?(sh|sz)(\d{6})\.html", text, re.I)
    if m:
        market = "1" if m.group(1).lower() == "sh" else "0"
        return f"{market}.{m.group(2)}"
    m = re.search(r"quote\.eastmoney\.com/unify/r/([01])\.(\d{6})", text, re.I)
    if m:
        return f"{m.group(1)}.{m.group(2)}"
    return None


def search_candidates(name: str, aliases: list[str] | None = None) -> list[str]:
    base = legal_short_name(name)
    candidates = []
    for value in [*(aliases or []), base, clean_text(name)]:
        value = clean_text(value)
        if value and value not in candidates:
            candidates.append(value)
    shrink_tokens = ["科技集团", "集团", "科技", "电气驱动", "股份", "有限", "公司"]
    for value in list(candidates):
        short = value
        for token in shrink_tokens:
            short = short.replace(token, "")
        short = short.strip()
        if len(short) >= 3 and short not in candidates:
            candidates.append(short)
    return candidates


def eastmoney_search(name: str, aliases: list[str] | None = None) -> dict | None:
    if not EASTMONEY_SEARCH_TOKEN:
        return None
    url = "https://searchapi.eastmoney.com/api/suggest/get"
    for candidate in search_candidates(name, aliases):
        params = {
            "input": candidate,
            "type": "14",
            "token": EASTMONEY_SEARCH_TOKEN,
            "count": "10",
        }
        data = SESSION.get(url, params=params, timeout=20).json()
        rows = (data.get("QuotationCodeTable") or {}).get("Data") or []
        if not rows:
            continue
        for row in rows:
            if row.get("Classify") == "AStock":
                return row
        return rows[0]
    return None


def company_survey(quote_id: str) -> dict | None:
    market, code = quote_id.split(".")
    prefix = "SH" if market == "1" else "SZ" if market == "0" else None
    if not prefix:
        return None
    url = f"https://emweb.securities.eastmoney.com/PC_HSF10/CompanySurvey/CompanySurveyAjax?code={prefix}{code}"
    data = SESSION.get(url, timeout=20).json()
    return data.get("jbzl")


def first_sentences(text: str, limit: int = 2) -> list[str]:
    raw = clean_text(text)
    parts = re.split(r"[。；!?！？]", raw)
    parts = [p for p in parts if p]
    return parts[:limit]


def summarize_product(profile: dict) -> str | None:
    intro = clean_text(profile.get("gsjj") or "")
    business_scope = clean_text(profile.get("jyfw") or "")
    short_name = profile.get("agjc") or profile.get("gsmc") or ""
    if intro:
        sents = first_sentences(intro, limit=5)
        for sent in sents:
            if any(skip in sent for skip in ["上市", "市值", "福布斯", "FT全球", "股票代码"]):
                continue
            if any(token in sent for token in ["专注于", "专注为", "主营", "产品矩阵", "设计", "研发", "品牌", "解决方案", "设备", "机械", "硬件", "软件", "饮品", "食品", "护理用品", "调味品", "医疗设备", "工程机械", "消费电子", "连接器", "家居"]):
                return f"{short_name}：{sent}。"
        if sents:
            sent = sents[0]
            if any(token in sent for token in ["产品", "品牌", "设备", "食品", "饮品", "用品", "家居", "器械", "机械"]):
                return f"{short_name}：{sent}。"
    if business_scope:
        parts = [p for p in re.split(r"[;,；、]", business_scope) if p][:4]
        if parts:
            return f"{short_name}经营范围涵盖" + "、".join(parts) + "。"
    return None


def summarize_business(profile: dict) -> str | None:
    intro = clean_text(profile.get("gsjj") or "")
    scope = clean_text(profile.get("jyfw") or "")
    short_name = profile.get("agjc") or profile.get("gsmc") or ""
    text = intro + scope
    if any(token in text for token in ["调味品", "酱油", "蚝油", "醋", "复合调味料", "食品", "乳制品", "饮料", "休闲食品", "卤味"]):
        return f"以{short_name}相关食品或调味品的生产和销售为核心，通过经销、零售或餐饮渠道覆盖市场。"
    if any(token in text for token in ["化妆品", "护肤", "彩妆", "个护", "护理用品"]):
        return f"以{short_name}相关美妆个护产品经营为核心，覆盖电商、零售终端或经销渠道。"
    if any(token in text for token in ["家居", "家纺", "家具", "床垫", "睡眠", "卫浴", "遮阳", "厨柜"]):
        return f"以{short_name}相关家居产品经营为核心，覆盖零售门店、电商平台与经销渠道。"
    if any(token in text for token in ["母婴", "婴童", "宠物", "宠粮", "宠物食品"]):
        return f"以{short_name}相关母婴、宠物或耐用品经营为核心，覆盖电商、零售终端与经销渠道。"
    if any(token in text for token in ["百货", "商场", "卖场", "超市", "门店", "连锁"]):
        return f"以{short_name}相关零售网络经营为核心，依赖总部、区域与门店协同。"
    if any(token in text for token in ["跨境", "Amazon", "亚马逊", "海外", "平台", "独立站", "站点"]):
        return f"以{short_name}相关跨境产品经营为核心，覆盖海外平台、站点与履约渠道。"
    if any(token in text for token in ["医院", "医学影像", "医疗设备", "器械"]):
        return f"以{short_name}相关医疗设备或解决方案经营为核心，面向医疗机构和专业客户。"
    if any(token in text for token in ["工程机械", "起重", "挖掘", "施工机械", "基础设施"]):
        return f"以{short_name}相关工程机械产品的研发、生产和销售为核心，覆盖直销、租赁或经销体系。"
    if any(token in text for token in ["自动化", "机器人", "半导体", "连接器", "驱动", "控制", "新能源", "储能", "材料", "装备"]):
        return f"以{short_name}相关工业产品或设备的研发、生产和销售为核心。"
    if intro:
        if all(token in intro for token in ["研发", "生产", "销售"]):
            return f"以{short_name}相关产品或解决方案的研发、生产和销售为核心。"
        if "设计" in intro and "研发" in intro and "销售" in intro:
            return f"以{short_name}相关产品的设计、研发和销售为核心。"
        if "解决方案" in intro and "客户" in intro:
            return f"以{short_name}相关产品与解决方案经营为核心，面向行业客户提供产品与服务。"
        if "线上平台" in intro or "Amazon" in intro or "跨境" in intro:
            return f"以多平台跨境销售为核心，覆盖线上平台与海外渠道经营。"
        if "经销" in intro and "零售" in intro:
            return f"以品牌经营为核心，通过经销、零售终端及线上线下渠道覆盖市场。"
        if any(token in intro for token in ["品牌", "消费者"]) and any(token in intro for token in ["经销", "渠道", "零售", "销售"]):
            return f"以品牌经营和渠道销售为核心，通过线上线下渠道覆盖目标市场。"
    if all(token in scope for token in ["生产", "销售", "维修"]):
        return f"以{short_name}相关产品的生产、销售与服务为核心。"
    if "生产" in scope and "销售" in scope:
        return f"以{short_name}相关产品的生产和销售为核心。"
    if "研发" in scope and "销售" in scope:
        return f"以{short_name}相关产品的研发和销售为核心。"
    return None


def summarize_customer(profile: dict) -> str | None:
    intro = clean_text(profile.get("gsjj") or "")
    scope = clean_text(profile.get("jyfw") or "")
    text = intro + scope
    if any(token in text for token in ["调味品", "酱油", "蚝油", "食品", "休闲食品", "饮料", "功能饮料", "乳业", "卤味", "火腿", "肉制品"]):
        return "家庭消费者、餐饮客户、经销商和零售渠道。"
    if any(token in text for token in ["纸尿裤", "卫生巾", "护理用品", "无纺", "个护", "护理"]):
        return "个人及家庭护理消费者、商超零售、电商平台与经销渠道。"
    if any(token in text for token in ["母婴", "婴童", "儿童用品"]):
        return "母婴家庭、儿童消费家庭、门店会员及零售渠道。"
    if any(token in text for token in ["宠物", "宠粮", "宠物食品"]):
        return "养宠家庭、电商平台、宠物零售渠道与经销网络。"
    if any(token in text for token in ["化妆品", "美妆", "护肤", "彩妆"]):
        return "终端消费者、电商平台、零售终端与经销渠道。"
    if any(token in text for token in ["家居", "家纺", "家具", "睡眠", "床垫", "厨柜", "卫浴", "遮阳"]):
        return "家庭消费者、家居零售门店、电商平台与经销渠道。"
    if any(token in text for token in ["零售", "百货", "商场", "门店", "连锁"]):
        return "终端消费者、门店会员、零售终端及区域门店网络。"
    if any(token in text for token in ["医院", "医疗机构", "科研机构", "医学中心", "医学影像", "医疗设备"]):
        return "医院、医学中心、科研院所及医疗设备采购相关机构。"
    if any(token in text for token in ["工程机械", "挖掘机械", "起重机械", "筑路机械", "桩工机械", "施工", "基础设施"]):
        return "工程施工客户、设备租赁及经销体系、基础设施建设相关客户。"
    if any(token in text for token in ["跨境", "Amazon", "亚马逊", "海外", "品牌出海", "站点", "全球市场", "平台"]):
        return "海外终端消费者、电商平台用户、海外零售及分销渠道。"
    if any(token in text for token in ["工业自动化", "连接器", "半导体", "显示", "芯片", "存储", "电子元器件", "电气", "材料", "新能源", "储能", "机器人", "自动化设备", "专用设备", "高端装备"]):
        return "工业制造客户、设备厂商、渠道伙伴及项目型客户。"
    return None


def target_fields() -> list[str]:
    mapping = {
        "all": ["公司产品与服务概述", "商业模式概述", "核心客户客群"],
        "product": ["公司产品与服务概述"],
        "business": ["商业模式概述"],
        "customer": ["核心客户客群"],
        "event": ["近一年重大事件"],
    }
    return mapping.get(ENRICH_TARGET, mapping["all"])


def is_generic_product_text(current: str) -> bool:
    current = str(current or "").strip()
    if not current:
        return True
    if current in GENERIC_PRODUCT_VALUES:
        return True
    generic_tokens = [
        "经营范围涵盖一般经营项目是",
        "主营",
        "属于多平台经营的跨境品牌或供应链出海主体",
        "属于研产销协同复杂的技术型制造主体",
        "属于高 SKU",
        "属于多工厂",
        "品牌消费品主体",
        "制造主体",
    ]
    return any(token in current for token in generic_tokens)


def is_generic_business_text(current: str) -> bool:
    current = str(current or "").strip()
    if not current:
        return True
    if current in GENERIC_BUSINESS_VALUES:
        return True
    generic_tokens = [
        "品牌零售或连锁经营，依赖总部、区域与门店协同管理。",
        "品牌消费品全渠道经营，依赖商品、供应链和渠道协同。",
        "技术驱动型制造，覆盖研发、制造、销售与项目交付协同。",
        "以多平台跨境销售为核心，覆盖线上平台与海外渠道经营。",
        "以品牌经营和渠道销售为核心，通过线上线下渠道覆盖目标市场。",
        "现有材料可支撑",
        "经营复杂度",
    ]
    return any(token in current for token in generic_tokens)


def should_refresh_existing(field: str, current: str, new_value: str | None) -> bool:
    current = str(current or "").strip()
    if not new_value:
        return False
    if current == FACT_PLACEHOLDER:
        return True
    if field == "核心客户客群" and current in GENERIC_CUSTOMER_VALUES and current != new_value:
        return True
    if field == "公司产品与服务概述" and "经营范围涵盖一般经营项目是" in current:
        return True
    if field == "公司产品与服务概述" and is_generic_product_text(current) and current != new_value:
        return True
    if field == "商业模式概述" and is_generic_business_text(current) and current != new_value:
        return True
    return False


def field_needs_refresh(field: str, current: str) -> bool:
    current = str(current or "").strip()
    if field == "公司产品与服务概述":
        return not current or current == FACT_PLACEHOLDER or is_generic_product_text(current)
    if field == "商业模式概述":
        return not current or current == FACT_PLACEHOLDER or is_generic_business_text(current)
    if field == "核心客户客群":
        return not current or current == FACT_PLACEHOLDER or current in GENERIC_CUSTOMER_VALUES
    if field == "近一年重大事件":
        return not current or current == FACT_PLACEHOLDER
    return False


def collect_and_update_profile(ws, target_levels: set[str], batch_limit: int = DEFAULT_BATCH_LIMIT) -> dict[str, dict]:
    headers = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}
    results: dict[str, dict] = {}
    updated = 0
    candidates = []
    fields = target_fields()
    for row in range(2, ws.max_row + 1):
        level = str(ws.cell(row, idx["静态潜客记录成熟度"]).value or "")
        if level not in target_levels:
            continue
        missing_fields = [
            field for field in fields
            if field in idx and field_needs_refresh(field, str(ws.cell(row, idx[field]).value or "").strip())
        ]
        current_product = str(ws.cell(row, idx["公司产品与服务概述"]).value or "").strip()
        current_business = str(ws.cell(row, idx["商业模式概述"]).value or "").strip()
        current_customer = str(ws.cell(row, idx["核心客户客群"]).value or "").strip()
        priority_count = (
            len(missing_fields)
            + (1 if current_customer in GENERIC_CUSTOMER_VALUES else 0)
            + (1 if "经营范围涵盖一般经营项目是" in current_product else 0)
            + (1 if is_generic_product_text(current_product) else 0)
            + (1 if is_generic_business_text(current_business) else 0)
        )
        if not missing_fields:
            continue
        if ENRICH_TARGET == "all" and priority_count < 2:
            continue
        refs = str(ws.cell(row, idx["primary_source_refs"]).value or "")
        score = (
            0 if "quote.eastmoney.com" in refs else
            1 if "cninfo.com.cn" in refs else
            2
        )
        level_score = 0 if level == "L1" else 1 if level == "L2" else 2
        candidates.append((level_score, score, -priority_count, row))

    for _, __, ___, row in sorted(candidates)[:batch_limit]:
        acc = str(ws.cell(row, idx["account_id"]).value or "").strip()
        name = str(ws.cell(row, idx["account_canonical_name"]).value or "").strip()
        brand_name = str(ws.cell(row, idx.get("brand_name", 0)).value or "").strip() if idx.get("brand_name") else ""
        if not name:
            continue
        try:
            refs_text = str(ws.cell(row, idx["primary_source_refs"]).value or "")
            quote_id = quote_id_from_refs(refs_text)
            found = None
            if quote_id:
                market, code = quote_id.split(".")
                found = {"QuoteID": quote_id, "Code": code}
            else:
                found = eastmoney_search(name, aliases=[brand_name] if brand_name else None)
            if not found or "." not in str(found.get("QuoteID") or ""):
                continue
            survey = company_survey(found["QuoteID"])
            if not survey:
                continue
        except Exception:
            continue

        product = summarize_product(survey) if "公司产品与服务概述" in fields else None
        business = summarize_business(survey) if "商业模式概述" in fields else None
        customer = summarize_customer(survey) if "核心客户客群" in fields else None
        changed = False
        current_product = str(ws.cell(row, idx["公司产品与服务概述"]).value or "").strip()
        if "公司产品与服务概述" in fields and should_refresh_existing("公司产品与服务概述", current_product, product):
            ws.cell(row, idx["公司产品与服务概述"]).value = product
            if idx.get("产品与服务长摘录"):
                ws.cell(row, idx["产品与服务长摘录"]).value = clean_text(survey.get("gsjj") or "")[:220] + "。"
            changed = True
        current_business = str(ws.cell(row, idx["商业模式概述"]).value or "").strip()
        if "商业模式概述" in fields and should_refresh_existing("商业模式概述", current_business, business):
            ws.cell(row, idx["商业模式概述"]).value = business
            if idx.get("商业模式长摘录"):
                ws.cell(row, idx["商业模式长摘录"]).value = clean_text(survey.get("gsjj") or "")[:220] + "。"
            changed = True
        current_customer = str(ws.cell(row, idx["核心客户客群"]).value or "").strip()
        if "核心客户客群" in fields and should_refresh_existing("核心客户客群", current_customer, customer):
            ws.cell(row, idx["核心客户客群"]).value = customer
            if idx.get("客户客群长摘录"):
                ws.cell(row, idx["客户客群长摘录"]).value = clean_text(survey.get("gsjj") or "")[:220] + "。"
            changed = True

        if idx.get("primary_source_refs"):
            existing = str(ws.cell(row, idx["primary_source_refs"]).value or "")
            quote_ref = f"https://quote.eastmoney.com/unify/r/{found['QuoteID']}"
            survey_ref = f"https://emweb.securities.eastmoney.com/PC_HSF10/CompanySurvey/Index?type=web&code={'SH' if found['QuoteID'].startswith('1.') else 'SZ'}{found['Code']}"
            refs = [x for x in existing.split("\n") if x.strip()]
            for ref in [quote_ref, survey_ref]:
                if ref not in refs:
                    refs.append(ref)
                    changed = True
            ws.cell(row, idx["primary_source_refs"]).value = "\n".join(refs)
        if idx.get("primary_source_types"):
            source_types = str(ws.cell(row, idx["primary_source_types"]).value or "")
            tokens = [x.strip() for x in source_types.split(",") if x.strip()]
            for token in ["上市公司公开披露", "东方财富F10公司概况"]:
                if token not in tokens:
                    tokens.append(token)
                    changed = True
            ws.cell(row, idx["primary_source_types"]).value = ",".join(tokens)
        if idx.get("official_source_count"):
            current = int(str(ws.cell(row, idx["official_source_count"]).value or 0) or 0)
            if current < 1:
                ws.cell(row, idx["official_source_count"]).value = 1
                changed = True
        if idx.get("last_profiled_at"):
            ws.cell(row, idx["last_profiled_at"]).value = TODAY
        if changed:
            updated += 1
            results[acc] = {
                "name": name,
                "product": product,
                "business": business,
                "customer": customer,
                "quote_id": found["QuoteID"],
                "code": found["Code"],
            }
    return {"updated": updated, "rows": results}


def apply_to_main(ws, fetched: dict[str, dict], target_levels: set[str]) -> int:
    headers = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}
    updated = 0
    fields = target_fields()
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, idx["静态潜客记录成熟度"]).value or "") not in target_levels:
            continue
        acc = str(ws.cell(row, idx["account_id"]).value or "").strip()
        payload = fetched.get(acc)
        if not payload:
            continue
        changed = False
        current_product = str(ws.cell(row, idx["公司产品与服务概述"]).value or "").strip()
        current_business = str(ws.cell(row, idx["商业模式概述"]).value or "").strip()
        current_customer = str(ws.cell(row, idx["核心客户客群"]).value or "").strip()
        if "公司产品与服务概述" in fields and payload.get("product") and should_refresh_existing("公司产品与服务概述", current_product, payload.get("product")):
            ws.cell(row, idx["公司产品与服务概述"]).value = payload["product"]
            changed = True
        if "商业模式概述" in fields and payload.get("business") and should_refresh_existing("商业模式概述", current_business, payload.get("business")):
            ws.cell(row, idx["商业模式概述"]).value = payload["business"]
            changed = True
        if "核心客户客群" in fields and payload.get("customer") and should_refresh_existing("核心客户客群", current_customer, payload.get("customer")):
            ws.cell(row, idx["核心客户客群"]).value = payload["customer"]
            changed = True
        if changed and idx.get("last_verified_at"):
            ws.cell(row, idx["last_verified_at"]).value = TODAY
            updated += 1
    return updated


def main() -> None:
    profile_wb = load_workbook(PROFILE_XLSX)
    result = collect_and_update_profile(profile_wb["account_profiles"], target_levels=TARGET_LEVELS)
    profile_wb.save(PROFILE_XLSX)

    main_wb = load_workbook(MAIN_XLSX)
    result_main = apply_to_main(main_wb["accounts_main"], result["rows"], target_levels=TARGET_LEVELS)
    main_wb.save(MAIN_XLSX)

    print({
        "profile_updated": result["updated"],
        "main_updated": result_main,
        "target_levels": sorted(TARGET_LEVELS),
        "enrich_target": ENRICH_TARGET,
        "examples": [v["name"] for v in list(result["rows"].values())[:10]],
    })


if __name__ == "__main__":
    main()
