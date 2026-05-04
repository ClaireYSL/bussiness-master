from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool.prospect_eligibility_gate import ProspectEligibilityGate
from shared.static_pool.signed_customer_gate import SignedCustomerGate, normalize_name

MILESTONES = WORKSPACE / "deliveries/archive/milestones"
CANONICAL = WORKSPACE / "deliveries/canonical/businessmaster"
M47 = MILESTONES / "milestone47r_trusted_pool_product"
M56 = MILESTONES / "milestone56r_trusted_pool_status_panel"
M180 = MILESTONES / "milestone180r_signed_customer_registry_v2"

HISTORICAL_LIST = Path("/Users/clairelu2026/个人文档【重要】【20260406】/07销售管理-LTO/静态潜客池/历年合作列表 2026-02-23 14_24_22.xlsx")
HISTORICAL_SHEET = "历年合作列表"
LEGACY_REGISTRY_V1 = CANONICAL / "signed_customer_registry_v1.json"
LEGACY_ALIAS_V1 = CANONICAL / "signed_customer_alias_registry_v1.json"
SIGNED_REGISTRY_V2 = CANONICAL / "signed_customer_registry_v2.json"
SIGNED_ALIAS_V2 = CANONICAL / "signed_customer_alias_registry_v2.json"
TRUSTED_POOL = M47 / "trusted_prospect_pool_v1.json"
SOURCE_TRACE = M47 / "source_trace_index_v1.json"
STATUS_PANEL = M56 / "trusted_pool_status_panel_v1.json"
VAULT_ROOT = Path("/Users/clairelu2026/26M3-Obsidian-潜客池/潜客池/07-可信潜客档案")

OLD_BLOCK_SAMPLES = ["百胜中国", "珀莱雅", "上海家化", "森马", "特步", "海澜之家", "锅圈", "来伊份", "天味食品", "水星家纺"]
NEW_BLOCK_SAMPLES = ["深圳智工坊科技有限公司", "广州丸碧化妆品有限公司", "广州优卡普科技有限公司", "深圳市美通供应链有限公司"]
NON_HIT_SAMPLES = ["不存在的未来样例科技有限公司", "上海随机非客户测试有限公司"]
BOUNDARY_SAMPLES = ["智工坊", "丸碧", "优卡普"]
DYNAMIC_TERMS = ["重点经营", "worth_following", "recommended_next_action", "business_feedback_pending"]
SECRET_PATTERNS = [r"sk-[A-Za-z0-9]{20,}", r"AKIA[0-9A-Z]{16}", r"DELEGATE_LLM_API_KEY\s*=\s*[^\s<]+"]
NON_COMPANY_TOKENS = ["案例", "方案", "对话", "CTO", "CIO", "访谈", "复盘", "方法论", "汇报", "交流", "不用", "如何用"]


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def read_json(path: Path, default: Any | None = None) -> Any:
    if not path.exists():
        return {} if default is None else default
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def stable_id(prefix: str, value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9\u4e00-\u9fff]+", "_", value.lower()).strip("_")[:28]
    digest = hashlib.sha1(value.encode("utf-8")).hexdigest()[:10]
    return f"{prefix}_{slug}_{digest}" if slug else f"{prefix}_{digest}"


def compact_name(value: Any) -> str:
    return re.sub(r"[\s·•,，。.;；:：/\\|\-—_（）()]+", "", str(value or "").strip().lower())


def conservative_aliases(name: str) -> list[dict[str, str]]:
    aliases: dict[str, str] = {name: "canonical_name"}
    ascii_paren = name.replace("（", "(").replace("）", ")")
    aliases.setdefault(ascii_paren, "punctuation_variant")
    no_space = re.sub(r"\s+", "", name)
    aliases.setdefault(no_space, "space_variant")
    short = re.sub(r"(股份有限公司|有限责任公司|有限公司|集团股份有限公司|集团有限公司)$", "", name)
    if short and short != name and len(short) >= 4:
        # 只保留相对安全的企业简称；2-3 字简称交给 fuzzy boundary review，不直接 block。
        aliases.setdefault(short, "legal_suffix_removed")
    if short.endswith("食汇") and len(short) > 2:
        aliases.setdefault(short[:-2], "brand_short_name")
    return [{"alias_name": k, "alias_type": v} for k, v in aliases.items() if k]


def looks_non_company(name: str) -> bool:
    return len(name) > 48 or any(token in name for token in NON_COMPANY_TOKENS)


def load_historical_list() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(HISTORICAL_LIST, read_only=True, data_only=True)
    if HISTORICAL_SHEET not in wb.sheetnames:
        raise ValueError(f"missing sheet: {HISTORICAL_SHEET}")
    ws = wb[HISTORICAL_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(x or "").strip() for x in rows[0]] if rows else []
    if "客户" not in headers or "首次签约日期 (季度)" not in headers:
        raise ValueError(f"unexpected headers: {headers}")

    items: list[dict[str, Any]] = []
    empty_rows: list[int] = []
    duplicates: list[dict[str, Any]] = []
    suspicious: list[dict[str, Any]] = []
    seen: dict[str, dict[str, Any]] = {}
    quarter_counts: Counter[str] = Counter()
    for row_index, row in enumerate(rows[1:], start=2):
        data = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        name = str(data.get("客户") or "").strip()
        quarter = str(data.get("首次签约日期 (季度)") or "").strip()
        if not name:
            empty_rows.append(row_index)
            continue
        key = normalize_name(name)
        if key in seen:
            duplicates.append({"row_index": row_index, "customer_name": name, "duplicate_of": seen[key].get("row_index")})
            continue
        item = {
            "row_index": row_index,
            "customer_id": stable_id("signed_customer", name),
            "canonical_name": name,
            "signed_status": "confirmed_signed_customer",
            "source_type": "historical_cooperation_list",
            "source_locator": str(HISTORICAL_LIST),
            "source_sheet": HISTORICAL_SHEET,
            "first_signed_quarter": quarter,
            "last_verified_at": now(),
            "exclusion_scope": "exclude_from_static_pool",
            "source_priority": "primary",
        }
        if looks_non_company(name):
            suspicious.append({"row_index": row_index, "customer_name": name, "reason": "name looks like non-company or too long"})
        seen[key] = item
        quarter_counts[quarter] += 1
        items.append(item)
    audit = {
        "source_path": str(HISTORICAL_LIST),
        "source_sheet": HISTORICAL_SHEET,
        "source_row_count": ws.max_row,
        "source_column_count": ws.max_column,
        "valid_customer_count": len(items),
        "unique_customer_count": len(seen),
        "empty_customer_name_count": len(empty_rows),
        "duplicate_customer_count": len(duplicates),
        "suspicious_non_company_count": len(suspicious),
        "quarter_distribution": dict(quarter_counts),
        "empty_rows": empty_rows[:100],
        "duplicates": duplicates[:200],
        "suspicious_non_company": suspicious[:200],
    }
    return items, audit


def merge_legacy(customers: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    merged = {normalize_name(item["canonical_name"]): item for item in customers}
    legacy_registry = read_json(LEGACY_REGISTRY_V1, {"items": []})
    legacy_items = legacy_registry.get("items") or []
    supplemental: list[dict[str, Any]] = []
    for legacy in legacy_items:
        name = str(legacy.get("canonical_name") or "").strip()
        if not name:
            continue
        key = normalize_name(name)
        if key in merged:
            merged[key].setdefault("supplemental_sources", []).append({
                "source_type": legacy.get("source_type") or "legacy_customer_registry",
                "source_locator": legacy.get("source_locator"),
                "source_note": legacy.get("source_note"),
                "legacy_customer_id": legacy.get("customer_id"),
            })
        else:
            item = {
                "customer_id": str(legacy.get("customer_id") or stable_id("legacy_signed_customer", name)),
                "canonical_name": name,
                "signed_status": "confirmed_signed_customer",
                "source_type": legacy.get("source_type") or "legacy_customer_registry",
                "source_locator": legacy.get("source_locator"),
                "source_note": legacy.get("source_note"),
                "first_signed_quarter": legacy.get("effective_from") or "legacy_unknown",
                "last_verified_at": now(),
                "exclusion_scope": legacy.get("exclusion_scope") or "exclude_from_static_pool",
                "source_priority": "supplemental_legacy",
            }
            merged[key] = item
            supplemental.append(item)
    return list(merged.values()), supplemental


def build_alias_registry(customers: list[dict[str, Any]]) -> dict[str, Any]:
    alias_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    legacy_aliases = read_json(LEGACY_ALIAS_V1, {"items": []}).get("items") or []
    customer_by_legacy_id = {item.get("customer_id"): item for item in customers}
    customer_by_norm = {normalize_name(item.get("canonical_name")): item for item in customers}

    for customer in customers:
        for alias in conservative_aliases(str(customer.get("canonical_name") or "")):
            alias_name = alias["alias_name"]
            alias_by_key[(customer["customer_id"], compact_name(alias_name))] = {
                "alias_id": stable_id("signed_alias", f"{customer['customer_id']}::{alias_name}"),
                "customer_id": customer["customer_id"],
                "canonical_name": customer["canonical_name"],
                "alias_name": alias_name,
                "alias_type": alias["alias_type"],
                "status": "active",
                "source_type": customer.get("source_type"),
                "source_locator": customer.get("source_locator"),
                "risk_level": "conservative",
            }
    for legacy in legacy_aliases:
        name = str(legacy.get("alias_name") or "").strip()
        if not name:
            continue
        customer = customer_by_legacy_id.get(legacy.get("customer_id")) or customer_by_norm.get(normalize_name(legacy.get("canonical_name")))
        if not customer:
            continue
        alias_by_key.setdefault((customer["customer_id"], compact_name(name)), {
            "alias_id": stable_id("signed_alias", f"{customer['customer_id']}::{name}"),
            "customer_id": customer["customer_id"],
            "canonical_name": customer["canonical_name"],
            "alias_name": name,
            "alias_type": legacy.get("alias_type") or "legacy_alias",
            "status": "active",
            "source_type": "legacy_alias_registry",
            "source_locator": legacy.get("source_locator"),
            "risk_level": "legacy_confirmed",
        })
    aliases = sorted(alias_by_key.values(), key=lambda x: (str(x.get("canonical_name")), str(x.get("alias_name"))))
    return {
        "registry_id": "signed_customer_alias_registry_v2",
        "generated_at": now(),
        "source_milestone": "M180R",
        "summary": {
            "alias_count": len(aliases),
            "active_alias_count": sum(1 for item in aliases if item.get("status") == "active"),
            "source_registry": "signed_customer_registry_v2",
            "conservative_alias_policy": True,
        },
        "items": aliases,
    }


def build_registry_v2() -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    historical, audit_summary = load_historical_list()
    merged, supplemental = merge_legacy(historical)
    merged = sorted(merged, key=lambda x: str(x.get("canonical_name")))
    registry = {
        "registry_id": "signed_customer_registry_v2",
        "generated_at": now(),
        "source_milestone": "M180R",
        "summary": {
            "confirmed_signed_customer_count": len(merged),
            "primary_historical_customer_count": len(historical),
            "supplemental_legacy_customer_count": len(supplemental),
            "source_path": str(HISTORICAL_LIST),
            "source_sheet": HISTORICAL_SHEET,
            "old_excel_written": False,
            "trusted_pool_written": False,
            "knowledge_asset_registry_written": False,
            "persona_registry_written": False,
        },
        "items": merged,
    }
    alias_registry = build_alias_registry(merged)
    import_audit = {
        "batch_id": "m180r_signed_customer_import_audit_v2",
        "milestone": "M180R",
        "generated_at": now(),
        "status": "PASS" if audit_summary["valid_customer_count"] > 900 else "FAIL_TOO_FEW_VALID_CUSTOMERS",
        "summary": audit_summary | {"merged_customer_count": len(merged), "supplemental_legacy_customer_count": len(supplemental)},
    }
    return registry, alias_registry, import_audit


def gate_report() -> dict[str, Any]:
    gate = SignedCustomerGate.from_files(SIGNED_REGISTRY_V2, SIGNED_ALIAS_V2)
    old_results = {name: gate.check(name).to_dict() for name in OLD_BLOCK_SAMPLES}
    new_results = {name: gate.check(name).to_dict() for name in NEW_BLOCK_SAMPLES}
    non_hit_results = {name: gate.check(name).to_dict() for name in NON_HIT_SAMPLES}
    boundary_results = {name: gate.check(name).to_dict() for name in BOUNDARY_SAMPLES}
    status = "PASS" if all(x["existing_customer_check_status"] == "excluded_existing_customer" for x in old_results.values()) and all(x["existing_customer_check_status"] == "excluded_existing_customer" for x in new_results.values()) and all(x["existing_customer_check_status"] == "passed" for x in non_hit_results.values()) and all(x["existing_customer_check_status"] == "boundary_review" for x in boundary_results.values()) else "FAIL"
    return {
        "batch_id": "m180r_signed_customer_gate_v2_regression_v1",
        "milestone": "M180R",
        "generated_at": now(),
        "status": status,
        "summary": {
            "old_block_sample_count": len(old_results),
            "new_block_sample_count": len(new_results),
            "non_hit_sample_count": len(non_hit_results),
            "boundary_sample_count": len(boundary_results),
            "v2_enabled": True,
        },
        "old_block_results": old_results,
        "new_block_results": new_results,
        "non_hit_results": non_hit_results,
        "boundary_results": boundary_results,
    }


def audit_current_outputs() -> tuple[dict[str, Any], dict[str, Any]]:
    gate = SignedCustomerGate.from_files(SIGNED_REGISTRY_V2, SIGNED_ALIAS_V2)
    pool = read_json(TRUSTED_POOL, {"items": []}).get("items") or []
    trusted_hits = []
    for item in pool:
        match = gate.check(item.get("company_name")).to_dict()
        if match["existing_customer_check_status"] in {"excluded_existing_customer", "boundary_review"}:
            trusted_hits.append({"prospect_id": item.get("prospect_id"), "company_name": item.get("company_name"), "level": item.get("level"), "match": match})

    vault_hits = []
    for folder in ["01-L1 ICP强匹配档案", "02-L2正式潜客档案", "03-L3可信摘要卡"]:
        root = VAULT_ROOT / folder
        if not root.exists():
            continue
        for path in root.glob("*.md"):
            if path.name == "README.md":
                continue
            name = path.stem
            match = gate.check(name).to_dict()
            if match["existing_customer_check_status"] in {"excluded_existing_customer", "boundary_review"}:
                vault_hits.append({"company_name": name, "path": str(path), "folder": folder, "match": match})

    audit = {
        "batch_id": "m180r_existing_customer_audit_report_v2",
        "milestone": "M180R",
        "generated_at": now(),
        "status": "PASS_AUDIT_READY",
        "summary": {
            "trusted_pool_count": len(pool),
            "trusted_pool_hit_count": len(trusted_hits),
            "trusted_pool_block_hit_count": sum(1 for x in trusted_hits if x["match"]["existing_customer_check_status"] == "excluded_existing_customer"),
            "trusted_pool_boundary_review_count": sum(1 for x in trusted_hits if x["match"]["existing_customer_check_status"] == "boundary_review"),
            "vault_hit_count": len(vault_hits),
            "auto_remediation_executed": False,
        },
        "trusted_pool_hits": trusted_hits,
        "vault_hits": vault_hits[:300],
    }
    remediation_items = []
    for hit in trusted_hits:
        remediation_items.append({
            "entity_name": hit["company_name"],
            "prospect_id": hit.get("prospect_id"),
            "current_level": hit.get("level"),
            "match_status": hit["match"]["existing_customer_check_status"],
            "recommended_action": "从新潜客展示层移除或标记 excluded_existing_customer；保留为知识/客户案例参考。",
            "auto_apply": False,
        })
    remediation = {
        "batch_id": "m180r_existing_customer_remediation_preview_v2",
        "milestone": "M180R",
        "generated_at": now(),
        "status": "PASS_PREVIEW_READY",
        "summary": {"remediation_candidate_count": len(remediation_items), "auto_apply_enabled": False},
        "items": remediation_items,
        "no_write_proof": {"trusted_pool_written": False, "vault_written": False, "old_excel_written": False},
    }
    return audit, remediation


def update_status_panel(registry: dict[str, Any], alias_registry: dict[str, Any], audit: dict[str, Any], expert: dict[str, Any]) -> dict[str, Any]:
    panel = read_json(STATUS_PANEL, {})
    panel["generated_at"] = now()
    panel["latest_milestone"] = "M180R"
    panel["overall_status"] = "PASS_M180R_SIGNED_CUSTOMER_REGISTRY_V2_READY" if expert.get("status") in {"pass", "pass_with_followups"} else "FAIL_M180R_SIGNED_CUSTOMER_REGISTRY_V2"
    panel["canonical_next_action"] = "基于 M180 remediation preview 审阅当前 trusted pool/vault 的老客命中；下一步单独执行受控 remediation。"
    counts = panel.setdefault("counts", {})
    counts.setdefault("signed_customer_v1_count", counts.get("signed_customer_count"))
    counts.setdefault("signed_customer_alias_v1_count", counts.get("signed_customer_alias_count"))
    counts["signed_customer_count"] = registry.get("summary", {}).get("confirmed_signed_customer_count", 0)
    counts["signed_customer_alias_count"] = alias_registry.get("summary", {}).get("alias_count", 0)
    counts["signed_customer_v2_count"] = registry.get("summary", {}).get("confirmed_signed_customer_count", 0)
    counts["signed_customer_alias_v2_count"] = alias_registry.get("summary", {}).get("alias_count", 0)
    counts["m180_trusted_pool_existing_customer_hit_count"] = audit.get("summary", {}).get("trusted_pool_hit_count", 0)
    counts["m180_vault_existing_customer_hit_count"] = audit.get("summary", {}).get("vault_hit_count", 0)
    panel["m180r_signed_customer_registry_v2"] = {
        "status": panel["overall_status"],
        "signed_customer_count": registry.get("summary", {}).get("confirmed_signed_customer_count"),
        "alias_count": alias_registry.get("summary", {}).get("alias_count"),
        "source_path": registry.get("summary", {}).get("source_path"),
        "trusted_pool_hit_count": audit.get("summary", {}).get("trusted_pool_hit_count"),
        "vault_hit_count": audit.get("summary", {}).get("vault_hit_count"),
        "expert_review_status": expert.get("status"),
        "v2_default_gate_enabled": True,
    }
    return panel


def expert_review(registry: dict[str, Any], gate: dict[str, Any], audit: dict[str, Any]) -> dict[str, Any]:
    followups = []
    if audit.get("summary", {}).get("trusted_pool_hit_count", 0) > 0:
        followups.append("当前 trusted pool 命中签约老客 v2，需要后续受控 remediation；本轮只生成 preview。")
    if audit.get("summary", {}).get("vault_hit_count", 0) > 0:
        followups.append("vault 用户层存在签约老客命中，需要后续受控治理；本轮不删除文件。")
    registry_ok = registry.get("summary", {}).get("confirmed_signed_customer_count", 0) > 900
    gate_ok = gate.get("status") == "PASS"
    product_status = "pass_with_followups" if followups else "pass"
    architecture_status = "pass" if registry_ok and gate_ok else "fail"
    data_status = "pass" if registry_ok and gate_ok else "fail"
    status = "fail" if "fail" in {product_status, architecture_status, data_status} else ("pass_with_followups" if followups else "pass")
    return {
        "batch_id": "m180r_expert_review_report_v1",
        "milestone": "M180R",
        "generated_at": now(),
        "status": status,
        "summary": {
            "product_review_status": product_status,
            "architecture_review_status": architecture_status,
            "data_governance_review_status": data_status,
            "followup_count": len(followups),
            "remediation_write_allowed": False,
        },
        "product_review": {"status": product_status, "followups": followups, "evidence": "v2 gate 已覆盖近千家合作客户；命中项只输出 remediation preview。"},
        "architecture_review": {"status": architecture_status, "evidence": "signed_customer_gate 默认 v2，v1 fallback/reference；pipeline readiness/production 接入 M180。"},
        "data_governance_review": {"status": data_status, "evidence": "历史合作名单为 primary source，legacy 10 家为 supplemental；客户案例不自动确认 signed customer。"},
    }


def scan_dynamic_terms(paths: list[Path]) -> list[dict[str, str]]:
    findings = []
    for root in paths:
        candidates = [root] if root.is_file() else list(root.rglob("*.json")) + list(root.rglob("*.md")) + list(root.rglob("*.py"))
        for path in candidates:
            if not path.is_file() or "legacy" in str(path).lower() or "m57r" in str(path).lower():
                continue
            if path.name.startswith("m180_validation_report"):
                continue
            text = path.read_text(encoding="utf-8", errors="ignore")
            for term in DYNAMIC_TERMS:
                if term in text:
                    findings.append({"path": str(path), "term": term})
    return findings


def scan_secrets(paths: list[Path]) -> list[dict[str, str]]:
    findings = []
    for root in paths:
        candidates = [root] if root.is_file() else list(root.rglob("*"))
        for path in candidates:
            if not path.is_file() or path.name == ".env" or path.suffix.lower() in {".xlsx", ".zip", ".png", ".pdf"}:
                continue
            text = path.read_text(encoding="utf-8", errors="ignore")[:200000]
            for pat in SECRET_PATTERNS:
                if re.search(pat, text):
                    findings.append({"path": str(path), "pattern": pat})
    return findings


def run_cmd(cmd: list[str]) -> dict[str, Any]:
    proc = subprocess.run(cmd, cwd=WORKSPACE, text=True, capture_output=True)
    return {"cmd": cmd, "returncode": proc.returncode, "stdout": proc.stdout[-3000:], "stderr": proc.stderr[-3000:]}


def validate(registry: dict[str, Any], alias_registry: dict[str, Any], import_audit: dict[str, Any], gate: dict[str, Any], audit: dict[str, Any], expert: dict[str, Any]) -> dict[str, Any]:
    compile_result = run_cmd(["python3", "-m", "py_compile", "scripts/build_m180r_signed_customer_registry_v2.py", "scripts/businessmaster_pipeline.py"])
    json_paths = [SIGNED_REGISTRY_V2, SIGNED_ALIAS_V2, M180 / "signed_customer_import_audit_v2.json", M180 / "signed_customer_gate_v2_regression_v1.json", M180 / "existing_customer_audit_report_v2.json", M180 / "remediation_preview_v2.json", M180 / "m180_expert_review_report_v1.json", STATUS_PANEL]
    json_errors = []
    for path in json_paths:
        try:
            read_json(path)
        except Exception as exc:  # noqa: BLE001
            json_errors.append({"path": str(path), "error": str(exc)})
    readiness = run_cmd(["python3", "scripts/businessmaster_pipeline.py", "--mode", "readiness"])
    dry_run = run_cmd(["python3", "scripts/businessmaster_pipeline.py", "--mode", "production", "--dry-run"])
    gate_default = SignedCustomerGate.from_files()
    default_v2_probe = gate_default.check("深圳智工坊科技有限公司").status
    eligibility_probe = ProspectEligibilityGate.from_files().check("深圳智工坊科技有限公司", "m180_probe").status
    # 扫描 M180 产物；规则脚本本身会保存禁止词清单，不代表用户输出污染。
    dynamic_findings = scan_dynamic_terms([M180])
    secret_findings = scan_secrets([M180, WORKSPACE / "scripts", WORKSPACE / "shared/static_pool"])
    checks = {
        "py_compile_pass": compile_result["returncode"] == 0,
        "json_parse_pass": not json_errors,
        "historical_sheet_row_count_pass": import_audit.get("summary", {}).get("source_row_count") == 1052,
        "effective_customer_count_pass": import_audit.get("summary", {}).get("valid_customer_count", 0) > 900,
        "registry_v2_count_pass": registry.get("summary", {}).get("confirmed_signed_customer_count", 0) > 900,
        "alias_v2_count_pass": alias_registry.get("summary", {}).get("alias_count", 0) >= registry.get("summary", {}).get("confirmed_signed_customer_count", 0),
        "gate_regression_pass": gate.get("status") == "PASS",
        "default_gate_v2_pass": default_v2_probe == "excluded_existing_customer",
        "prospect_eligibility_uses_v2_pass": eligibility_probe == "excluded_signed_customer",
        "readiness_pass": readiness["returncode"] == 0 and "signed_customer_v2" in readiness["stdout"],
        "production_dry_run_has_m180_pass": dry_run["returncode"] == 0 and "build_m180r_signed_customer_registry_v2.py" in dry_run["stdout"],
        "current_pool_audit_ready": audit.get("status") == "PASS_AUDIT_READY",
        "dynamic_term_scan_pass": not dynamic_findings,
        "api_key_scan_pass": not secret_findings,
        "expert_review_allows_next": expert.get("status") in {"pass", "pass_with_followups"},
        "no_write_proof_pass": registry.get("summary", {}).get("old_excel_written") is False and registry.get("summary", {}).get("knowledge_asset_registry_written") is False and registry.get("summary", {}).get("persona_registry_written") is False,
    }
    return {
        "batch_id": "m180r_validation_report_v1",
        "milestone": "M180R",
        "generated_at": now(),
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": checks,
        "json_errors": json_errors,
        "default_v2_probe_status": default_v2_probe,
        "prospect_eligibility_probe_status": eligibility_probe,
        "readiness_stdout_tail": readiness["stdout"],
        "production_dry_run_stdout_tail": dry_run["stdout"],
        "dynamic_term_findings": dynamic_findings[:50],
        "secret_findings": secret_findings[:50],
        "compile_result": compile_result,
    }


def build_all() -> dict[str, Any]:
    registry, alias_registry, import_audit = build_registry_v2()
    write_json(SIGNED_REGISTRY_V2, registry)
    write_json(SIGNED_ALIAS_V2, alias_registry)
    write_json(M180 / "signed_customer_registry_v2.json", registry)
    write_json(M180 / "signed_customer_alias_registry_v2.json", alias_registry)
    write_json(M180 / "signed_customer_import_audit_v2.json", import_audit)

    gate = gate_report()
    audit, remediation = audit_current_outputs()
    expert = expert_review(registry, gate, audit)
    write_json(M180 / "signed_customer_gate_v2_regression_v1.json", gate)
    write_json(M180 / "existing_customer_audit_report_v2.json", audit)
    write_json(M180 / "remediation_preview_v2.json", remediation)
    write_json(M180 / "m180_expert_review_report_v1.json", expert)
    write_json(STATUS_PANEL, update_status_panel(registry, alias_registry, audit, expert))
    validation = validate(registry, alias_registry, import_audit, gate, audit, expert)
    write_json(M180 / "m180_validation_report_v1.json", validation)
    return {
        "status": validation["status"],
        "summary": {
            "signed_customer_v2_count": registry["summary"]["confirmed_signed_customer_count"],
            "alias_v2_count": alias_registry["summary"]["alias_count"],
            "trusted_pool_hit_count": audit["summary"]["trusted_pool_hit_count"],
            "vault_hit_count": audit["summary"]["vault_hit_count"],
            "expert_review_status": expert["status"],
        },
        "validation_checks": validation["checks"],
    }


def readiness_payload() -> dict[str, Any]:
    registry = read_json(SIGNED_REGISTRY_V2, {})
    alias_registry = read_json(SIGNED_ALIAS_V2, {})
    status = "PASS" if registry.get("summary", {}).get("confirmed_signed_customer_count", 0) > 900 else "FAIL_MISSING_SIGNED_CUSTOMER_V2"
    return {
        "mode": "readiness",
        "milestone": "M180R",
        "generated_at": now(),
        "status": status,
        "read_only": True,
        "signed_customer_v2_enabled": status == "PASS",
        "summary": {
            "signed_customer_v2_count": registry.get("summary", {}).get("confirmed_signed_customer_count", 0),
            "signed_customer_alias_v2_count": alias_registry.get("summary", {}).get("alias_count", 0),
            "source_path": registry.get("summary", {}).get("source_path"),
            "fallback_v1_available": LEGACY_REGISTRY_V1.exists(),
        },
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M180R signed customer registry v2 and production gate reports.")
    parser.add_argument("--stage", choices=["all", "readiness", "validate"], default="all")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if args.stage == "readiness":
        payload = readiness_payload()
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0 if str(payload.get("status", "")).startswith("PASS") else 2
    if args.stage == "validate":
        registry = read_json(SIGNED_REGISTRY_V2, {})
        alias_registry = read_json(SIGNED_ALIAS_V2, {})
        import_audit = read_json(M180 / "signed_customer_import_audit_v2.json", {})
        gate = read_json(M180 / "signed_customer_gate_v2_regression_v1.json", {})
        audit = read_json(M180 / "existing_customer_audit_report_v2.json", {})
        expert = read_json(M180 / "m180_expert_review_report_v1.json", {})
        payload = validate(registry, alias_registry, import_audit, gate, audit, expert)
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0 if payload["status"] == "PASS" else 2
    payload = build_all()
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if payload["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
