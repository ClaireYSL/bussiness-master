from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - dependency is present in project runtime
    load_workbook = None

WORKSPACE = Path(__file__).resolve().parents[1]

DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone41r_source_material_inventory"
DEFAULT_REVIEW_MD = "docs/03-执行与校验/Milestone 41R-知识源盘点与可学习素材分层-v1.md"
REGISTRY_XLSX = Path("/Users/clairelu2026/26M3-Obsidian-潜客池/潜客池/知识资产注册表.xlsx")

SOURCE_ROOTS = {
    "dingtalk_2026_03_28": WORKSPACE / "static-pool-deps-20260419_151148/static-pool-raw-materials/dingtalk_2026_03_28",
    "geo_l1_cases": WORKSPACE / "static-pool-deps-20260419_151148/static-pool-raw-materials/geo_l1_cases",
    "l1_customer_cases_20260313": Path("/Users/clairelu2026/26M1-ObsidianVaultGeneral/Obsidian知识库极简版/L1客户案例20260313"),
}

ALLOWED_EXTENSIONS = {".md", ".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".png", ".jpg", ".jpeg", ".gif", ".webp"}
SUPPORTING_ATTACHMENT_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
STATUS_SEGMENTS = ("已转正式资产", "已入学习队列", "待评估", "明确不学")

PERSONA_KEYWORDS = {
    "retail_high_sku_brand": ["品牌零售", "美妆", "个护", "母婴", "宠物", "鞋服", "服饰", "SKU", "花西子", "自然堂", "珀莱雅", "森马", "安踏", "太平鸟", "宝洁", "联合利华", "Fila", "GAP", "babycare"],
    "retail_multi_store": ["门店", "连锁", "商超", "便利", "分货", "巡店", "全家", "7-Eleven", "鲜丰", "来伊份", "博士眼镜", "正新", "珠宝", "茶叶连锁"],
    "fnb_chain_beverage_coffee": ["茶饮", "咖啡", "蜜雪", "COCO", "CoCo", "奈雪", "Linlee", "柠檬茶", "饮料", "茶业", "八马"],
    "fnb_chain_standardized": ["餐饮", "老乡鸡", "张亮", "夸父", "遇见小面", "老娘舅", "乐凯撒", "快乐蜂", "五谷渔粉", "赛百味", "德克士"],
    "cbec_multi_platform_brand": ["跨境", "出海", "亚马逊", "Amazon", "独立站", "SmallRig", "乐其", "安克", "倍思", "出海标杆"],
    "cbec_platform_operator": ["TP", "DP", "代运营", "电商代运营", "平台运营", "东南亚"],
    "mfg_multi_factory_group": ["制造", "工厂", "多工厂", "产线", "新能源", "汽车", "装备", "工业", "晶科", "零跑", "昊志", "首帆", "CDMO", "医疗器械"],
    "mfg_rnd_sales_complex": ["研产销", "研发", "定制化", "订单交付", "PLM", "MES", "交付", "供应链"],
    "finance_data_democratization": ["银行", "金融", "保险", "证券", "农商", "信贷", "宜信", "数禾", "科技金融"],
    "internet_platform_operation": ["互联网", "招聘", "在线教育", "雪球", "快看", "西瓜创客", "智联招聘", "新东方"],
}

TRACK_BY_PERSONA = {
    "retail_high_sku_brand": "retail_consumer",
    "retail_multi_store": "retail_consumer",
    "fnb_chain_beverage_coffee": "retail_consumer",
    "fnb_chain_standardized": "retail_consumer",
    "cbec_multi_platform_brand": "cross_border_ecommerce",
    "cbec_platform_operator": "cross_border_ecommerce",
    "mfg_multi_factory_group": "advanced_manufacturing",
    "mfg_rnd_sales_complex": "advanced_manufacturing",
    "finance_data_democratization": "finance",
    "internet_platform_operation": "internet_platform",
}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M41R source material inventory and learnable queue.")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--review-md", default=DEFAULT_REVIEW_MD)
    parser.add_argument("--registry-xlsx", default=str(REGISTRY_XLSX))
    return parser


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _write_json(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str | Path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _normalize_title(value: str) -> str:
    text = Path(value or "").stem
    text = re.sub(r"[\s_\-—｜|·（）()【】\[\]「」『』]+", "", text)
    return text.lower()


def _safe_file_type(path: Path) -> str:
    suffix = path.suffix.lower().lstrip(".")
    return suffix or "unknown"


def _iter_source_files(root: Path) -> list[Path]:
    if not root.exists():
        return []
    files: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if any(part.startswith(".") for part in path.relative_to(root).parts):
            continue
        if path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        files.append(path)
    return sorted(files, key=lambda p: str(p).lower())


def _source_folder_status(source_root: str, rel: Path) -> str:
    if rel.suffix.lower() in SUPPORTING_ATTACHMENT_EXTENSIONS:
        return "supporting_attachment"
    if source_root == "l1_customer_cases_20260313":
        return "legacy_case_library"
    for part in rel.parts:
        if part in STATUS_SEGMENTS:
            return part
    return "unclassified"


def _status_from_folder(source_root: str, folder_status: str) -> str:
    if folder_status == "supporting_attachment":
        return "附件材料"
    if source_root == "l1_customer_cases_20260313":
        return "待评估"
    if folder_status in STATUS_SEGMENTS:
        return folder_status
    return "未分层"


def _learnability(source_root: str, folder_status: str) -> str:
    if folder_status == "supporting_attachment":
        return "supporting_attachment"
    if folder_status == "已转正式资产":
        return "already_asset"
    if folder_status == "已入学习队列":
        return "learnable_now"
    if folder_status == "待评估":
        return "needs_triage"
    if folder_status == "明确不学":
        return "excluded"
    if source_root == "l1_customer_cases_20260313":
        return "needs_triage"
    return "needs_triage"


def _next_action(learnability: str) -> str:
    return {
        "already_asset": "保留资产映射；后续用于画像证据映射，不重复学习。",
        "learnable_now": "进入学习抽取队列，优先提炼客户案例/解决方案信号。",
        "needs_triage": "先人工或脚本复核是否为真实客户案例、解决方案或权威行业材料。",
        "excluded": "保持排除，不进入知识资产与画像校准。",
        "supporting_attachment": "作为正文素材附件保留索引，不单独进入学习队列。",
    }.get(learnability, "待复核")


def _guess_persona_and_track(text: str) -> tuple[str, str, list[str]]:
    scores: Counter[str] = Counter()
    matched: list[str] = []
    for persona, keywords in PERSONA_KEYWORDS.items():
        for keyword in keywords:
            if keyword.lower() in text.lower():
                scores[persona] += 1
                matched.append(keyword)
    if not scores:
        return "unknown", "unknown", []
    persona, _ = scores.most_common(1)[0]
    return persona, TRACK_BY_PERSONA.get(persona, "unknown"), sorted(set(matched))[:8]


def _priority(learnability: str, persona_guess: str, source_root: str) -> str:
    if learnability in {"excluded", "supporting_attachment"}:
        return "P9"
    if learnability == "already_asset":
        return "P2"
    if learnability == "learnable_now" and persona_guess != "unknown":
        return "P0"
    if source_root == "l1_customer_cases_20260313" and persona_guess != "unknown":
        return "P1"
    if persona_guess != "unknown":
        return "P1"
    return "P2"


def _material_id(source_root: str, rel: Path) -> str:
    digest = hashlib.sha1(f"{source_root}/{rel.as_posix()}".encode("utf-8")).hexdigest()[:12]
    return f"m41r_{source_root}_{digest}"


def _load_registry(path: Path) -> dict[str, Any]:
    empty = {"available": False, "by_title": {}, "counts": {}}
    if load_workbook is None or not path.exists():
        return empty
    wb = load_workbook(path, read_only=True, data_only=True)
    result: dict[str, Any] = {"available": True, "by_title": defaultdict(list), "counts": {}}
    for sheet_name in ("raw_material_inventory", "knowledge_assets", "learning_queue"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        rows = []
        for r in range(2, ws.max_row + 1):
            item = {str(headers[c - 1]): ws.cell(r, c).value for c in range(1, ws.max_column + 1) if headers[c - 1]}
            title = item.get("material_title") or item.get("title") or item.get("proposed_title") or ""
            path_or_url = item.get("material_path_or_url") or item.get("source_path_or_url") or ""
            key = _normalize_title(str(title or path_or_url))
            if key:
                result["by_title"][key].append({"sheet": sheet_name, "row": item})
            rows.append(item)
        result["counts"][sheet_name] = len(rows)
    result["by_title"] = dict(result["by_title"])
    return result


def _registry_match(path: Path, registry: dict[str, Any]) -> dict[str, Any]:
    key = _normalize_title(path.name)
    matches = (registry.get("by_title") or {}).get(key, [])
    if not matches:
        return {"matched": False}
    sheets = sorted({m["sheet"] for m in matches})
    first = matches[0]["row"]
    return {
        "matched": True,
        "matched_sheets": sheets,
        "registry_material_status": first.get("material_status") or first.get("extraction_status") or "",
        "registry_persona_guess": first.get("persona_guess") or first.get("persona_ids") or first.get("trigger_persona_id") or "",
        "registry_track_guess": first.get("track_guess") or first.get("track_ids") or first.get("trigger_track_id") or "",
        "registry_priority": first.get("priority") or "",
        "registry_match_count": len(matches),
    }


def _build_inventory(registry: dict[str, Any]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for source_root, root in SOURCE_ROOTS.items():
        for file_path in _iter_source_files(root):
            rel = file_path.relative_to(root)
            folder_status = _source_folder_status(source_root, rel)
            material_status = _status_from_folder(source_root, folder_status)
            learnability = _learnability(source_root, folder_status)
            text_for_guess = f"{rel.as_posix()} {file_path.stem}"
            persona_guess, track_guess, matched_keywords = _guess_persona_and_track(text_for_guess)
            registry_info = _registry_match(file_path, registry)
            record = {
                "material_id": _material_id(source_root, rel),
                "material_title": file_path.stem,
                "material_path_or_url": str(file_path),
                "relative_path": rel.as_posix(),
                "source_root": source_root,
                "source_root_path": str(root),
                "file_type": _safe_file_type(file_path),
                "source_folder_status": folder_status,
                "track_guess": track_guess,
                "persona_guess": persona_guess,
                "matched_keywords": matched_keywords,
                "material_status": material_status,
                "learnability_status": learnability,
                "is_learnable": learnability in {"learnable_now", "needs_triage"},
                "already_formal_asset": learnability == "already_asset" or "knowledge_assets" in registry_info.get("matched_sheets", []),
                "priority": _priority(learnability, persona_guess, source_root),
                "next_action": _next_action(learnability),
                "exists": file_path.exists(),
                "file_size_bytes": file_path.stat().st_size,
                "registry_match": registry_info,
            }
            records.append(record)
    return records


def _build_summary(records: list[dict[str, Any]], registry: dict[str, Any]) -> dict[str, Any]:
    by_root = Counter(r["source_root"] for r in records)
    by_type = Counter(r["file_type"] for r in records)
    by_folder = Counter(r["source_folder_status"] for r in records)
    by_learnability = Counter(r["learnability_status"] for r in records)
    by_persona = Counter(r["persona_guess"] for r in records)
    source_roots = {
        name: {
            "path": str(path),
            "exists": path.exists(),
            "file_count": by_root.get(name, 0),
        }
        for name, path in SOURCE_ROOTS.items()
    }
    return {
        "source_root_count": len(SOURCE_ROOTS),
        "source_roots": source_roots,
        "total_material_count": len(records),
        "learnable_queue_count": sum(1 for r in records if r["is_learnable"]),
        "excluded_material_count": sum(1 for r in records if r["learnability_status"] == "excluded"),
        "already_asset_count": sum(1 for r in records if r["already_formal_asset"]),
        "registry_available": bool(registry.get("available")),
        "registry_counts": registry.get("counts") or {},
        "registry_matched_count": sum(1 for r in records if r["registry_match"].get("matched")),
        "counts_by_source_root": dict(by_root),
        "counts_by_file_type": dict(by_type),
        "counts_by_source_folder_status": dict(by_folder),
        "counts_by_learnability_status": dict(by_learnability),
        "top_persona_guesses": dict(by_persona.most_common(12)),
        "no_write_proof": {
            "old_workbook_write_enabled": False,
            "knowledge_asset_write_enabled": False,
            "persona_registry_write_enabled": False,
            "prospect_generation_enabled": False,
        },
        "next_milestone": {
            "milestone": "M42R",
            "name": "ICP/画像体系重校准",
            "input": "M41R learnable material queue + existing persona registry",
        },
    }


def _render_md(payload: dict[str, Any]) -> str:
    summary = payload["source_coverage_summary"]["summary"]
    lines = [
        "# Milestone 41R-知识源盘点与可学习素材分层-v1",
        "",
        "## 结论",
        "",
        f"- 素材总数：`{summary['total_material_count']}`",
        f"- 可学习/待分诊素材：`{summary['learnable_queue_count']}`",
        f"- 已转正式资产或已匹配资产素材：`{summary['already_asset_count']}`",
        f"- 明确不学素材：`{summary['excluded_material_count']}`",
        f"- 旧注册表可读：`{summary['registry_available']}`；本机路径匹配数：`{summary['registry_matched_count']}`",
        "",
        "## 源根目录",
        "",
    ]
    for name, info in summary["source_roots"].items():
        lines.append(f"- `{name}`：exists=`{info['exists']}`，file_count=`{info['file_count']}`，path=`{info['path']}`")
    lines.extend([
        "",
        "## 分层口径",
        "",
        "- `already_asset`：原素材目录已标为正式资产，或与知识资产表匹配；后续用于 M42R 画像证据映射，不重复学习。",
        "- `learnable_now`：原素材目录已入学习队列，可优先进入客户案例/解决方案提炼。",
        "- `needs_triage`：真实素材存在，但需要先判断是否可学习，尤其是大批 L1 md 案例库。",
        "- `excluded`：明确不学，保持排除。",
        "",
        "## 边界",
        "",
        "- 本轮只做源治理，不生成潜客。",
        "- 本轮不写旧主表、不写正式知识资产、不改画像 registry。",
        "- 潜客产出不能反向覆盖知识资产；M42R 只能生成画像刷新 proposal。",
    ])
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    registry = _load_registry(Path(args.registry_xlsx))
    inventory_records = _build_inventory(registry)
    learnable_queue = [r for r in inventory_records if r["is_learnable"]]
    excluded_register = [r for r in inventory_records if r["learnability_status"] == "excluded"]
    summary = _build_summary(inventory_records, registry)

    payload = {
        "batch_id": "milestone41r_source_material_inventory_v2",
        "generated_at": _now(),
        "inputs": {
            "source_roots": {name: str(path) for name, path in SOURCE_ROOTS.items()},
            "registry_xlsx": str(Path(args.registry_xlsx)),
        },
        "summary": summary,
        "materials": inventory_records,
    }
    learnable_payload = {
        "batch_id": "milestone41r_learnable_material_queue_v1",
        "generated_at": _now(),
        "summary": {
            "queue_count": len(learnable_queue),
            "by_priority": dict(Counter(r["priority"] for r in learnable_queue)),
            "by_source_root": dict(Counter(r["source_root"] for r in learnable_queue)),
            "by_persona_guess": dict(Counter(r["persona_guess"] for r in learnable_queue).most_common(12)),
        },
        "items": learnable_queue,
    }
    excluded_payload = {
        "batch_id": "milestone41r_excluded_material_register_v1",
        "generated_at": _now(),
        "summary": {
            "excluded_count": len(excluded_register),
            "by_source_root": dict(Counter(r["source_root"] for r in excluded_register)),
        },
        "items": excluded_register,
    }
    coverage_payload = {
        "batch_id": "milestone41r_source_coverage_summary_v1",
        "generated_at": _now(),
        "summary": summary,
    }
    review_payload = {
        "source_coverage_summary": coverage_payload,
    }

    output_dir = Path(args.output_dir)
    _write_json(output_dir / "milestone41r_source_material_inventory_v2.json", payload)
    _write_json(output_dir / "milestone41r_learnable_material_queue_v1.json", learnable_payload)
    _write_json(output_dir / "milestone41r_excluded_material_register_v1.json", excluded_payload)
    _write_json(output_dir / "milestone41r_source_coverage_summary_v1.json", coverage_payload)
    _write_text(args.review_md, _render_md(review_payload))

    print(json.dumps({
        "output_dir": str(output_dir),
        "review_md": args.review_md,
        "summary": summary,
    }, ensure_ascii=False, indent=2))

    roots_ok = all(info["exists"] and info["file_count"] > 0 for info in summary["source_roots"].values())
    minimum_counts_ok = (
        summary["counts_by_source_root"].get("dingtalk_2026_03_28", 0) >= 100
        and summary["counts_by_source_root"].get("geo_l1_cases", 0) >= 50
        and summary["counts_by_source_root"].get("l1_customer_cases_20260313", 0) >= 500
    )
    no_write_ok = not any(summary["no_write_proof"].values())
    return 0 if roots_ok and minimum_counts_ok and no_write_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
