from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

VAULT = Path.home() / "Documents/Obsidian-Codex/潜客池"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
SHARED_XLSX = VAULT / "内部运营-静态潜客池-共享版.xlsx"

SHARED_SHEETS = [
    "全量主表",
    "主线汇总",
    "高质量层",
    "L4_L5扩展层",
    "边界主体",
    "画像汇总",
]

MAIN_HEADERS = [
    "account_id",
    "account_canonical_name",
    "brand_name",
    "primary_track",
    "persona_tag",
    "management_persona_tags",
    "信息扎实度",
    "ICP匹配概率",
    "静态潜客记录成熟度",
    "公司产品与服务概述",
    "商业模式概述",
    "核心客户客群",
    "收入规模区间",
    "营收增长概述",
    "已上线系统概况",
    "数字化项目动态",
    "近一年重大事件",
    "admission_reason_summary",
    "knowledge_asset_refs",
    "talk_track_refs",
    "validation_gap",
    "source_note",
]


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Rebuild 静态潜客主表.xlsx from the current shared workbook.")
    parser.add_argument("--output-file", help="Optional JSON summary output.")
    return parser


def backup_if_exists(path: Path) -> str | None:
    if not path.exists():
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.stem}.recovery_backup_{timestamp}{path.suffix}")
    shutil.copy2(path, backup)
    return str(backup)


def append_sheet_values(source_ws, target_ws) -> None:
    for row in source_ws.iter_rows(values_only=True):
        target_ws.append(list(row))


def build_accounts_main(shared_ws, wb: Workbook) -> None:
    ws = wb.create_sheet("accounts_main")
    ws.append(MAIN_HEADERS)
    headers = [cell.value for cell in next(shared_ws.iter_rows(min_row=1, max_row=1))]
    idx = {str(value): i for i, value in enumerate(headers) if value}
    for row in shared_ws.iter_rows(min_row=2, values_only=True):
        record = {
            "account_id": "",
            "account_canonical_name": row[idx["公司主体"]] if "公司主体" in idx else "",
            "brand_name": row[idx["品牌名"]] if "品牌名" in idx else "",
            "primary_track": row[idx["主线"]] if "主线" in idx else "",
            "persona_tag": row[idx["业务形态画像"]] if "业务形态画像" in idx else "",
            "management_persona_tags": row[idx["管理诉求画像"]] if "管理诉求画像" in idx else "",
            "信息扎实度": row[idx["信息扎实度"]] if "信息扎实度" in idx else "",
            "ICP匹配概率": row[idx["ICP匹配概率"]] if "ICP匹配概率" in idx else "",
            "静态潜客记录成熟度": row[idx["静态潜客记录成熟度"]] if "静态潜客记录成熟度" in idx else "",
            "公司产品与服务概述": row[idx["公司产品与服务概述"]] if "公司产品与服务概述" in idx else "",
            "商业模式概述": row[idx["商业模式概述"]] if "商业模式概述" in idx else "",
            "核心客户客群": row[idx["核心客户客群"]] if "核心客户客群" in idx else "",
            "收入规模区间": row[idx["收入规模区间"]] if "收入规模区间" in idx else "",
            "营收增长概述": row[idx["营收增长概述"]] if "营收增长概述" in idx else "",
            "已上线系统概况": row[idx["已上线系统概况"]] if "已上线系统概况" in idx else "",
            "数字化项目动态": row[idx["数字化项目动态"]] if "数字化项目动态" in idx else "",
            "近一年重大事件": row[idx["近一年重大事件"]] if "近一年重大事件" in idx else "",
            "admission_reason_summary": row[idx["一话入池理由"]] if "一话入池理由" in idx else "",
            "knowledge_asset_refs": row[idx["主要知识资产引用"]] if "主要知识资产引用" in idx else "",
            "talk_track_refs": row[idx["主要切入话术引用"]] if "主要切入话术引用" in idx else "",
            "validation_gap": row[idx["待验证项"]] if "待验证项" in idx else "",
            "source_note": "2026-04-08 主事实表恢复：由共享主表重建",
        }
        ws.append([record.get(header, "") for header in MAIN_HEADERS])


def main() -> int:
    args = build_parser().parse_args()
    backup = backup_if_exists(MAIN_XLSX)

    shared_wb = load_workbook(SHARED_XLSX, read_only=True, data_only=True)
    restored_wb = Workbook()
    default = restored_wb.active
    restored_wb.remove(default)

    build_accounts_main(shared_wb["全量主表"], restored_wb)
    for name in SHARED_SHEETS:
        source_ws = shared_wb[name]
        target_ws = restored_wb.create_sheet(name)
        append_sheet_values(source_ws, target_ws)

    restored_wb.save(MAIN_XLSX)

    result = {
        "main_file": str(MAIN_XLSX),
        "backup_file": backup,
        "restored_from": str(SHARED_XLSX),
        "sheet_names": restored_wb.sheetnames,
    }
    if args.output_file:
        Path(args.output_file).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
