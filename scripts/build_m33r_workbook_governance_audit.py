from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import check_workbook_integrity, resolve_static_pool_paths


DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone33r_workbook_governance_audit"
DEFAULT_REVIEW_MD = "docs/03-执行与校验/Milestone 33R-三表口径治理审计-v1.md"

MAIN_TO_SHARED = {
    "account_canonical_name": "公司主体",
    "brand_name": "品牌名",
    "primary_track": "主线",
    "persona_tag": "业务形态画像",
    "信息扎实度": "信息扎实度",
    "ICP匹配概率": "ICP匹配概率",
    "静态潜客记录成熟度": "静态潜客记录成熟度",
    "公司产品与服务概述": "公司产品与服务概述",
    "商业模式概述": "商业模式概述",
    "核心客户客群": "核心客户客群",
    "收入规模区间": "收入规模区间",
    "营收增长概述": "营收增长概述",
    "已上线系统概况": "已上线系统概况",
    "数字化项目动态": "数字化项目动态",
    "近一年重大事件": "近一年重大事件",
    "admission_reason_summary": "一话入池理由",
    "knowledge_asset_refs": "主要知识资产引用",
    "validation_gap": "待验证项",
}
MAIN_PROFILE_COMPARE_FIELDS = [
    "account_canonical_name",
    "primary_track",
    "persona_tag",
    "静态潜客记录成熟度",
    "信息扎实度",
    "ICP匹配概率",
    "admission_reason_summary",
    "公司产品与服务概述",
    "商业模式概述",
]
HIGH_QUALITY_LEVELS = {"L1", "L2", "L3"}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M33R workbook governance audit without mutating workbooks.")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--review-md", default=DEFAULT_REVIEW_MD)
    return parser


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(value: object) -> str:
    return str(value or "").strip()


def _norm_name(value: object) -> str:
    return _clean(value).replace("（", "(").replace("）", ")").replace(" ", "")


def _write_json(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str | Path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        return [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()


def _by_key(rows: list[dict[str, Any]], key: str) -> dict[str, dict[str, Any]]:
    return {_clean(row.get(key)): row for row in rows if _clean(row.get(key))}


def _duplicates(rows: list[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        value = _clean(row.get(key))
        if value:
            grouped[value].append(row)
    return [{"key": key, "value": value, "count": len(items)} for value, items in sorted(grouped.items()) if len(items) > 1]


def _field_mismatches(
    left: dict[str, dict[str, Any]],
    right: dict[str, dict[str, Any]],
    fields: list[str],
) -> list[dict[str, Any]]:
    mismatches: list[dict[str, Any]] = []
    for account_id in sorted(set(left) & set(right)):
        left_row = left[account_id]
        right_row = right[account_id]
        diffs = []
        for field in fields:
            left_value = _clean(left_row.get(field))
            right_value = _clean(right_row.get(field))
            if left_value != right_value:
                diffs.append({"field": field, "main_value": left_value, "profile_value": right_value})
        if diffs:
            mismatches.append(
                {
                    "account_id": account_id,
                    "account_name": _clean(left_row.get("account_canonical_name") or right_row.get("account_canonical_name")),
                    "diffs": diffs,
                }
            )
    return mismatches


def _shared_mismatches(main_rows: list[dict[str, Any]], shared_by_name: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    mismatches: list[dict[str, Any]] = []
    for main in main_rows:
        name = _clean(main.get("account_canonical_name"))
        if not name:
            continue
        shared = shared_by_name.get(_norm_name(name))
        if not shared:
            continue
        diffs = []
        for main_field, shared_field in MAIN_TO_SHARED.items():
            main_value = _clean(main.get(main_field))
            shared_value = _clean(shared.get(shared_field))
            if main_value != shared_value:
                diffs.append(
                    {
                        "main_field": main_field,
                        "shared_field": shared_field,
                        "main_value": main_value,
                        "shared_value": shared_value,
                    }
                )
        if diffs:
            mismatches.append({"account_id": _clean(main.get("account_id")), "account_name": name, "diffs": diffs})
    return mismatches


def _main_shared_missing(main_rows: list[dict[str, Any]], shared_by_name: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    missing = []
    for row in main_rows:
        name = _clean(row.get("account_canonical_name"))
        if name and _norm_name(name) not in shared_by_name:
            missing.append(
                {
                    "account_id": _clean(row.get("account_id")),
                    "account_name": name,
                    "level": _clean(row.get("静态潜客记录成熟度")),
                    "track": _clean(row.get("primary_track")),
                    "persona": _clean(row.get("persona_tag")),
                }
            )
    return missing


def _shared_not_in_main(shared_rows: list[dict[str, Any]], main_by_name: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    missing = []
    for row in shared_rows:
        name = _clean(row.get("公司主体"))
        if name and _norm_name(name) not in main_by_name:
            missing.append(
                {
                    "company_name": name,
                    "level": _clean(row.get("静态潜客记录成熟度")),
                    "track": _clean(row.get("主线")),
                    "persona": _clean(row.get("业务形态画像")),
                }
            )
    return missing


def _missing_profile_queue(main_only: list[dict[str, Any]]) -> list[dict[str, Any]]:
    items = []
    for row in main_only:
        level = _clean(row.get("静态潜客记录成熟度"))
        priority = "P0" if level in HIGH_QUALITY_LEVELS else "P2"
        items.append(
            {
                "queue_item_id": f"m33r_missing_profile_{_clean(row.get('account_id'))}_v1",
                "queue_type": "missing_profile",
                "account_id": _clean(row.get("account_id")),
                "account_name": _clean(row.get("account_canonical_name")),
                "level": level,
                "priority": priority,
                "status": "open",
                "recommended_action": "为主表账户补建 account_profiles 档案行；高质量层优先。",
            }
        )
    return items


def _shared_rebuild_plan(main_rows: list[dict[str, Any]]) -> dict[str, Any]:
    level_counts = Counter(_clean(row.get("静态潜客记录成熟度")) for row in main_rows if _clean(row.get("静态潜客记录成熟度")))
    high_quality = [row for row in main_rows if _clean(row.get("静态潜客记录成熟度")) in HIGH_QUALITY_LEVELS]
    extended = [row for row in main_rows if _clean(row.get("静态潜客记录成熟度")) not in HIGH_QUALITY_LEVELS]
    return {
        "plan_id": "milestone33r_shared_view_rebuild_plan_v1",
        "source_of_truth": "静态潜客主表.xlsx/accounts_main",
        "target_workbook": "内部运营-静态潜客池-共享版.xlsx",
        "mode": "plan_only_no_write",
        "expected_full_rows": len(main_rows),
        "expected_high_quality_rows": len(high_quality),
        "expected_extended_rows": len(extended),
        "expected_level_counts": dict(level_counts),
        "rules": [
            "全量主表应由主表按公司主体投影生成。",
            "高质量层应包含 L1/L2/L3。",
            "L4_L5扩展层应包含 L4/L5。",
            "共享版不应作为 account/persona/level 的 source of truth。",
        ],
    }


def _safe_sync_candidates(
    *,
    missing_profiles: list[dict[str, Any]],
    shared_missing: list[dict[str, Any]],
    profile_mismatches: list[dict[str, Any]],
    shared_mismatches: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "package_id": "milestone33r_safe_sync_candidate_package_v1",
        "mode": "candidate_only_no_write",
        "actions": {
            "create_missing_profile_rows": missing_profiles,
            "rebuild_shared_view_from_main": {
                "candidate_count": len(shared_missing) + len(shared_mismatches),
                "recommended_action": "先备份共享版，再由主表和档案库投影重建共享版；不要手工逐行修。",
            },
            "review_profile_mismatches": profile_mismatches[:200],
            "review_shared_mismatches": shared_mismatches[:200],
        },
        "guards": [
            "真实写回前必须备份三表。",
            "先修 identity/account_id，再修字段差异。",
            "共享版只能从主表/档案库生成，不反向覆盖主表。",
            "删除任何行必须单独确认。",
        ],
    }


def _render_md(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 33R-三表口径治理审计-v1",
        "",
        "## 摘要",
        "",
        f"- 主表行数：`{summary['main_row_count']}`，唯一 account_id：`{summary['main_unique_account_ids']}`",
        f"- 档案库行数：`{summary['profile_row_count']}`，唯一 account_id：`{summary['profile_unique_account_ids']}`",
        f"- 共享版全量行数：`{summary['shared_row_count']}`，唯一公司主体：`{summary['shared_unique_company_names']}`",
        f"- 主表有、档案库无：`{summary['main_not_in_profile_count']}`",
        f"- 主表有、共享版无：`{summary['main_not_in_shared_count']}`",
        f"- 共享版有、主表无：`{summary['shared_not_in_main_count']}`",
        f"- 主表/档案字段差异对象：`{summary['profile_mismatch_account_count']}`",
        f"- 主表/共享字段差异对象：`{summary['shared_mismatch_account_count']}`",
        f"- 工作簿完整性：`{summary['workbook_integrity_ok']}`",
        "",
        "## 判断",
        "",
        "- 三表治理应优先执行。主表应是总账，档案库是事实承载，共享版是投影视图。",
        "- 当前共享版不是主表完整镜像，不能作为项目主统计口径。",
        "- 本包只审计和生成同步候选，不写入工作簿。",
        "",
        "## 推荐治理顺序",
        "",
        "1. 补齐主表有、档案库无的 profile 行。",
        "2. 处理主表/档案库核心字段差异。",
        "3. 从主表投影重建共享版全量、高质量层和扩展层。",
        "4. 重跑 workbook integrity 和层级统计。",
    ]
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    pool = resolve_static_pool_paths()
    main_rows = _rows(pool["main"], "accounts_main")
    profile_rows = _rows(pool["profile"], "account_profiles")
    shared_rows = _rows(pool["main_shared"], "全量主表")
    integrity = check_workbook_integrity([pool["main"], pool["profile"], pool["main_shared"], pool["governance"]], deep_scan=True)

    main_by_id = _by_key(main_rows, "account_id")
    profile_by_id = _by_key(profile_rows, "account_id")
    main_by_name = {_norm_name(row.get("account_canonical_name")): row for row in main_rows if _clean(row.get("account_canonical_name"))}
    shared_by_name = {_norm_name(row.get("公司主体")): row for row in shared_rows if _clean(row.get("公司主体"))}

    main_not_in_profile = [main_by_id[account_id] for account_id in sorted(set(main_by_id) - set(profile_by_id))]
    profile_not_in_main = [profile_by_id[account_id] for account_id in sorted(set(profile_by_id) - set(main_by_id))]
    main_not_in_shared = _main_shared_missing(main_rows, shared_by_name)
    shared_not_in_main = _shared_not_in_main(shared_rows, main_by_name)
    profile_mismatches = _field_mismatches(main_by_id, profile_by_id, MAIN_PROFILE_COMPARE_FIELDS)
    shared_mismatches = _shared_mismatches(main_rows, shared_by_name)
    missing_profile_queue = _missing_profile_queue(main_not_in_profile)
    shared_plan = _shared_rebuild_plan(main_rows)
    safe_sync = _safe_sync_candidates(
        missing_profiles=missing_profile_queue,
        shared_missing=main_not_in_shared,
        profile_mismatches=profile_mismatches,
        shared_mismatches=shared_mismatches,
    )
    summary = {
        "main_row_count": len(main_rows),
        "main_unique_account_ids": len(main_by_id),
        "profile_row_count": len(profile_rows),
        "profile_unique_account_ids": len(profile_by_id),
        "shared_row_count": len(shared_rows),
        "shared_unique_company_names": len(shared_by_name),
        "main_level_counts": dict(Counter(_clean(row.get("静态潜客记录成熟度")) for row in main_rows if _clean(row.get("静态潜客记录成熟度")))),
        "profile_level_counts": dict(Counter(_clean(row.get("静态潜客记录成熟度")) for row in profile_rows if _clean(row.get("静态潜客记录成熟度")))),
        "shared_level_counts": dict(Counter(_clean(row.get("静态潜客记录成熟度")) for row in shared_rows if _clean(row.get("静态潜客记录成熟度")))),
        "main_duplicate_account_id_count": len(_duplicates(main_rows, "account_id")),
        "profile_duplicate_account_id_count": len(_duplicates(profile_rows, "account_id")),
        "shared_duplicate_company_name_count": len(_duplicates(shared_rows, "公司主体")),
        "main_not_in_profile_count": len(main_not_in_profile),
        "profile_not_in_main_count": len(profile_not_in_main),
        "main_not_in_shared_count": len(main_not_in_shared),
        "shared_not_in_main_count": len(shared_not_in_main),
        "profile_mismatch_account_count": len(profile_mismatches),
        "shared_mismatch_account_count": len(shared_mismatches),
        "missing_profile_queue_count": len(missing_profile_queue),
        "workbook_integrity_ok": bool(integrity.get("ok")),
        "true_writeback_enabled": False,
        "destructive_action_enabled": False,
    }
    payload = {
        "batch_id": "milestone33r_workbook_governance_audit_package_v1",
        "generated_at": _now(),
        "static_pool_root": str(pool["root"]),
        "policy": {
            "main_is_source_of_truth": True,
            "profile_is_fact_archive": True,
            "shared_is_generated_view": True,
            "true_writeback_enabled": False,
            "destructive_action_enabled": False,
        },
        "summary": summary,
        "identity_mismatch_report": {
            "main_not_in_profile": [
                {
                    "account_id": _clean(row.get("account_id")),
                    "account_name": _clean(row.get("account_canonical_name")),
                    "level": _clean(row.get("静态潜客记录成熟度")),
                    "track": _clean(row.get("primary_track")),
                    "persona": _clean(row.get("persona_tag")),
                }
                for row in main_not_in_profile
            ],
            "profile_not_in_main": [
                {
                    "account_id": _clean(row.get("account_id")),
                    "account_name": _clean(row.get("account_canonical_name")),
                    "level": _clean(row.get("静态潜客记录成熟度")),
                    "track": _clean(row.get("primary_track")),
                    "persona": _clean(row.get("persona_tag")),
                }
                for row in profile_not_in_main
            ],
            "main_not_in_shared": main_not_in_shared,
            "shared_not_in_main": shared_not_in_main,
        },
        "level_mismatch_report": {
            "profile_mismatches": profile_mismatches,
            "shared_mismatches": shared_mismatches,
        },
        "duplicate_report": {
            "main_duplicate_account_ids": _duplicates(main_rows, "account_id"),
            "profile_duplicate_account_ids": _duplicates(profile_rows, "account_id"),
            "shared_duplicate_company_names": _duplicates(shared_rows, "公司主体"),
        },
        "missing_profile_queue": missing_profile_queue,
        "shared_view_rebuild_plan": shared_plan,
        "safe_sync_candidate_package": safe_sync,
        "workbook_integrity": integrity,
    }
    output_dir = Path(args.output_dir)
    _write_json(output_dir / "milestone33r_workbook_governance_audit_package_v1.json", payload)
    _write_json(output_dir / "milestone33r_identity_mismatch_report_v1.json", payload["identity_mismatch_report"])
    _write_json(output_dir / "milestone33r_level_mismatch_report_v1.json", payload["level_mismatch_report"])
    _write_json(output_dir / "milestone33r_missing_profile_queue_v1.json", {"batch_id": "milestone33r_missing_profile_queue_v1", "items": missing_profile_queue})
    _write_json(output_dir / "milestone33r_shared_view_rebuild_plan_v1.json", shared_plan)
    _write_json(output_dir / "milestone33r_safe_sync_candidate_package_v1.json", safe_sync)
    _write_text(args.review_md, _render_md(payload))
    print(json.dumps({"output_json": str(output_dir / "milestone33r_workbook_governance_audit_package_v1.json"), "review_md": args.review_md, "summary": summary}, ensure_ascii=False, indent=2))
    return 0 if summary["workbook_integrity_ok"] and not summary["true_writeback_enabled"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
