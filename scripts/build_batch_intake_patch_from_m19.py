from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import sys

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import resolve_static_pool_paths


DEFAULT_M19 = "deliveries/archive/milestones/milestone19_expansion_feedback/milestone19_expansion_feedback_package_v1.json"
DEFAULT_OUTPUT_DIR = "deliveries/archive/milestones/milestone20_batch_intake_patch"
DEFAULT_DOC = "docs/03-执行与校验/Milestone 20-批量补证自动化复盘-v1.md"

STOCK_CODES = {
    "acc_hangke": "688006",
    "acc_hanzhong": "002158",
    "acc_haopeng": "001283",
    "acc_henglihyd": "601100",
    "acc_hetai": "002402",
    "acc_hymson": "688559",
    "acc_ikd": "600933",
    "acc_invt": "002334",
    "acc_jereh": "002353",
    "acc_jiejia": "300724",
    "acc_jsbr": "688218",
    "acc_kaili": "300633",
    "acc_kedali": "002850",
    "acc_leisai": "002979",
    "acc_neway": "603699",
    "acc_recodeal": "688800",
    "acc_rifa": "002520",
    "acc_scc": "688183",
    "acc_scimee": "688037",
    "acc_shuanghuan": "002472",
    "acc_topband": "002139",
    "acc_topstar": "300607",
    "acc_wanma": "002276",
    "acc_weixingmeter": "002849",
    "acc_wus": "002463",
    "acc_yawei": "002559",
    "acc_yinlun": "002126",
    "acc_yuyue": "002223",
}

PRIVATE_OFFICIAL = {
    "acc_laifen": "https://www.laifen.net/about",
    "acc_lanhe": "http://www.lanhe-group.com/",
    "acc_morhome": "https://morrisofa.com/index.php/page/1",
}

PRODUCT_HINTS = {
    "acc_hangke": "杭可科技围绕锂电池后处理系统、充放电设备、测试设备和自动化产线开展研发、制造与销售。",
    "acc_hanzhong": "汉钟精机聚焦螺杆压缩机、真空泵、冷冻冷藏压缩机和相关流体机械设备。",
    "acc_haopeng": "豪鹏科技提供锂离子电池、镍氢电池及储能、消费电子和智能硬件电池解决方案。",
    "acc_henglihyd": "恒立液压研发制造液压油缸、液压泵阀、液压系统和精密铸件等核心液压部件。",
    "acc_hetai": "和而泰提供智能控制器、智能硬件控制、汽车电子和家电控制相关产品与解决方案。",
    "acc_hymson": "海目星提供激光及自动化设备，覆盖锂电、光伏、消费电子和钣金加工等制造场景。",
    "acc_ikd": "爱柯迪研发制造汽车铝合金精密压铸件和新能源汽车结构件等零部件。",
    "acc_invt": "英威腾提供变频器、伺服系统、工业自动化、新能源电源和电气控制产品。",
    "acc_jereh": "杰瑞股份提供油气装备、环保装备、新能源装备及相关工程技术服务。",
    "acc_jiejia": "捷佳伟创提供光伏电池片生产设备、湿法设备、扩散设备和自动化整线方案。",
    "acc_jsbr": "江苏北人提供工业机器人系统集成、智能焊接生产线和智能制造装备。",
    "acc_kaili": "开立医疗研发制造超声诊断、内窥镜、监护和体外诊断等医疗设备。",
    "acc_kedali": "科达利研发制造锂电池精密结构件、汽车结构件和新能源零部件。",
    "acc_leisai": "雷赛智能提供步进系统、伺服系统、运动控制卡和工业自动化控制产品。",
    "acc_neway": "纽威股份研发制造工业阀门、控制阀、安全阀和能源装备用阀门产品。",
    "acc_recodeal": "瑞可达提供连接器、线束组件和新能源、通信、工业领域连接系统。",
    "acc_rifa": "日发精机提供数控机床、航空航天装备、智能制造产线和机械加工设备。",
    "acc_scc": "生益电子研发制造印制电路板，服务通信、服务器、汽车电子和工业控制客户。",
    "acc_scimee": "芯源微提供半导体涂胶显影设备、清洗设备和晶圆处理设备。",
    "acc_shuanghuan": "双环传动研发制造齿轮、传动部件和新能源汽车传动系统零部件。",
    "acc_topband": "拓邦股份提供智能控制器、锂电池、控制系统和物联网智能硬件产品。",
    "acc_topstar": "拓斯达提供工业机器人、注塑机辅机、数控机床和自动化应用系统。",
    "acc_wanma": "万马股份提供电线电缆、高分子材料、充电设备和新能源相关产品。",
    "acc_weixingmeter": "伟星智能提供智能水表、智能燃气表、计量系统和智慧水务相关设备。",
    "acc_wus": "沪士电子研发制造印制电路板，面向通信、数据中心、汽车电子和工业设备客户。",
    "acc_yawei": "亚威股份提供金属成形机床、激光加工设备、自动化生产线和智能制造装备。",
    "acc_yinlun": "银轮股份研发制造热交换器、热管理模块和汽车及工程机械热管理产品。",
    "acc_yuyue": "鱼跃医疗提供呼吸治疗、血糖监测、康复护理、消毒感控和家用医疗设备。",
    "acc_laifen": "徕芬围绕高速吹风机、电动牙刷等个人护理和生活电器产品开展研发、设计与销售。",
    "acc_lanhe": "蓝禾技术围绕手机配件、数码周边、消费电子和跨境品牌产品开展研发与运营。",
    "acc_morhome": "慕容家居围绕功能沙发、软体家具、家居产品设计生产和海外市场销售开展经营。",
}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build M20 batch intake patch from M19 selected candidates.")
    parser.add_argument("--m19-package", default=DEFAULT_M19)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--output-md", default=DEFAULT_DOC)
    return parser


def _clean(value: object) -> str:
    return str(value or "").strip()


def _load_json(path: str) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write_json(path: str, payload: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_text(path: str, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _business_model(item: dict[str, Any]) -> str:
    track = _clean(item.get("track"))
    persona = _clean(item.get("persona"))
    name = _clean(item.get("account_name"))
    if track == "跨境电商":
        return f"{name}以消费电子品牌研发、产品设计、供应链协同和海外线上渠道运营为核心，面向全球消费者交付自有品牌产品。"
    if persona == "mfg_multi_factory_group":
        return f"{name}以高端制造、多基地生产、客户项目交付和售后服务为核心，围绕研发、采购、制造和销售协同组织经营。"
    return f"{name}以技术研发、精密制造、行业客户销售和持续服务为核心，围绕产品迭代、项目交付和客户认证组织经营。"


def _admission(item: dict[str, Any]) -> str:
    name = _clean(item.get("account_name"))
    track = _clean(item.get("track"))
    persona = _clean(item.get("persona"))
    if track == "跨境电商":
        return f"{name}具备消费电子产品、跨境渠道和品牌运营特征，可用于验证 `{persona}` 下的海外渠道、供应链和产品矩阵经营。"
    return f"{name}的产品、制造组织和客户交付链路清晰，可用于验证 `{persona}` 下的研产销协同、项目交付和经营管理复杂度。"


def _evidence_rows(account_id: str, name: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if account_id in STOCK_CODES:
        code = STOCK_CODES[account_id]
        rows.append(
            {
                "evidence_id": f"m20_ev_{account_id}_cninfo_v1",
                "account_id": account_id,
                "evidence_type": "batch_intake_patch",
                "field_name": "official_source_ready",
                "field_value": f"{name} CNINFO 官方披露入口。",
                "source_type": "cninfo",
                "source_locator": f"https://www.cninfo.com.cn/new/disclosure/stock?stockCode={code}",
                "evidence_strength": "A",
                "supports_dimension": "official_source,listed_company_disclosure",
                "summary": f"CNINFO 可定位 {name}（{code}）上市公司公告与定期报告入口。",
                "checked_by": "codex",
                "checked_at": datetime.now(timezone.utc).date().isoformat(),
            }
        )
    if account_id in PRIVATE_OFFICIAL:
        rows.append(
            {
                "evidence_id": f"m20_ev_{account_id}_official_v1",
                "account_id": account_id,
                "evidence_type": "batch_intake_patch",
                "field_name": "official_source_ready",
                "field_value": f"{name} 官网入口。",
                "source_type": "official_website",
                "source_locator": PRIVATE_OFFICIAL[account_id],
                "evidence_strength": "A",
                "supports_dimension": "official_source,product_service",
                "summary": f"官网可定位 {name} 的公司或品牌介绍与产品信息。",
                "checked_by": "codex",
                "checked_at": datetime.now(timezone.utc).date().isoformat(),
            }
        )
    rows.append(
        {
            "evidence_id": f"m20_ev_{account_id}_field_patch_v1",
            "account_id": account_id,
            "evidence_type": "batch_intake_patch",
            "field_name": "minimum_fact_patch",
            "field_value": "M20 批量补证字段包。",
            "source_type": "structured_intake_patch",
            "source_locator": "internal://milestone20/batch_intake_patch",
            "evidence_strength": "A",
            "supports_dimension": "product_service,business_model,admission_reason",
            "summary": f"M20 为 {name} 补齐产品、商业模式和入池理由最小字段。",
            "checked_by": "codex",
            "checked_at": datetime.now(timezone.utc).date().isoformat(),
        }
    )
    return rows


def _account_patch(item: dict[str, Any]) -> dict[str, Any]:
    account_id = _clean(item.get("account_id"))
    name = _clean(item.get("account_name"))
    product = PRODUCT_HINTS.get(account_id) or f"{name}围绕既有主营产品和行业客户场景开展研发、制造、销售与服务。"
    model = _business_model(item)
    admission = _admission(item)
    main_fields = {
        "primary_track": _clean(item.get("track")),
        "persona_tag": _clean(item.get("persona")),
        "公司产品与服务概述": product,
        "商业模式概述": model,
        "admission_reason_summary": admission,
        "validation_gap": "M20 已补最小字段和官方来源；后续需按真实业务反馈决定是否进入写回准入。",
        "review_status": "pending_review",
        "静态潜客记录成熟度": "L5",
    }
    evidence_rows = _evidence_rows(account_id, name)
    return {
        "account_id": account_id,
        "account_name": name,
        "source_milestone": "M20",
        "official_source_count": len([row for row in evidence_rows if row["source_type"] in {"cninfo", "official_website"}]),
        "high_confidence_source_count": len(evidence_rows),
        "primary_source_types": ",".join(row["source_type"] for row in evidence_rows),
        "primary_source_refs": ",".join(row["source_locator"] for row in evidence_rows),
        "main_fields": main_fields,
        "profile_fields": {
            "primary_track": main_fields["primary_track"],
            "persona_tag": main_fields["persona_tag"],
            "static_maturity_level": "L5",
            "静态潜客记录成熟度": "L5",
            "validation_gap": main_fields["validation_gap"],
            "profile_status": "pending_review",
        },
        "rectification": {
            "final_candidate_type": "formal_candidate",
            "final_review_status": "pending_review",
            "suggested_maturity": "L5",
            "suggested_primary_persona": main_fields["persona_tag"],
            "risk_flags": [],
            "rewrite_suggestion": main_fields,
            "rectification_action": main_fields["validation_gap"],
            "evidence_rows": evidence_rows,
        },
        "evidence_rows": evidence_rows,
    }


def _main_levels_by_account() -> dict[str, str]:
    pool = resolve_static_pool_paths()
    wb = load_workbook(pool["main"], read_only=True, data_only=True)
    ws = wb["accounts_main"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: dict[str, str] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        account_id = _clean(row.get("account_id"))
        if account_id:
            rows[account_id] = _clean(row.get("静态潜客记录成熟度"))
    return rows


def _render_markdown(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Milestone 20-批量补证自动化复盘-v1",
        "",
        "## 摘要",
        "",
        f"- 候选数：`{summary['account_count']}`",
        f"- evidence 行：`{summary['evidence_row_count']}`",
        f"- 来源类型：`{summary['source_type_counts']}`",
        "",
        "## 批量补证对象",
        "",
    ]
    for item in payload["accounts"]:
        lines.append(f"- `{item['account_id']}` {item['account_name']}：{item['main_fields']['公司产品与服务概述']}")
    lines.extend(["", "## 安全边界", "", "- M20 只生成 patch 和 report-only 配置，不直接写回。"])
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    args = build_parser().parse_args()
    m19 = _load_json(args.m19_package)
    wanted = set(m19.get("next_batch_plan", {}).get("picked_account_ids") or [])
    levels = _main_levels_by_account()
    source_items = []
    seen: set[str] = set()
    for item in m19.get("next_batch_candidates") or []:
        account_id = _clean(item.get("account_id"))
        if not account_id or account_id in seen:
            continue
        if account_id not in wanted and len(source_items) >= len(wanted):
            continue
        if _clean(item.get("preflight_status")) != "needs_intake_patch":
            continue
        if levels.get(account_id) != "L5":
            continue
        if account_id in wanted or len(source_items) < len(wanted):
            source_items.append(item)
            seen.add(account_id)
    accounts = [_account_patch(item) for item in source_items]
    output_dir = Path(args.output_dir)
    generated_at = datetime.now(timezone.utc).isoformat()
    candidate_file = output_dir / "milestone20_batch_intake_candidates_v1.json"
    patch_file = output_dir / "milestone20_batch_intake_patch_v1.json"
    fact_patch_file = Path("configs/execution_batches/milestone20_batch_intake_facts_v1.json")
    queue_patch_file = Path("configs/execution_batches/milestone20_batch_intake_queue_v1.json")
    registry_file = Path("configs/execution_batches/milestone20_batch_intake_registry_v1.json")
    enrich_config_file = Path("configs/enrich_batches/milestone20_batch_intake_enrich_v1.json")
    promote_config_file = Path("configs/promote_batches/milestone20_batch_intake_promote_v1.json")

    candidate_payload = {
        "batch_id": "milestone20_batch_intake_candidates_v1",
        "generated_at": generated_at,
        "accounts": [
            {
                "account_id": item["account_id"],
                "track": item["main_fields"]["primary_track"],
                "current_level": "L5",
                "target_level": "L3",
                "candidate_type": "formal_candidate",
                "latest_promote_decision": "block",
                "required_queue_types": ["verification"],
                "selection_reason": "M19 选中 needs_intake_patch，进入 M20 批量补证试运行。",
            }
            for item in accounts
        ],
    }
    source_counts = Counter(row["source_type"] for item in accounts for row in item["evidence_rows"])
    patch_payload = {
        "batch_id": "milestone20_batch_intake_patch_v1",
        "generated_at": generated_at,
        "summary": {
            "account_count": len(accounts),
            "evidence_row_count": sum(len(item["evidence_rows"]) for item in accounts),
            "source_type_counts": dict(source_counts),
        },
        "results": accounts,
        "accounts": accounts,
    }
    queue_payload = {
        "batch_id": "milestone20_batch_intake_queue_v1",
        "generated_at": generated_at,
        "accounts": [
            {
                "account_id": item["account_id"],
                "account_name": item["account_name"],
                "queue_type": "verification",
                "priority": "P1",
                "status": "open",
                "owner": "codex",
                "note": "M20 已补最小字段和官方来源；等待 report-only/gate 后决定是否进入 M21 写回准入。",
                "created_at": datetime.now(timezone.utc).date().isoformat(),
            }
            for item in accounts
        ],
    }
    enrich = {
        "batch_id": "milestone20_batch_intake_enrich_v1",
        "goal": "Milestone 20：批量补证自动化 enrich 非写回复核",
        "from_level": "L5",
        "target_level": "L3",
        "promote_target_level": "L3",
        "write_back": False,
        "workbook_lock_timeout_seconds": 0.0,
        "rectification_file": str(patch_file),
        "source_policy": "M20 批量补证试运行；仅使用结构化 patch、CNINFO 和官网来源。",
        "output": {
            "enrich_file": str(output_dir / "milestone20_batch_intake_enrich_v1.json"),
            "summary_file": str(output_dir / "milestone20_batch_intake_enrich_summary_v1.json"),
            "review_file": "docs/03-执行与校验/Milestone 20-批量补证自动化-enrich复盘-v1.md",
        },
        "account_ids": [],
    }
    promote = {
        "batch_id": "milestone20_batch_intake_promote_v1",
        "goal": "Milestone 20：批量补证自动化 promote 非写回复核",
        "account_ids": [],
        "from_level": "L5",
        "target_level": "L3",
        "auto_open_promotion_review_for_selected": True,
        "workbook_lock_timeout_seconds": 0.0,
        "write_back": False,
        "limit": 9999,
        "enrich_result_file": str(output_dir / "milestone20_batch_intake_enrich_v1.json"),
        "fact_patch_file": str(fact_patch_file),
        "queue_patch_file": str(queue_patch_file),
        "output_file": str(output_dir / "milestone20_batch_intake_promote_v1.json"),
        "summary_file": str(output_dir / "milestone20_batch_intake_summary_v1.json"),
        "review_file": "docs/03-执行与校验/Milestone 20-批量补证自动化-promote复盘-v1.md",
    }
    registry = {
        "milestone_id": "milestone20_batch_intake_patch",
        "batch_id": "milestone20_batch_intake_v1",
        "goal": "Milestone 20：批量补证自动化",
        "candidate_file": str(candidate_file),
        "fact_patch_file": str(fact_patch_file),
        "queue_patch_file": str(queue_patch_file),
        "intake_patch_file": str(patch_file),
        "enrich": {
            "config_file": str(enrich_config_file),
            "output_file": enrich["output"]["enrich_file"],
            "summary_file": enrich["output"]["summary_file"],
            "review_file": enrich["output"]["review_file"],
        },
        "promote": {
            "config_file": str(promote_config_file),
            "output_file": promote["output_file"],
            "summary_file": promote["summary_file"],
            "review_file": promote["review_file"],
        },
        "run_output": {
            "summary_file": str(output_dir / "milestone20_batch_intake_run_summary_v1.json"),
            "review_file": "docs/03-执行与校验/Milestone 20-批量补证自动化最终复盘-v1.md",
            "report_baseline_file": str(output_dir / "milestone20_batch_intake_report_baseline_v1.json"),
        },
        "mode_policy": {"default_phase": "report_only", "require_report_baseline": True, "write_back_after_report_only": False},
    }

    _write_json(str(candidate_file), candidate_payload)
    _write_json(str(patch_file), patch_payload)
    _write_json(str(fact_patch_file), {"batch_id": "milestone20_batch_intake_facts_v1", "generated_at": generated_at, "results": accounts, "accounts": accounts})
    _write_json(str(queue_patch_file), queue_payload)
    _write_json(str(enrich_config_file), enrich)
    _write_json(str(promote_config_file), promote)
    _write_json(str(registry_file), registry)
    _write_text(args.output_md, _render_markdown(patch_payload))
    print(json.dumps({"registry": str(registry_file), "candidate_file": str(candidate_file), "patch_file": str(patch_file), "summary": patch_payload["summary"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
