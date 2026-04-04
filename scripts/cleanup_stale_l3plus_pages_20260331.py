from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/clairaipartner")
VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
L12_DIR = VAULT / "07-L3以上客户档案/01-L1-L2档案"
L3_DIR = VAULT / "07-L3以上客户档案/02-L3档案"
TRASH_BASE = ROOT / ".Trash/codex-l3plus-cleanup-20260331"


def normalize_filename(name: str) -> str:
    return str(name).replace("/", "-")


def main() -> None:
    wb = load_workbook(PROFILE_XLSX, data_only=True)
    ws = wb["account_profiles"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    expected_l12 = set()
    expected_l3 = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        level = str(row[idx["静态潜客记录成熟度"]] or "")
        if level not in {"L1", "L2", "L3"}:
            continue
        name = normalize_filename(str(row[idx["account_canonical_name"]] or "").strip()) + ".md"
        if level in {"L1", "L2"}:
            expected_l12.add(name)
        else:
            expected_l3.add(name)

    extra_l12 = [p for p in L12_DIR.glob("*.md") if p.name not in expected_l12]
    extra_l3 = [p for p in L3_DIR.glob("*.md") if p.name not in expected_l3]

    for bucket, files in [("01-L1-L2档案", extra_l12), ("02-L3档案", extra_l3)]:
        target_dir = TRASH_BASE / bucket
        target_dir.mkdir(parents=True, exist_ok=True)
        for p in files:
            dest = target_dir / p.name
            if dest.exists():
                dest = target_dir / f"{p.stem}__dup{p.suffix}"
            shutil.move(str(p), str(dest))

    print(
        {
            "moved_l12": len(extra_l12),
            "moved_l3": len(extra_l3),
            "trash_dir": str(TRASH_BASE),
        }
    )


if __name__ == "__main__":
    main()
