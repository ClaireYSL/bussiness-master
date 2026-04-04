from __future__ import annotations

import json
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


REGISTRY_XLSX = Path("/Users/clairaipartner/Documents/Obsidian-Codex/潜客池/知识资产注册表.xlsx")
DEFAULT_DEST = Path("/Users/clairaipartner/Codex/static-pool-raw-materials")


def load_raw_material_rows() -> list[dict[str, object]]:
    wb = load_workbook(REGISTRY_XLSX, read_only=True, data_only=True)
    ws = wb["raw_material_inventory"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[i]: row[i] for i in range(len(headers))}
        path = str(item.get("material_path_or_url") or "").strip()
        if path.startswith("/"):
            rows.append(item)
    return rows


def safe_relative_name(path: Path) -> str:
    return path.name.replace("/", "_")


def build_destination(base_dir: Path, row: dict[str, object], seen: Counter[str]) -> Path:
    source_root = str(row.get("source_root") or "unknown_root").strip() or "unknown_root"
    file_type = str(row.get("file_type") or "").strip().lower() or "unknown"
    status = str(row.get("material_status") or "").strip() or "unknown_status"
    source_path = Path(str(row["material_path_or_url"]))
    target_dir = base_dir / source_root / status / file_type
    target_dir.mkdir(parents=True, exist_ok=True)

    name = safe_relative_name(source_path)
    destination = target_dir / name
    key = str(destination)
    if seen[key]:
        destination = target_dir / f"{source_path.stem}__{row.get('material_id')}{source_path.suffix}"
    seen[str(destination)] += 1
    return destination


def copy_rows(rows: list[dict[str, object]], base_dir: Path) -> dict[str, object]:
    seen_destinations: Counter[str] = Counter()
    copied = []
    missing = []
    root_counter: Counter[str] = Counter()
    status_counter: Counter[str] = Counter()

    for row in rows:
        source_path = Path(str(row["material_path_or_url"]))
        source_root = str(row.get("source_root") or "").strip() or "unknown_root"
        status = str(row.get("material_status") or "").strip() or "unknown_status"
        root_counter[source_root] += 1
        status_counter[status] += 1

        if not source_path.exists():
            missing.append(str(source_path))
            continue

        destination = build_destination(base_dir, row, seen_destinations)
        shutil.copy2(source_path, destination)
        copied.append(
            {
                "material_id": row.get("material_id"),
                "material_title": row.get("material_title"),
                "source_root": source_root,
                "material_status": status,
                "source_path": str(source_path),
                "destination_path": str(destination),
            }
        )

    manifest = {
        "registry_path": str(REGISTRY_XLSX),
        "destination_root": str(base_dir),
        "copied_count": len(copied),
        "missing_count": len(missing),
        "source_root_counter": dict(root_counter),
        "status_counter": dict(status_counter),
        "missing_paths": missing,
        "copied_items": copied,
    }
    return manifest


def write_outputs(base_dir: Path, manifest: dict[str, object]) -> None:
    manifest_dir = base_dir / "_manifest"
    manifest_dir.mkdir(parents=True, exist_ok=True)
    (manifest_dir / "copy_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        "# 已登记原素材复制清单",
        "",
        f"- 注册表：`{manifest['registry_path']}`",
        f"- 目标根目录：`{manifest['destination_root']}`",
        f"- 已复制：`{manifest['copied_count']}`",
        f"- 缺失：`{manifest['missing_count']}`",
        "",
        "## 来源根分布",
        "",
    ]
    for key, value in sorted(manifest["source_root_counter"].items()):
        lines.append(f"- `{key}`：`{value}`")
    lines.extend(["", "## 状态分布", ""])
    for key, value in sorted(manifest["status_counter"].items()):
        lines.append(f"- `{key}`：`{value}`")
    lines.extend(["", "## 说明", "", "- 本次仅复制，不删除旧文件。", "- 本次未回写 `raw_material_inventory / learning_queue / knowledge_assets` 路径。"])
    (manifest_dir / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str]) -> int:
    destination_root = Path(argv[1]) if len(argv) > 1 else DEFAULT_DEST
    destination_root.mkdir(parents=True, exist_ok=True)
    rows = load_raw_material_rows()
    manifest = copy_rows(rows, destination_root)
    write_outputs(destination_root, manifest)
    print(json.dumps({k: manifest[k] for k in ("destination_root", "copied_count", "missing_count", "source_root_counter", "status_counter")}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
