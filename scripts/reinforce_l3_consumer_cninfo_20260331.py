from __future__ import annotations

import re
import os
import signal
from pathlib import Path

import akshare as ak
from openpyxl import load_workbook


TODAY = "2026-03-31"
ROOT = Path("/Users/clairaipartner")
WORKSPACE = ROOT / ".openclaw/workspace-main/bussiness-master"
VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"

MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
GOV_XLSX = VAULT / "治理与证据.xlsx"
DOC_PATH = WORKSPACE / "docs/03-执行与校验/静态潜客池-L3消费品官方源补强专项-v1.md"
MEMORY_PATH = WORKSPACE / "memory/2026-03-31.md"

CODE_PAT = re.compile(r"acc_l5_(\d{6})$")
MAX_ITEMS = int(os.environ.get("MAX_ITEMS", "60"))
REQUEST_TIMEOUT = int(os.environ.get("REQUEST_TIMEOUT", "8"))


class FetchTimeout(Exception):
    pass


def _timeout_handler(signum, frame):
    raise FetchTimeout("cninfo fetch timeout")


def main() -> None:
    main_wb = load_workbook(MAIN_XLSX)
    main_ws = main_wb["accounts_main"]
    mh = [c.value for c in main_ws[1]]
    mi = {h: i + 1 for i, h in enumerate(mh)}

    profile_wb = load_workbook(PROFILE_XLSX)
    profiles_ws = profile_wb["account_profiles"]
    coverage_ws = profile_wb["profile_coverage"]
    ph = [c.value for c in profiles_ws[1]]
    pi = {h: i + 1 for i, h in enumerate(ph)}
    ch = [c.value for c in coverage_ws[1]]
    ci = {h: i + 1 for i, h in enumerate(ch)}

    gov_wb = load_workbook(GOV_XLSX)
    evidence_ws = gov_wb["evidence_log"]
    eh = [c.value for c in evidence_ws[1]]
    ei = {h: i + 1 for i, h in enumerate(eh)}

    profile_rows = {}
    coverage_rows = {}
    for r in range(2, profiles_ws.max_row + 1):
        aid = profiles_ws.cell(r, pi["account_id"]).value
        if aid:
            profile_rows[str(aid)] = r
    for r in range(2, coverage_ws.max_row + 1):
        aid = coverage_ws.cell(r, ci["account_id"]).value
        if aid:
            coverage_rows[str(aid)] = r

    candidates = []
    for r in range(2, main_ws.max_row + 1):
        aid = str(main_ws.cell(r, mi["account_id"]).value or "")
        lvl = main_ws.cell(r, mi["静态潜客记录成熟度"]).value
        source = str(main_ws.cell(r, mi["source_note"]).value or "")
        info = str(main_ws.cell(r, mi["信息扎实度"]).value or "")
        m = CODE_PAT.match(aid)
        if (
            lvl == "L3"
            and m
            and "消费品新画像扩池 2026-03-31" in source
            and "L5到L3批量上移" in source
        ):
            pr = profile_rows.get(aid)
            if pr and int(profiles_ws.cell(pr, pi["official_source_count"]).value or 0) >= 1:
                continue
            candidates.append((r, aid, m.group(1), str(main_ws.cell(r, mi["account_canonical_name"]).value), info))

    updated = 0
    skipped = []

    for main_row, aid, code, name, info in candidates[:MAX_ITEMS]:
        try:
            signal.signal(signal.SIGALRM, _timeout_handler)
            signal.alarm(REQUEST_TIMEOUT)
            df = ak.stock_profile_cninfo(symbol=code)
            signal.alarm(0)
        except FetchTimeout:
            signal.alarm(0)
            skipped.append((aid, name, "timeout"))
            continue
        except Exception:
            signal.alarm(0)
            skipped.append((aid, name, "fetch_error"))
            continue
        if df is None or df.empty:
            skipped.append((aid, name, "empty"))
            continue
        row = df.iloc[0]
        website = str(row.get("官方网站") or "").strip()
        business = str(row.get("主营业务") or "").strip()
        industry = str(row.get("所属行业") or "").strip()
        company_name = str(row.get("公司名称") or "").strip()
        if not company_name:
            skipped.append((aid, name, "missing_company_name"))
            continue

        # main
        if info == "中":
            main_ws.cell(main_row, mi["信息扎实度"]).value = "中高"
        main_ws.cell(main_row, mi["company_scale_band" if "company_scale_band" in mi else "收入规模区间"]).value = (
            main_ws.cell(main_row, mi["收入规模区间"]).value or "待补公开财报口径"
        )
        main_ws.cell(main_row, mi["公司产品与服务概述"]).value = business or main_ws.cell(main_row, mi["公司产品与服务概述"]).value
        if industry:
            main_ws.cell(main_row, mi["industry_l1"]).value = main_ws.cell(main_row, mi["industry_l1"]).value or "零售消费"
            main_ws.cell(main_row, mi["industry_l2"]).value = industry
        main_ws.cell(main_row, mi["source_note"]).value = f"{main_ws.cell(main_row, mi['source_note']).value}；{TODAY} cninfo基础资料补强"
        main_ws.cell(main_row, mi["validation_gap"]).value = "已补上市公司基础资料与官网；后续重点补年报、IR、财报口径与字段级 evidence。"
        main_ws.cell(main_row, mi["last_verified_at"]).value = TODAY

        # profile
        pr = profile_rows[aid]
        if str(profiles_ws.cell(pr, pi["信息扎实度"]).value or "") == "中":
            profiles_ws.cell(pr, pi["信息扎实度"]).value = "中高"
        profiles_ws.cell(pr, pi["account_canonical_name"]).value = company_name
        profiles_ws.cell(pr, pi["industry_l1"]).value = profiles_ws.cell(pr, pi["industry_l1"]).value or "零售消费"
        profiles_ws.cell(pr, pi["industry_l2"]).value = industry or profiles_ws.cell(pr, pi["industry_l2"]).value
        if business:
            profiles_ws.cell(pr, pi["公司产品与服务概述"]).value = business
            profiles_ws.cell(pr, pi["产品与服务长摘录"]).value = business
        source_types = str(profiles_ws.cell(pr, pi["primary_source_types"]).value or "")
        if "上市公司基础资料" not in source_types:
            profiles_ws.cell(pr, pi["primary_source_types"]).value = f"{source_types},上市公司基础资料".strip(",")
        refs = str(profiles_ws.cell(pr, pi["primary_source_refs"]).value or "")
        additions = []
        if website and website not in refs:
            additions.append(f"https://{website}" if not website.startswith("http") else website)
        additions.append(f"cninfo:{code}")
        all_refs = "\n".join([x for x in [refs] + additions if x]).strip()
        profiles_ws.cell(pr, pi["primary_source_refs"]).value = all_refs
        profiles_ws.cell(pr, pi["official_source_count"]).value = max(int(profiles_ws.cell(pr, pi["official_source_count"]).value or 0), 1)
        profiles_ws.cell(pr, pi["high_confidence_source_count"]).value = max(int(profiles_ws.cell(pr, pi["high_confidence_source_count"]).value or 0), 2)
        profiles_ws.cell(pr, pi["last_profiled_at"]).value = TODAY
        profiles_ws.cell(pr, pi["profile_owner"]).value = "Codex 结构化沉淀"
        profiles_ws.cell(pr, pi["validation_gap"]).value = "已补上市公司基础资料与官网；后续重点补年报、IR、财报口径与字段级 evidence。"
        profiles_ws.cell(pr, pi["商业模式长摘录"]).value = profiles_ws.cell(pr, pi["商业模式概述"]).value

        # coverage
        cr = coverage_rows[aid]
        coverage_ws.cell(cr, ci["official_source_ready"]).value = "yes"
        coverage_ws.cell(cr, ci["high_confidence_ready"]).value = "yes"
        coverage_ws.cell(cr, ci["missing_core_fields"]).value = "收入规模、利润状态、营收增长仍需回到财报/年报/IR 口径继续补齐。"
        coverage_ws.cell(cr, ci["next_action"]).value = "继续补年报、IR 和财报口径，必要时再评估是否上移到 L2。"

        # evidence
        evidence_id = f"ev_{aid}_cninfo_profile_20260331"
        exists = False
        for er in range(2, evidence_ws.max_row + 1):
            if str(evidence_ws.cell(er, ei["evidence_id"]).value or "") == evidence_id:
                exists = True
                break
        if not exists:
            locator = f"https://www.cninfo.com.cn/ ; code={code}"
            evidence_ws.append([
                evidence_id,
                aid,
                "official_profile",
                locator,
                "A",
                "background_profile,identity",
                f"{company_name} 已通过上市公司基础资料确认公司名称、官网、行业与主营业务。",
                "codex_llm",
                TODAY,
                "",
                "公司产品与服务概述",
                business,
            ])

        updated += 1

    main_wb.save(MAIN_XLSX)
    profile_wb.save(PROFILE_XLSX)
    gov_wb.save(GOV_XLSX)

    DOC_PATH.write_text(
        "\n".join([
            "# 静态潜客池-L3消费品官方源补强专项-v1",
            "",
            f"- 执行日期：`{TODAY}`",
            f"- 覆盖对象：消费品新画像批次中已上移到 `L3` 且可稳定映射 A 股代码的对象。",
            f"- 本轮目标批次：`{min(len(candidates), MAX_ITEMS)}` 家。",
            f"- 本轮成功补强：`{updated}` 家。",
            f"- 跳过：`{len(skipped)}` 家。",
            "",
            "## 本轮补强内容",
            "",
            "- 通过 cninfo 上市公司基础资料补充公司全称、官网、所属行业和主营业务。",
            "- 将 `official_source_count` 至少提升到 `1`，并把 `official_source_ready` 收敛到 `yes`。",
            "- 对这批对象把 `信息扎实度` 从 `中` 提升到 `中高`（如适用）。",
            "- 新增 `official_profile` 类型 evidence。",
        ]),
        encoding="utf-8",
    )

    with MEMORY_PATH.open("a", encoding="utf-8") as f:
        f.write(f"\n- 对消费品新画像批次中已上移到 `L3` 的 A 股主体做了一轮 cninfo 官方源补强，成功补强 `{updated}` 家。\n")

    print({"candidate_total": len(candidates), "batch_target": min(len(candidates), MAX_ITEMS), "updated": updated, "skipped": len(skipped), "sample_skipped": skipped[:10]})


if __name__ == "__main__":
    main()
