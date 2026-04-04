from __future__ import annotations

import importlib.util
import json
import time
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


ROOT = Path("/Users/clairaipartner")
WORKSPACE = ROOT / ".openclaw/workspace-main/bussiness-master"
VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"

PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
DELIVERIES_DIR = WORKSPACE / "deliveries"

FACT_PLACEHOLDER = "待补官网/年报/IR口径"
EVENT_PLACEHOLDER = "待补更强官方披露"
MARKET_PLACEHOLDER = "待补公司级市场参考"

ENRICH_FACTS = WORKSPACE / "scripts/enrich_l1_company_facts_20260402.py"
ENRICH_MAPPING = WORKSPACE / "scripts/enrich_icp_mapping_fields_20260402.py"

INVALID_TALKTRACKS = {
    "ka_narrative_group_dashboard_mfg_v1",
    "待补话术资产",
}

PERSONA_TALKTRACK_FALLBACK = {
    "mfg_multi_factory_group": "ka_talktrack_group_coordination_mfg_v1",
    "mfg_rnd_sales_complex": "ka_talktrack_group_coordination_mfg_v1",
    "cbec_multi_platform_brand": "ka_talktrack_profit_governance_cbec_v1",
    "cbec_supply_chain_complex": "ka_talktrack_profit_governance_cbec_v1",
    "cbec_platform_operator": "ka_talktrack_profit_governance_cbec_v1",
    "retail_multi_store": "ka_talktrack_hq_visibility_retail_v1",
    "retail_multi_store_chain": "ka_talktrack_hq_visibility_retail_v1",
    "retail_chain_fnb": "ka_talktrack_frontline_loop_retail_v1",
    "fnb_chain_standardized": "ka_talktrack_frontline_loop_retail_v1",
    "fnb_chain_beverage_coffee": "ka_talktrack_frontline_loop_retail_v1",
    "retail_high_sku_brand": "ka_talktrack_inventory_replenishment_retail_v1",
}


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


facts_mod = load_module(ENRICH_FACTS, "enrich_facts")
mapping_mod = load_module(ENRICH_MAPPING, "enrich_mapping")


def sanitize_filename_suffix(num: int) -> str:
    return f"{num:02d}"


def clean(text) -> str:
    return str(text or "").strip()


def safe_load_workbook(path: Path, **kwargs):
    last_error = None
    for attempt in range(6):
        try:
            return load_workbook(path, **kwargs)
        except (EOFError, BadZipFile) as exc:
            last_error = exc
            time.sleep(0.5 * (attempt + 1))
    raise last_error


def is_placeholder(value: str) -> bool:
    value = clean(value)
    return value in {
        "",
        FACT_PLACEHOLDER,
        EVENT_PLACEHOLDER,
        MARKET_PLACEHOLDER,
        "待补字段级 evidence",
        "待补公开财报口径",
    }


def choose_talktrack(current: str, persona: str, records: list[dict]) -> str:
    current = clean(current)
    if current and current not in INVALID_TALKTRACKS:
        return current
    chosen = mapping_mod.choose_assets(records, ("sales_narrative",), limit=1)
    if chosen:
        return chosen[0]
    return PERSONA_TALKTRACK_FALLBACK.get(persona, "")


def choose_knowledge_assets(current: str, records: list[dict]) -> str:
    current = clean(current)
    if current:
        return current
    chosen = mapping_mod.choose_assets(
        records,
        ("customer_case", "industry_insight", "scenario_pack", "solution_playbook"),
        limit=3,
    )
    return ",".join(chosen)


def choose_mgmt_tags(current: str, track: str, persona: str) -> str:
    current = clean(current)
    if current:
        return current
    return mapping_mod.infer_mgmt_tags(track, persona)


def quote_id_for(main_row: dict, profile_row: dict) -> str | None:
    refs = "\n".join(
        x for x in [
            clean(main_row.get("source_urls")),
            clean(profile_row.get("primary_source_refs")),
        ] if x
    )
    q = facts_mod.quote_id_from_refs(refs)
    if q:
        return q
    search = facts_mod.eastmoney_search(profile_row["account_canonical_name"])
    if not search:
        return None
    code = clean(search.get("Code"))
    if not code:
        return None
    market = "1" if clean(search.get("SecurityTypeName")).upper() == "SH" else None
    if market is None:
        sec_code = clean(search.get("SecurityType"))
        market = "1" if sec_code in {"SH", "1"} else "0"
    return f"{market}.{code}"


def company_profile(main_row: dict, profile_row: dict) -> tuple[dict | None, list[str], list[str]]:
    refs = []
    types = []
    quote_id = quote_id_for(main_row, profile_row)
    survey = None
    if quote_id:
        try:
            survey = facts_mod.company_survey(quote_id)
        except Exception:
            survey = None
        if survey:
            market, code = quote_id.split(".")
            prefix = "SH" if market == "1" else "SZ"
            refs.append(f"https://emweb.securities.eastmoney.com/PC_HSF10/CompanySurvey/Index?type=web&code={prefix}{code}")
            types.append("东方财富F10公司概况")
    name = profile_row["account_canonical_name"]
    refs.append(f"https://www.cninfo.com.cn/new/fulltextSearch?notautosubmit=&keyWord={name}")
    types.append("上市公司公开披露")
    return survey, types, refs


def build_customer_text(profile_survey: dict | None) -> str:
    if not profile_survey:
        return FACT_PLACEHOLDER
    text = facts_mod.summarize_customer(profile_survey)
    if not text:
        return FACT_PLACEHOLDER
    generic_bad = {
        "工业制造客户、设备厂商、渠道伙伴及项目型客户。",
        "终端消费者、经销/零售渠道及线上线下平台。",
        "工业企业、行业客户及项目型客户。",
    }
    if text in generic_bad:
        return FACT_PLACEHOLDER
    return text


def build_competitor_text(persona_meta: dict | None, current: str) -> tuple[str, str]:
    current = clean(current)
    similar = current if current and current != "待补公司级市场参考" else ""
    if not persona_meta:
        return similar or "待补公司级市场参考", "待补公司级市场参考"
    refs = [x.strip() for x in clean(persona_meta.get("reference_customers")).split(",") if x.strip()]
    if not similar and refs:
        similar = "、".join(refs[:3])
    if refs:
        competitor = f"{'、'.join(refs[:2])}及相关企业。"
    else:
        competitor = "待补公司级市场参考"
    return similar or "待补公司级市场参考", competitor


def build_repair_note(name: str, product: str) -> str:
    core = product.replace("主营", "").replace("，是", "与").replace("。", "")
    return f"把{name.replace('股份有限公司','').replace('集团股份有限公司','').replace('集团股份公司','').replace('有限公司','')}收回到{core}的公司级事实。"


def load_rows():
    profile_wb = safe_load_workbook(PROFILE_XLSX, read_only=True, data_only=True)
    profile_ws = profile_wb["account_profiles"]
    ph = [c.value for c in next(profile_ws.iter_rows(min_row=1, max_row=1))]
    pi = {h: i for i, h in enumerate(ph)}
    profiles = {}
    for row in profile_ws.iter_rows(min_row=2, values_only=True):
        name = clean(row[pi["account_canonical_name"]])
        if not name:
            continue
        profiles[name] = {h: row[i] for h, i in pi.items()}

    main_wb = safe_load_workbook(MAIN_XLSX, read_only=True, data_only=True)
    main_ws = main_wb["accounts_main"]
    mh = [c.value for c in next(main_ws.iter_rows(min_row=1, max_row=1))]
    mi = {h: i for i, h in enumerate(mh)}
    main_by_id = {}
    for row in main_ws.iter_rows(min_row=2, values_only=True):
        aid = clean(row[mi["account_id"]])
        if aid:
            main_by_id[aid] = {h: row[i] for h, i in mi.items()}
    return profiles, main_by_id


def candidate_score(profile: dict, main: dict) -> int | None:
    status = clean(profile.get("archive_repair_status"))
    if status and status != "未修复":
        return None
    if clean(profile.get("静态潜客记录成熟度")) != "L1":
        return None

    sturdiness = clean(profile.get("信息扎实度"))
    sturdiness_score = {
        "高": 4,
        "中高": 3,
        "中": 2,
        "中低": 0,
    }.get(sturdiness, 0)
    if sturdiness_score <= 0:
        return None

    product = clean(profile.get("公司产品与服务概述"))
    business = clean(profile.get("商业模式概述"))
    product_ready = not is_placeholder(product)
    business_ready = not is_placeholder(business)
    refs = "\n".join(
        x for x in [
            clean(main.get("source_urls")),
            clean(profile.get("primary_source_refs")),
        ] if x
    )
    has_survey_ref = "emweb.securities.eastmoney.com/PC_HSF10/CompanySurvey" in refs
    if not product_ready and not has_survey_ref:
        return None
    if not business_ready:
        return None
    if not product_ready and not business_ready:
        return None

    score = sturdiness_score * 10
    score += 8 if product_ready else 0
    score += 8 if business_ready else 0

    official = int(profile.get("official_source_count") or 0)
    high_conf = int(profile.get("high_confidence_source_count") or 0)
    score += min(official, 3) * 2
    score += min(high_conf, 3) * 2

    if has_survey_ref:
        score += 6
    if "quote.eastmoney.com/unify" in refs:
        score += 3
    if "cninfo.com.cn" in refs:
        score += 2

    return score


def next_l1_accounts(limit: int) -> list[dict]:
    profiles, main_by_id = load_rows()
    rows = []
    for name, row in profiles.items():
        aid = clean(row["account_id"])
        main = main_by_id.get(aid, {})
        score = candidate_score(row, main)
        if score is None:
            continue
        rows.append({
            "account_id": aid,
            "account_name": name,
            "primary_track": clean(row.get("primary_track")),
            "profile": row,
            "main": main,
            "score": score,
        })
    rows.sort(key=lambda x: (-x["score"], x["account_name"]))
    return rows[:limit]


def generate_batch(batch_num: int, limit: int = 12) -> Path:
    persona_meta = mapping_mod.load_persona_meta()
    assets_by_persona = mapping_mod.load_assets()
    candidates = next_l1_accounts(limit)
    batch_id = f"archive_repair_2026_04_batch_{sanitize_filename_suffix(batch_num)}"
    payload = {
        "batch_id": batch_id,
        "batch_title": f"客户档案修复第{batch_num}批",
        "account_names": [x["account_name"] for x in candidates],
        "accounts": [],
    }

    for item in candidates:
        profile = item["profile"]
        main = item["main"]
        persona = clean(profile.get("persona_tag"))
        track = clean(profile.get("primary_track"))
        persona_info, records = mapping_mod.persona_with_fallback(persona_meta, assets_by_persona, persona)
        survey, source_types_add, source_refs_add = company_profile(main, profile)

        product = (facts_mod.summarize_product(survey) if survey else None) or clean(profile.get("公司产品与服务概述")) or FACT_PLACEHOLDER
        if facts_mod.is_generic_product_text(product):
            product = FACT_PLACEHOLDER
        business = (facts_mod.summarize_business(survey) if survey else None) or clean(profile.get("商业模式概述")) or FACT_PLACEHOLDER
        if facts_mod.is_generic_business_text(business):
            business = FACT_PLACEHOLDER
        customer = clean(profile.get("核心客户客群"))
        if not customer or customer == FACT_PLACEHOLDER:
            customer = build_customer_text(survey)
        similar, competitor = build_competitor_text(persona_info, clean(profile.get("相似客户线索")))

        knowledge_refs = choose_knowledge_assets(clean(profile.get("knowledge_asset_refs")), records)
        talk_refs = choose_talktrack(clean(profile.get("talk_track_refs")), persona, records)
        mgmt = choose_mgmt_tags(clean(profile.get("management_persona_tags")), track, persona)

        account = {
            "account_canonical_name": item["account_name"],
            "fields": {
                "公司产品与服务概述": product,
                "商业模式概述": business,
                "核心客户客群": customer,
                "主要竞品概述": competitor,
                "相似客户线索": similar,
            },
            "matching": {},
            "excerpts": {
                "产品与服务长摘录": f"公开披露可支撑其{product}" if product != FACT_PLACEHOLDER else FACT_PLACEHOLDER,
                "商业模式长摘录": f"公开资料可支撑其{business}" if business != FACT_PLACEHOLDER else FACT_PLACEHOLDER,
            },
            "source_types_add": source_types_add,
            "source_refs_add": source_refs_add,
            "repair_note": build_repair_note(item["account_name"], product if product != FACT_PLACEHOLDER else "公司级事实"),
        }
        if mgmt:
            account["matching"]["management_persona_tags"] = mgmt
        if knowledge_refs:
            account["matching"]["knowledge_asset_refs"] = knowledge_refs
        if talk_refs:
            account["matching"]["talk_track_refs"] = talk_refs
        payload["accounts"].append(account)

    out = DELIVERIES_DIR / f"customer_archive_repair_batch_2026_04_03_{sanitize_filename_suffix(batch_num)}.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        raise SystemExit("Usage: generate_customer_archive_repair_batch_20260403.py <batch_num> [limit]")
    batch_num = int(sys.argv[1])
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 12
    path = generate_batch(batch_num, limit)
    print(path)
