from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


REGISTRY_XLSX = Path("/Users/clairaipartner/Documents/Obsidian-Codex/潜客池/知识资产注册表.xlsx")
MANIFEST_JSON = Path("/Users/clairaipartner/Codex/static-pool-raw-materials/_manifest/copy_manifest.json")


def load_mapping() -> dict[str, str]:
    data = json.loads(MANIFEST_JSON.read_text(encoding="utf-8"))
    return {item["source_path"]: item["destination_path"] for item in data["copied_items"]}


def backup_registry() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = REGISTRY_XLSX.with_name(f"{REGISTRY_XLSX.stem}.backup_{timestamp}{REGISTRY_XLSX.suffix}")
    shutil.copy2(REGISTRY_XLSX, backup_path)
    return backup_path


def update_sheet_paths(ws, source_column: str, mapping: dict[str, str]) -> int:
    headers = [cell.value for cell in ws[1]]
    if source_column not in headers:
        return 0
    col_idx = headers.index(source_column) + 1
    updated = 0
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, col_idx)
        value = str(cell.value or "").strip()
        if value in mapping:
            cell.value = mapping[value]
            updated += 1
    return updated


def main() -> int:
    mapping = load_mapping()
    backup_path = backup_registry()
    wb = load_workbook(REGISTRY_XLSX)

    updated = {
        "raw_material_inventory": update_sheet_paths(wb["raw_material_inventory"], "material_path_or_url", mapping),
        "learning_queue": update_sheet_paths(wb["learning_queue"], "material_path_or_url", mapping),
        "knowledge_assets": update_sheet_paths(wb["knowledge_assets"], "source_path_or_url", mapping),
    }

    wb.save(REGISTRY_XLSX)
    print(
        json.dumps(
            {
                "registry_path": str(REGISTRY_XLSX),
                "backup_path": str(backup_path),
                "updated": updated,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
