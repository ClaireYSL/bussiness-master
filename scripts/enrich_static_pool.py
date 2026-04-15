from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
ROOT = Path.home()
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))

from shared.static_pool import (
    build_enrich_results,
    build_enrich_summary_payload,
    dataclass_to_dict,
    render_enrich_review_markdown,
)

VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
MAIN_SHARED_XLSX = VAULT / "内部运营-静态潜客池-共享版.xlsx"
GOV_XLSX = VAULT / "治理与证据.xlsx"
DEFAULT_RECTIFICATION = WORKSPACE / "deliveries/phase1_rectification_package_v1.json"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Static pool enrich entry for fact strengthening and knowledge attachment.")
    parser.add_argument("--config-file", help="Optional JSON batch config file.")
    parser.add_argument("--account-id", action="append", default=[], help="Single account_id to enrich. Can be repeated.")
    parser.add_argument("--batch-file", help="Optional JSON or newline-delimited file containing account_ids.")
    parser.add_argument("--track", help="Optional track filter, for example 零售消费.")
    parser.add_argument("--from-level", help="Optional current level filter, for example L4.")
    parser.add_argument("--rectification-file", help="Optional rectification package JSON.")
    parser.add_argument("--output-file", help="Where to write the enrich result package.")
    parser.add_argument("--report-only", action="store_true", help="Only build enrich results, do not write back.")
    parser.add_argument("--write-back", action="store_true", help="Write enrich results back for selected accounts.")
    return parser


def _clean(value: object) -> str:
    return str(value or "").strip()


def load_account_ids(path: str | None) -> list[str]:
    if not path:
        return []
    payload = Path(path).read_text(encoding="utf-8").strip()
    if not payload:
        return []
    if payload.startswith("["):
        values = json.loads(payload)
        return [str(item).strip() for item in values if str(item).strip()]
    return [line.strip() for line in payload.splitlines() if line.strip()]


def resolve_output_file(path: str | None, batch_hint: str) -> Path:
    if path:
        return Path(path)
    base = Path(tempfile.gettempdir()) / "codex-static-pool-runs"
    base.mkdir(parents=True, exist_ok=True)
    return base / f"{batch_hint}.json"


def optional_output_path(path: str | None) -> Path | None:
    if not path:
        return None
    return Path(path)


def backup_once(path: Path) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = path.with_name(f"{path.stem}.enrich_backup_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    return str(backup_path)


def ensure_column(ws, header: str) -> int:
    headers = [cell.value for cell in ws[1]]
    for idx, value in enumerate(headers, start=1):
        if value == header:
            return idx
    idx = len(headers) + 1
    ws.cell(1, idx).value = header
    return idx


def build_row_index(ws, key_header: str) -> tuple[dict[str, int], dict[str, int]]:
    headers = [cell.value for cell in ws[1]]
    header_index = {str(value): idx + 1 for idx, value in enumerate(headers) if value}
    if key_header not in header_index:
        raise KeyError(f"missing key header: {key_header}")
    row_index: dict[str, int] = {}
    key_col = header_index[key_header]
    for row in range(2, ws.max_row + 1):
        value = _clean(ws.cell(row, key_col).value)
        if value:
            row_index[value] = row
    return header_index, row_index


def set_if_header(ws, header_index: dict[str, int], row: int, header: str, value: object) -> None:
    if header in header_index:
        ws.cell(row, header_index[header]).value = value


def ensure_review_queue_item(ws, header_index: dict[str, int], account_id: str, queue_type: str, note: str) -> bool:
    account_col = header_index["account_id"]
    queue_col = header_index["queue_type"]
    status_col = header_index["status"]
    open_status = {"", "open", "in_progress"}
    for row in range(2, ws.max_row + 1):
        if _clean(ws.cell(row, account_col).value) != account_id:
            continue
        if _clean(ws.cell(row, queue_col).value) != queue_type:
            continue
        if _clean(ws.cell(row, status_col).value) not in open_status:
            continue
        if "note" in header_index:
            ws.cell(row, header_index["note"]).value = note
        return False

    target_row = ws.max_row + 1
    queue_item_id = f"enrich_{account_id}_{queue_type}"
    set_if_header(ws, header_index, target_row, "queue_item_id", queue_item_id)
    set_if_header(ws, header_index, target_row, "queue_type", queue_type)
    set_if_header(ws, header_index, target_row, "account_id", account_id)
    set_if_header(ws, header_index, target_row, "priority", "P1" if queue_type == "boundary_review" else "P2")
    set_if_header(ws, header_index, target_row, "status", "open")
    set_if_header(ws, header_index, target_row, "owner", "codex")
    set_if_header(ws, header_index, target_row, "note", note)
    set_if_header(ws, header_index, target_row, "created_at", datetime.now().strftime("%Y-%m-%d"))
    return True


def ensure_evidence_item(ws, header_index: dict[str, int], result: dict[str, object]) -> bool:
    account_id = _clean(result["account_id"])
    marker = "execution_rectification"
    account_col = header_index["account_id"]
    field_name_col = header_index["field_name"]
    locator_col = header_index["source_locator"]
    for row in range(2, ws.max_row + 1):
        if _clean(ws.cell(row, account_col).value) != account_id:
            continue
        if _clean(ws.cell(row, field_name_col).value) != marker:
            continue
        ws.cell(row, locator_col).value = "phase1_rectification_package_v1"
        ws.cell(row, header_index["summary"]).value = _clean(result["summary"])
        ws.cell(row, header_index["related_asset_ids"]).value = ",".join(result.get("knowledge_asset_refs") or [])
        ws.cell(row, header_index["field_value"]).value = json.dumps(
            {
                "candidate_type": result.get("candidate_type"),
                "review_status": result.get("review_status"),
                "persona_tag": result.get("persona_tag"),
            },
            ensure_ascii=False,
        )
        return False

    target_row = ws.max_row + 1
    set_if_header(ws, header_index, target_row, "evidence_id", f"enrich_{account_id}")
    set_if_header(ws, header_index, target_row, "account_id", account_id)
    set_if_header(ws, header_index, target_row, "evidence_type", "governance_enrich_rectification")
    set_if_header(ws, header_index, target_row, "source_locator", "phase1_rectification_package_v1")
    set_if_header(ws, header_index, target_row, "evidence_strength", "A")
    set_if_header(ws, header_index, target_row, "supports_dimension", "execution_layer_rectification")
    set_if_header(ws, header_index, target_row, "summary", _clean(result["summary"]))
    set_if_header(ws, header_index, target_row, "checked_by", "codex")
    set_if_header(ws, header_index, target_row, "checked_at", datetime.now().strftime("%Y-%m-%d"))
    set_if_header(ws, header_index, target_row, "related_asset_ids", ",".join(result.get("knowledge_asset_refs") or []))
    set_if_header(ws, header_index, target_row, "field_name", "execution_rectification")
    set_if_header(
        ws,
        header_index,
        target_row,
        "field_value",
        json.dumps(
            {
                "candidate_type": result.get("candidate_type"),
                "review_status": result.get("review_status"),
                "persona_tag": result.get("persona_tag"),
            },
            ensure_ascii=False,
        ),
    )
    return True


def write_back_results(results: list[dict[str, object]]) -> dict[str, object]:
    backups = {
        "profile": backup_once(PROFILE_XLSX),
        "main": backup_once(MAIN_XLSX),
        "main_shared": backup_once(MAIN_SHARED_XLSX),
        "governance": backup_once(GOV_XLSX),
    }

    profile_wb = load_workbook(PROFILE_XLSX)
    profile_ws = profile_wb["account_profiles"]
    profile_secondary_col = ensure_column(profile_ws, "secondary_persona_tags")
    profile_headers, profile_rows = build_row_index(profile_ws, "account_id")
    profile_headers["secondary_persona_tags"] = profile_secondary_col

    main_wb = load_workbook(MAIN_XLSX)
    main_ws = main_wb["accounts_main"]
    main_headers, main_rows = build_row_index(main_ws, "account_canonical_name")

    shared_wb = load_workbook(MAIN_SHARED_XLSX)
    shared_ws = shared_wb["全量主表"]
    shared_headers, shared_rows = build_row_index(shared_ws, "公司主体")

    gov_wb = load_workbook(GOV_XLSX)
    queue_ws = gov_wb["review_queue"]
    evidence_ws = gov_wb["evidence_log"]
    queue_headers, _queue_rows = build_row_index(queue_ws, "queue_item_id")
    evidence_headers, _evidence_rows = build_row_index(evidence_ws, "evidence_id")

    profile_updates = 0
    main_updates = 0
    shared_updates = 0
    queue_created = 0
    evidence_created = 0

    for result in results:
        account_id = _clean(result["account_id"])
        account_name = _clean(result["account_canonical_name"])
        if account_id in profile_rows:
            row = profile_rows[account_id]
            set_if_header(profile_ws, profile_headers, row, "primary_track", result.get("primary_track"))
            set_if_header(profile_ws, profile_headers, row, "persona_tag", result.get("persona_tag"))
            set_if_header(profile_ws, profile_headers, row, "secondary_persona_tags", ",".join(result.get("secondary_persona_tags") or []))
            set_if_header(profile_ws, profile_headers, row, "static_maturity_level", result.get("suggested_maturity"))
            set_if_header(profile_ws, profile_headers, row, "静态潜客记录成熟度", result.get("suggested_maturity"))
            set_if_header(profile_ws, profile_headers, row, "validation_gap", result.get("validation_gap"))
            set_if_header(profile_ws, profile_headers, row, "profile_status", result.get("review_status"))
            set_if_header(profile_ws, profile_headers, row, "knowledge_asset_refs", ",".join(result.get("knowledge_asset_refs") or []))
            set_if_header(profile_ws, profile_headers, row, "talk_track_refs", ",".join(result.get("talk_track_refs") or []))
            set_if_header(profile_ws, profile_headers, row, "last_profiled_at", datetime.now().strftime("%Y-%m-%d"))
            profile_updates += 1

        if account_name in main_rows:
            row = main_rows[account_name]
            set_if_header(main_ws, main_headers, row, "primary_track", result.get("primary_track"))
            set_if_header(main_ws, main_headers, row, "persona_tag", result.get("persona_tag"))
            set_if_header(main_ws, main_headers, row, "静态潜客记录成熟度", result.get("suggested_maturity"))
            set_if_header(main_ws, main_headers, row, "knowledge_asset_refs", ",".join(result.get("knowledge_asset_refs") or []))
            set_if_header(main_ws, main_headers, row, "talk_track_refs", ",".join(result.get("talk_track_refs") or []))
            set_if_header(main_ws, main_headers, row, "validation_gap", result.get("validation_gap"))
            rewrite = result.get("rewrite_suggestion") or {}
            if isinstance(rewrite, dict):
                if rewrite.get("admission_reason_summary"):
                    set_if_header(main_ws, main_headers, row, "admission_reason_summary", rewrite["admission_reason_summary"])
                if rewrite.get("公司产品与服务概述"):
                    set_if_header(main_ws, main_headers, row, "公司产品与服务概述", rewrite["公司产品与服务概述"])
                if rewrite.get("商业模式概述"):
                    set_if_header(main_ws, main_headers, row, "商业模式概述", rewrite["商业模式概述"])
                if rewrite.get("核心客户客群"):
                    set_if_header(main_ws, main_headers, row, "核心客户客群", rewrite["核心客户客群"])
            main_updates += 1

        if account_name in shared_rows:
            row = shared_rows[account_name]
            set_if_header(shared_ws, shared_headers, row, "主线", result.get("primary_track"))
            set_if_header(shared_ws, shared_headers, row, "业务形态画像", result.get("persona_tag"))
            set_if_header(shared_ws, shared_headers, row, "静态潜客记录成熟度", result.get("suggested_maturity"))
            set_if_header(shared_ws, shared_headers, row, "主要知识资产引用", ",".join(result.get("knowledge_asset_refs") or []))
            set_if_header(shared_ws, shared_headers, row, "主要切入话术引用", ",".join(result.get("talk_track_refs") or []))
            set_if_header(shared_ws, shared_headers, row, "待验证项", result.get("validation_gap"))
            rewrite = result.get("rewrite_suggestion") or {}
            if isinstance(rewrite, dict):
                if rewrite.get("admission_reason_summary"):
                    set_if_header(shared_ws, shared_headers, row, "一话入池理由", rewrite["admission_reason_summary"])
                if rewrite.get("公司产品与服务概述"):
                    set_if_header(shared_ws, shared_headers, row, "公司产品与服务概述", rewrite["公司产品与服务概述"])
                if rewrite.get("商业模式概述"):
                    set_if_header(shared_ws, shared_headers, row, "商业模式概述", rewrite["商业模式概述"])
                if rewrite.get("核心客户客群"):
                    set_if_header(shared_ws, shared_headers, row, "核心客户客群", rewrite["核心客户客群"])
            shared_updates += 1

        note = f"{result.get('summary')} | required_queue_type={result.get('required_queue_type')}"
        if ensure_review_queue_item(queue_ws, queue_headers, account_id, _clean(result["required_queue_type"]), note):
            queue_created += 1
        if ensure_evidence_item(evidence_ws, evidence_headers, result):
            evidence_created += 1

    profile_wb.save(PROFILE_XLSX)
    main_wb.save(MAIN_XLSX)
    shared_wb.save(MAIN_SHARED_XLSX)
    gov_wb.save(GOV_XLSX)
    return {
        "backups": backups,
        "profile_updates": profile_updates,
        "main_updates": main_updates,
        "main_shared_updates": shared_updates,
        "queue_items_created": queue_created,
        "evidence_items_created": evidence_created,
    }


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    config: dict[str, object] = {}
    if args.config_file:
        config = json.loads(Path(args.config_file).read_text(encoding="utf-8"))
    config_output = config.get("output") or {}

    account_ids = list(
        dict.fromkeys(
            [
                *args.account_id,
                *load_account_ids(args.batch_file),
                *[str(item).strip() for item in config.get("account_ids") or [] if str(item).strip()],
            ]
        )
    )
    rectification_file = args.rectification_file or str(config.get("rectification_file") or str(DEFAULT_RECTIFICATION))
    rectification_path = Path(rectification_file) if rectification_file else None
    results = [dataclass_to_dict(item) for item in build_enrich_results(account_ids=account_ids or None, rectification_path=rectification_path)]
    track = args.track or str(config.get("track") or "")
    from_level = args.from_level or str(config.get("from_level") or "")
    output_file = args.output_file or str(config_output.get("enrich_file") or "")
    summary_file = str(config_output.get("summary_file") or config.get("summary_file") or "")
    review_file = str(config_output.get("review_file") or config.get("review_file") or "")
    output_path = resolve_output_file(output_file or None, str(config.get("batch_id") or "enrich_static_pool_run"))
    if track:
        results = [item for item in results if _clean(item.get("primary_track")) == _clean(track)]
    if from_level:
        results = [item for item in results if _clean(item.get("current_level")) == _clean(from_level)]
    payload = {
        "batch_id": str(config.get("batch_id") or output_path.stem),
        "goal": str(config.get("goal") or ""),
        "rectification_file": str(rectification_path) if rectification_path else "",
        "selection": {
            "account_ids": account_ids,
            "track": track,
            "from_level": from_level,
        },
        "summary": {
            "result_count": len(results),
            "ready_for_promote": sum(1 for item in results if item.get("enrich_ready_for_promote")),
            "observation": sum(1 for item in results if item.get("candidate_type") == "observation"),
        },
        "results": results,
    }
    if args.write_back and not args.report_only:
        payload["write_back"] = write_back_results(results)
    else:
        payload["write_back"] = {"enabled": False}

    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    summary_path = optional_output_path(summary_file)
    summary_payload = None
    if summary_path:
        summary_path.parent.mkdir(parents=True, exist_ok=True)
        summary_payload = build_enrich_summary_payload(
            config,
            payload,
            config_path=str(args.config_file or ""),
            enrich_result_path=str(output_path),
        )
        summary_path.write_text(json.dumps(summary_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    review_path = optional_output_path(review_file)
    if review_path:
        review_path.parent.mkdir(parents=True, exist_ok=True)
        if summary_payload is None:
            summary_payload = build_enrich_summary_payload(
                config,
                payload,
                config_path=str(args.config_file or ""),
                enrich_result_path=str(output_path),
            )
        review_path.write_text(render_enrich_review_markdown(summary_payload), encoding="utf-8")
    print(
        json.dumps(
            {
                "output_file": str(output_path),
                "summary": payload["summary"],
                "write_back": payload["write_back"],
                "summary_file": str(summary_path) if summary_path else "",
                "review_file": str(review_path) if review_path else "",
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
