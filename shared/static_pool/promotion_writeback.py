from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def append_semicolon_note(existing: object, addition: str) -> str:
    current = str(existing or "").strip().strip("；")
    extra = str(addition or "").strip().strip("；")
    if not extra:
        return current
    if not current:
        return extra
    if extra in current:
        return current
    return f"{current}；{extra}"


def prepend_gap_once(existing: object, prefix: str) -> str:
    current = str(existing or "").strip()
    marker = str(prefix or "").strip()
    if not marker:
        return current
    if marker in current:
        return current
    return f"{marker}{current}" if current else marker


def ensure_evidence_row(
    ws: Any,
    header_index: dict[str, int],
    evidence_id: str,
    values: dict[str, object],
) -> bool:
    evidence_col = header_index["evidence_id"]
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, evidence_col).value or "") == evidence_id:
            return False

    row_values = [values.get(header) for header in header_index]
    ws.append(row_values)
    return True


def apply_row_updates(
    ws: Any,
    header_index: dict[str, int],
    row_num: int,
    updates: dict[str, object],
) -> None:
    for key, value in updates.items():
        if key not in header_index:
            continue
        ws.cell(row_num, header_index[key]).value = value


def update_main_promotion_core(
    ws: Any,
    header_index: dict[str, int],
    row_num: int,
    *,
    maturity_level: str,
    source_note_suffix: str | None = None,
    validation_gap: str | None = None,
    extra_updates: dict[str, object] | None = None,
) -> None:
    updates = dict(extra_updates or {})
    updates["静态潜客记录成熟度"] = maturity_level
    if source_note_suffix and "source_note" in header_index:
        updates["source_note"] = append_semicolon_note(
            ws.cell(row_num, header_index["source_note"]).value,
            source_note_suffix,
        )
    if validation_gap is not None and "validation_gap" in header_index:
        updates["validation_gap"] = validation_gap
    apply_row_updates(ws, header_index, row_num, updates)


def update_profile_promotion_core(
    ws: Any,
    header_index: dict[str, int],
    row_num: int,
    *,
    maturity_level: str,
    profile_status: str | None = None,
    validation_gap: str | None = None,
    last_profiled_at: str | None = None,
    extra_updates: dict[str, object] | None = None,
) -> None:
    updates = dict(extra_updates or {})
    updates["静态潜客记录成熟度"] = maturity_level
    if "static_maturity_level" in header_index:
        updates["static_maturity_level"] = maturity_level
    if profile_status is not None and "profile_status" in header_index:
        updates["profile_status"] = profile_status
    if validation_gap is not None and "validation_gap" in header_index:
        updates["validation_gap"] = validation_gap
    if last_profiled_at is not None and "last_profiled_at" in header_index:
        updates["last_profiled_at"] = last_profiled_at
    apply_row_updates(ws, header_index, row_num, updates)


def resolve_open_queue_rows(
    ws: Any,
    header_index: dict[str, int],
    *,
    account_id: str,
    queue_type: str = "promotion_review",
    resolved_at: str | None = None,
    note_suffix: str | None = None,
) -> int:
    resolved = 0
    for row in range(2, ws.max_row + 1):
        row_account_id = str(ws.cell(row, header_index["account_id"]).value or "")
        row_queue_type = str(ws.cell(row, header_index["queue_type"]).value or "")
        row_status = str(ws.cell(row, header_index["status"]).value or "")
        if row_account_id != account_id or row_queue_type != queue_type or row_status != "open":
            continue
        ws.cell(row, header_index["status"]).value = "resolved"
        if resolved_at and "resolved_at" in header_index:
            ws.cell(row, header_index["resolved_at"]).value = resolved_at
        if note_suffix and "note" in header_index:
            note_cell = ws.cell(row, header_index["note"])
            note_cell.value = append_semicolon_note(note_cell.value, note_suffix)
        resolved += 1
    return resolved


def backup_once(path: Path, suffix: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = path.with_name(f"{path.stem}.{suffix}_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    return str(backup_path)


def build_header_index(ws: Any) -> dict[str, int]:
    headers = [cell.value for cell in ws[1]]
    return {str(value): idx + 1 for idx, value in enumerate(headers) if value}


def build_row_index(ws: Any, key_header: str) -> tuple[dict[str, int], dict[str, int]]:
    header_index = build_header_index(ws)
    if key_header not in header_index:
        raise KeyError(f"missing key header: {key_header}")
    row_index: dict[str, int] = {}
    key_col = header_index[key_header]
    for row in range(2, ws.max_row + 1):
        value = str(ws.cell(row, key_col).value or "").strip()
        if value:
            row_index[value] = row
    return header_index, row_index


def _coverage_updates_for_target(target_level: str) -> dict[str, object]:
    if target_level == "L2":
        return {
            "target_scope": "L1-L2",
            "official_source_ready": "yes",
            "high_confidence_ready": "yes",
            "profile_complete_status": "high_quality_ready",
            "next_action": "继续补更强官方披露和财报口径，按高质量样本标准持续压实。",
            "missing_core_fields": "收入规模、利润状态、营收增长仍需回到财报/年报/IR 口径继续补齐。",
        }
    if target_level == "L3":
        return {
            "target_scope": "L3",
            "official_source_ready": "yes",
            "high_confidence_ready": "yes",
            "profile_complete_status": "standard_ready",
            "next_action": "继续补官网、年报、IR 与字段级 evidence，再评估是否进入 L2。",
            "missing_core_fields": "收入规模、利润状态、营收增长仍需回到财报/年报/IR 口径继续补齐。",
        }
    return {}


def _profile_status_for_target(target_level: str, existing: object) -> str:
    if target_level in {"L1", "L2"}:
        return "high_quality_ready"
    if target_level == "L3":
        return "standard_ready"
    clean = str(existing or "").strip()
    return clean or "active"


def _validation_gap_for_target(target_level: str, existing: object) -> str:
    current = str(existing or "").strip()
    if target_level == "L2":
        return prepend_gap_once(current, "已进入L2；后续继续补官网、年报、IR和财报口径字段。")
    if target_level == "L3":
        return prepend_gap_once(current, "已进入L3；后续继续补官网、年报、IR和财报口径字段，再评估是否进入L2。")
    return current


def write_back_promotion_results(
    results: list[dict[str, object]],
    *,
    batch_id: str,
    profile_xlsx: Path,
    main_xlsx: Path,
    main_shared_xlsx: Path,
    gov_xlsx: Path,
) -> dict[str, object]:
    backups = {
        "profile": backup_once(profile_xlsx, "promote_backup"),
        "main": backup_once(main_xlsx, "promote_backup"),
        "main_shared": backup_once(main_shared_xlsx, "promote_backup"),
        "governance": backup_once(gov_xlsx, "promote_backup"),
    }

    profile_wb = load_workbook(profile_xlsx)
    profile_ws = profile_wb["account_profiles"]
    coverage_ws = profile_wb["profile_coverage"]
    profile_headers, profile_rows = build_row_index(profile_ws, "account_id")
    coverage_headers, coverage_rows = build_row_index(coverage_ws, "account_id")

    main_wb = load_workbook(main_xlsx)
    main_ws = main_wb["accounts_main"]
    main_headers, main_rows = build_row_index(main_ws, "account_canonical_name")

    shared_wb = load_workbook(main_shared_xlsx)
    shared_ws = shared_wb["全量主表"]
    shared_headers, shared_rows = build_row_index(shared_ws, "公司主体")

    gov_wb = load_workbook(gov_xlsx)
    queue_ws = gov_wb["review_queue"]
    evidence_ws = gov_wb["evidence_log"]
    queue_headers = build_header_index(queue_ws)
    evidence_headers = build_header_index(evidence_ws)

    promoted = 0
    skipped = 0
    queue_resolved = 0
    evidence_created = 0
    profile_updates = 0
    main_updates = 0
    shared_updates = 0
    coverage_updates = 0
    samples: list[dict[str, object]] = []
    today = datetime.now().strftime("%Y-%m-%d")

    for item in results:
        gate = item.get("promotion_gate") or {}
        decision = str(gate.get("decision") or "").strip()
        account_id = str(item.get("account_id") or "").strip()
        account_name = str(item.get("account_canonical_name") or "").strip()
        target_level = str(item.get("target_level") or "").strip()
        if decision != "allow" or not account_id or not account_name or not target_level:
            skipped += 1
            continue

        profile_row_num = profile_rows.get(account_id)
        main_row_num = main_rows.get(account_name)
        if profile_row_num is None or main_row_num is None:
            skipped += 1
            continue

        existing_profile_status = ""
        if "profile_status" in profile_headers:
            existing_profile_status = profile_ws.cell(profile_row_num, profile_headers["profile_status"]).value
        profile_status = _profile_status_for_target(target_level, existing_profile_status)
        profile_gap = _validation_gap_for_target(target_level, profile_ws.cell(profile_row_num, profile_headers["validation_gap"]).value)
        main_gap_header = "validation_gap" if "validation_gap" in main_headers else "待验证项"
        main_gap = _validation_gap_for_target(target_level, main_ws.cell(main_row_num, main_headers[main_gap_header]).value)

        update_profile_promotion_core(
            profile_ws,
            profile_headers,
            profile_row_num,
            maturity_level=target_level,
            profile_status=profile_status,
            validation_gap=profile_gap,
            last_profiled_at=today,
        )
        profile_updates += 1

        update_main_promotion_core(
            main_ws,
            main_headers,
            main_row_num,
            maturity_level=target_level,
            source_note_suffix=f"{today} {batch_id} 升层写回",
            validation_gap=main_gap,
        )
        main_updates += 1

        shared_row_num = shared_rows.get(account_name)
        if shared_row_num is not None:
            update_main_promotion_core(
                shared_ws,
                shared_headers,
                shared_row_num,
                maturity_level=target_level,
                validation_gap=main_gap,
            )
            shared_updates += 1

        coverage_row_num = coverage_rows.get(account_id)
        if coverage_row_num is not None:
            apply_row_updates(
                coverage_ws,
                coverage_headers,
                coverage_row_num,
                _coverage_updates_for_target(target_level),
            )
            coverage_updates += 1

        evidence_id = f"ev_{account_id}_{batch_id}_{target_level}"
        if ensure_evidence_row(
            evidence_ws,
            evidence_headers,
            evidence_id,
            {
                "evidence_id": evidence_id,
                "account_id": account_id,
                "evidence_type": "promotion_assessment",
                "source_locator": f"internal://promotion/{batch_id}",
                "evidence_strength": "A",
                "supports_dimension": "promotion_writeback",
                "summary": f"{account_name} 已在 {batch_id} 中完成升层写回，当前层级调整为 {target_level}。",
                "checked_by": "codex",
                "checked_at": today,
                "related_asset_ids": "",
                "field_name": "静态潜客记录成熟度",
                "field_value": target_level,
            },
        ):
            evidence_created += 1

        queue_resolved += resolve_open_queue_rows(
            queue_ws,
            queue_headers,
            account_id=account_id,
            queue_type="promotion_review",
            resolved_at=today,
            note_suffix=f"{today} 已完成升层写回，当前层级={target_level}。",
        )

        promoted += 1
        samples.append(
            {
                "account_id": account_id,
                "account_canonical_name": account_name,
                "from_level": item.get("from_level"),
                "target_level": target_level,
            }
        )

    profile_wb.save(profile_xlsx)
    main_wb.save(main_xlsx)
    shared_wb.save(main_shared_xlsx)
    gov_wb.save(gov_xlsx)

    return {
        "enabled": True,
        "backups": backups,
        "promoted": promoted,
        "skipped": skipped,
        "profile_updates": profile_updates,
        "main_updates": main_updates,
        "main_shared_updates": shared_updates,
        "coverage_updates": coverage_updates,
        "evidence_created": evidence_created,
        "promotion_review_resolved": queue_resolved,
        "samples": samples,
    }
