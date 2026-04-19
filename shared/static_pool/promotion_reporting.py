from __future__ import annotations

import json
from typing import Any

from openpyxl import load_workbook

from .pathing import resolve_static_pool_paths

POOL_PATHS = resolve_static_pool_paths()
PROFILE_XLSX = POOL_PATHS["profile"]
MAIN_XLSX = POOL_PATHS["main"]
GOV_XLSX = POOL_PATHS["governance"]


def _clean(value: object) -> str:
    return str(value or "").strip()


def load_json(path: str | Path) -> dict[str, object]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def load_sheet_rows(path: Path, sheet_name: str) -> list[dict[str, object]]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    return [{headers[i]: row[i] for i in range(len(headers))} for row in ws.iter_rows(min_row=2, values_only=True)]


def build_lookup(rows: list[dict[str, object]], key: str) -> dict[str, dict[str, object]]:
    result: dict[str, dict[str, object]] = {}
    for row in rows:
        value = _clean(row.get(key))
        if value:
            result[value] = row
    return result


def build_grouped_lookup(rows: list[dict[str, object]], key: str) -> dict[str, list[dict[str, object]]]:
    result: dict[str, list[dict[str, object]]] = {}
    for row in rows:
        value = _clean(row.get(key))
        if value:
            result.setdefault(value, []).append(row)
    return result


def build_promote_summary_payload(config: dict[str, object], promote: dict[str, object], *, config_path: str = "", promote_result_path: str = "") -> dict[str, object]:
    profile_rows = load_sheet_rows(PROFILE_XLSX, "account_profiles")
    main_rows = load_sheet_rows(MAIN_XLSX, "accounts_main")
    queue_rows = load_sheet_rows(GOV_XLSX, "review_queue")

    profile_by_id = build_lookup(profile_rows, "account_id")
    main_by_name = build_lookup(main_rows, "account_canonical_name")
    queue_by_account = build_grouped_lookup(queue_rows, "account_id")

    write_back = promote.get("write_back") or {}
    promoted_samples = {
        _clean(item.get("account_id")): item for item in write_back.get("samples") or [] if _clean(item.get("account_id"))
    }

    items: list[dict[str, object]] = []
    summary = {
        "account_count": 0,
        "allow": 0,
        "warn": 0,
        "block": 0,
        "writeback_promoted": int(write_back.get("promoted") or 0),
        "writeback_skipped": int(write_back.get("skipped") or 0),
        "target_level_reached": 0,
        "level_changed": 0,
        "queue_resolved": int(write_back.get("promotion_review_resolved") or 0),
        "evidence_created": int(write_back.get("evidence_created") or 0),
    }

    for result in promote.get("results") or []:
        gate = result.get("promotion_gate") or {}
        account_id = _clean(result.get("account_id"))
        account_name = _clean(result.get("account_canonical_name"))
        from_level = _clean(result.get("from_level"))
        target_level = _clean(result.get("target_level") or promote.get("target_level"))
        decision = _clean(gate.get("decision"))
        profile = profile_by_id.get(account_id, {})
        main = main_by_name.get(account_name, {})
        current_level = _clean(profile.get("静态潜客记录成熟度") or main.get("静态潜客记录成熟度"))
        level_changed = bool(current_level and from_level and current_level != from_level)
        target_level_reached = bool(target_level and current_level == target_level)
        open_queue_types = sorted(
            {
                _clean(row.get("queue_type"))
                for row in queue_by_account.get(account_id, [])
                if _clean(row.get("status")) in {"", "open", "in_progress"}
            }
        )
        remaining_gaps = [
            _clean(issue.get("message")) for issue in (gate.get("blocking_issues") or []) if _clean(issue.get("message"))
        ] + [
            _clean(issue.get("message")) for issue in (gate.get("warning_issues") or []) if _clean(issue.get("message"))
        ]
        items.append(
            {
                "account_id": account_id,
                "account_name": account_name,
                "track": _clean(result.get("primary_track")),
                "persona_tag": _clean(result.get("persona_tag")),
                "from_level": from_level,
                "target_level": target_level,
                "current_level": current_level,
                "decision": decision,
                "level_changed": level_changed,
                "target_level_reached": target_level_reached,
                "writeback_applied": account_id in promoted_samples,
                "profile_status": _clean(profile.get("profile_status")),
                "validation_gap": _clean(profile.get("validation_gap") or main.get("validation_gap")),
                "promote_summary": _clean(gate.get("summary")),
                "blocking_issues": gate.get("blocking_issues") or [],
                "warning_issues": gate.get("warning_issues") or [],
                "remaining_gaps": remaining_gaps,
                "open_queue_types": open_queue_types,
            }
        )
        summary["account_count"] += 1
        if decision in summary:
            summary[decision] += 1
        if level_changed:
            summary["level_changed"] += 1
        if target_level_reached:
            summary["target_level_reached"] += 1

    return {
        "batch_id": _clean(config.get("batch_id") or promote.get("batch_id")),
        "goal": _clean(config.get("goal") or promote.get("goal")),
        "from_level": _clean(config.get("from_level") or promote.get("from_level")),
        "target_level": _clean(config.get("target_level") or promote.get("target_level")),
        "selection": promote.get("selection") or {"account_ids": list(config.get("account_ids") or [])},
        "source_files": {
            "config_file": config_path,
            "promote_result_file": promote_result_path,
        },
        "summary": summary,
        "write_back": write_back,
        "items": items,
    }


def render_issue_lines(issues: list[dict[str, object]]) -> list[str]:
    lines: list[str] = []
    for issue in issues:
        message = _clean(issue.get("message"))
        code = _clean(issue.get("code"))
        if message:
            lines.append(f"- `{code}`：{message}" if code else f"- {message}")
    return lines


def render_promote_review_markdown(payload: dict[str, object]) -> str:
    summary = payload.get("summary") or {}
    items = payload.get("items") or []
    lines = [
        f"# {payload.get('batch_id')} promote复盘-v1",
        "",
        "## 批次概览",
        "",
        f"- 目标：`{_clean(payload.get('goal'))}`",
        f"- 当前层级：`{_clean(payload.get('from_level'))}`",
        f"- 目标层级：`{_clean(payload.get('target_level'))}`",
        f"- 样本数：`{summary.get('account_count', 0)}`",
        f"- 判定结果：`allow={summary.get('allow', 0)} / warn={summary.get('warn', 0)} / block={summary.get('block', 0)}`",
        f"- 写回结果：`promoted={summary.get('writeback_promoted', 0)} / skipped={summary.get('writeback_skipped', 0)}`",
        f"- 实际升层：`level_changed={summary.get('level_changed', 0)} / target_level_reached={summary.get('target_level_reached', 0)}`",
        f"- 队列与证据：`queue_resolved={summary.get('queue_resolved', 0)} / evidence_created={summary.get('evidence_created', 0)}`",
        "",
        "## 分对象复盘",
        "",
    ]
    for item in items:
        lines.extend(
            [
                f"### {_clean(item.get('account_name'))} / `{_clean(item.get('account_id'))}`",
                "",
                f"- 主线 / 画像：`{_clean(item.get('track'))} / {_clean(item.get('persona_tag'))}`",
                f"- 层级变化：`{_clean(item.get('from_level'))} -> {_clean(item.get('current_level'))}`，目标=`{_clean(item.get('target_level'))}`",
                f"- promote 判定：`{_clean(item.get('decision'))}`",
                f"- 写回状态：`writeback_applied={item.get('writeback_applied')} / target_level_reached={item.get('target_level_reached')}`",
                f"- 当前档案状态：`profile_status={_clean(item.get('profile_status'))}`",
                f"- 结论：{_clean(item.get('promote_summary'))}",
            ]
        )
        validation_gap = _clean(item.get("validation_gap"))
        if validation_gap:
            lines.append(f"- 待验证项：{validation_gap}")
        open_queue_types = item.get("open_queue_types") or []
        if open_queue_types:
            lines.append(f"- 当前仍开放队列：`{', '.join(open_queue_types)}`")
        blocking = render_issue_lines(item.get("blocking_issues") or [])
        warning = render_issue_lines(item.get("warning_issues") or [])
        remaining = item.get("remaining_gaps") or []
        if blocking:
            lines.extend(["- 阻塞项：", *blocking])
        if warning:
            lines.extend(["- 风险项：", *warning])
        if remaining:
            lines.extend(["- 剩余缺口：", *[f"- {gap}" for gap in remaining]])
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"
