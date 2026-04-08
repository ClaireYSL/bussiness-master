from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .constants import LEGACY_PERSONA_TAG_MAP, STANDARD_PERSONA_IDS
from .models import EnrichResult, KnowledgeMatch, ValidationIssue
from .promote_engine import attach_account_ids, load_main_rows_with_fallback, load_sheet_rows
from .validators import classify_l5_candidate, evaluate_minimum_fact_set, normalize_persona_tag, normalize_secondary_persona_tags

ROOT = Path("/Users/clairaipartner")
VAULT = ROOT / "Documents/Obsidian-Codex/潜客池"
MAIN_XLSX = VAULT / "静态潜客主表.xlsx"
MAIN_SHARED_XLSX = VAULT / "内部运营-静态潜客池-共享版.xlsx"
PROFILE_XLSX = VAULT / "潜客档案库.xlsx"
GOV_XLSX = VAULT / "治理与证据.xlsx"
TRACK_PERSONA_XLSX = VAULT / "主线与画像注册表.xlsx"
KNOWLEDGE_XLSX = VAULT / "知识资产注册表.xlsx"

TRACK_ID_TO_NAME = {
    "retail_consumer": "零售消费",
    "cross_border_ecommerce": "跨境电商",
    "advanced_manufacturing": "先进制造",
}
TRACK_NAME_TO_ID = {value: key for key, value in TRACK_ID_TO_NAME.items()}


def _clean(value: object) -> str:
    return str(value or "").strip()


def _split_multi(value: object) -> list[str]:
    text = _clean(value).replace("，", ",").replace("；", ",").replace(";", ",").replace("\n", ",")
    return [item.strip() for item in text.split(",") if item.strip()]


def _track_name(value: object) -> str:
    text = _clean(value)
    return TRACK_ID_TO_NAME.get(text, text)


def _track_id(value: object) -> str:
    text = _clean(value)
    return TRACK_NAME_TO_ID.get(text, text)


def load_active_personas() -> dict[str, dict[str, str]]:
    wb = load_workbook(TRACK_PERSONA_XLSX, read_only=True, data_only=True)
    ws = wb["personas"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {header: pos for pos, header in enumerate(headers)}
    active: dict[str, dict[str, str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        persona_id = _clean(row[idx["persona_id"]])
        status = _clean(row[idx["status"]])
        if not persona_id or status != "active":
            continue
        active[persona_id] = {
            "track_id": _clean(row[idx["track_id"]]),
            "track_name": _track_name(row[idx["track_id"]]),
            "persona_name": _clean(row[idx["persona_name"]]),
            "typical_jtbd": _clean(row[idx["typical_jtbd"]]),
            "reference_customers": _clean(row[idx["reference_customers"]]),
            "reference_cases": _clean(row[idx["reference_cases"]]),
            "reference_solutions": _clean(row[idx["reference_solutions"]]),
        }
    return active


def load_knowledge_assets() -> list[dict[str, Any]]:
    _headers, rows = load_sheet_rows(KNOWLEDGE_XLSX, "knowledge_assets")
    assets: list[dict[str, Any]] = []
    for row in rows:
        assets.append(
            {
                "asset_id": _clean(row.get("asset_id")),
                "asset_type": _clean(row.get("asset_type")),
                "title": _clean(row.get("title")),
                "track_ids": [_track_name(item) for item in _split_multi(row.get("track_ids"))],
                "persona_ids": [normalize_persona_tag(item) for item in _split_multi(row.get("persona_ids"))],
                "jtbd_tags": _split_multi(row.get("jtbd_tags")),
                "complexity_tags": _split_multi(row.get("complexity_tags")),
                "summary": _clean(row.get("summary")),
            }
        )
    return assets


def load_learning_queue() -> list[dict[str, Any]]:
    _headers, rows = load_sheet_rows(KNOWLEDGE_XLSX, "learning_queue")
    return rows


def load_state() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    main_rows = load_main_rows_with_fallback(MAIN_XLSX, "accounts_main", MAIN_SHARED_XLSX, "全量主表")
    _profile_headers, profile_rows = load_sheet_rows(PROFILE_XLSX, "account_profiles")
    _queue_headers, queue_rows = load_sheet_rows(GOV_XLSX, "review_queue")
    _evidence_headers, evidence_rows = load_sheet_rows(GOV_XLSX, "evidence_log")
    main_rows = attach_account_ids(main_rows, profile_rows)
    return main_rows, profile_rows, queue_rows, evidence_rows


def load_rectification_results(path: Path) -> dict[str, dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {str(item.get("account_id") or ""): item for item in payload.get("results", [])}


def _build_lookup(rows: list[dict[str, Any]], key: str) -> dict[str, dict[str, Any]]:
    mapping: dict[str, dict[str, Any]] = {}
    for row in rows:
        value = _clean(row.get(key))
        if value:
            mapping[value] = row
    return mapping


def _build_grouped_lookup(rows: list[dict[str, Any]], key: str) -> dict[str, list[dict[str, Any]]]:
    mapping: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        value = _clean(row.get(key))
        if not value:
            continue
        mapping.setdefault(value, []).append(row)
    return mapping


def _score_asset(
    asset: dict[str, Any],
    *,
    track_name: str,
    primary_persona: str,
    secondary_personas: list[str],
    current_refs: list[str],
    persona_meta: dict[str, dict[str, str]],
) -> tuple[int, list[str]]:
    score = 0
    matched_on: list[str] = []
    if asset["asset_id"] in current_refs:
        score += 4
        matched_on.append("current_refs")
    if track_name and track_name in asset["track_ids"]:
        score += 2
        matched_on.append("track")
    if primary_persona and primary_persona in asset["persona_ids"]:
        score += 6
        matched_on.append("primary_persona")
    secondary_hit = [tag for tag in secondary_personas if tag in asset["persona_ids"]]
    if secondary_hit:
        score += min(4, len(secondary_hit) * 2)
        matched_on.append("secondary_persona")

    jtbd_tokens: list[str] = []
    if primary_persona and persona_meta.get(primary_persona, {}).get("typical_jtbd"):
        jtbd_tokens.extend(_split_multi(persona_meta[primary_persona]["typical_jtbd"]))
    for tag in secondary_personas:
        if persona_meta.get(tag, {}).get("typical_jtbd"):
            jtbd_tokens.extend(_split_multi(persona_meta[tag]["typical_jtbd"]))
    if jtbd_tokens and any(token and token in asset["summary"] for token in jtbd_tokens):
        score += 1
        matched_on.append("jtbd")
    return score, matched_on


def match_related_knowledge_assets(
    *,
    track_name: str,
    primary_persona: str,
    secondary_personas: list[str],
    current_refs: list[str],
    assets: list[dict[str, Any]],
    persona_meta: dict[str, dict[str, str]],
) -> tuple[list[KnowledgeMatch], list[str], list[str]]:
    scored: list[KnowledgeMatch] = []
    for asset in assets:
        score, matched_on = _score_asset(
            asset,
            track_name=track_name,
            primary_persona=primary_persona,
            secondary_personas=secondary_personas,
            current_refs=current_refs,
            persona_meta=persona_meta,
        )
        if score <= 0:
            continue
        reason = " + ".join(matched_on) if matched_on else "track"
        scored.append(
            KnowledgeMatch(
                asset_id=asset["asset_id"],
                asset_type=asset["asset_type"],
                title=asset["title"],
                match_reason=f"{track_name} / {primary_persona or '待定主画像'} 命中 {reason}",
                match_strength=score,
                matched_on=matched_on,
                summary=asset["summary"],
            )
        )
    scored.sort(key=lambda item: (-item.match_strength, item.asset_id))
    strong_assets: list[str] = []
    strong_talks: list[str] = []
    knowledge_matches: list[KnowledgeMatch] = []
    for item in scored:
        if item.asset_type == "sales_narrative":
            if item.asset_id not in strong_talks and len(strong_talks) < 3:
                strong_talks.append(item.asset_id)
                knowledge_matches.append(item)
        else:
            if item.asset_id not in strong_assets and len(strong_assets) < 5:
                strong_assets.append(item.asset_id)
                knowledge_matches.append(item)
        if len(strong_assets) >= 5 and len(strong_talks) >= 3:
            break
    return knowledge_matches, strong_assets, strong_talks


def find_matching_learning_queue(
    *,
    track_name: str,
    persona_tags: list[str],
    queue_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    matched: list[dict[str, Any]] = []
    track_id = _track_id(track_name)
    wanted = {tag for tag in persona_tags if tag}
    seen: set[str] = set()
    for row in queue_rows:
        if _clean(row.get("extraction_status")) != "queued":
            continue
        queue_id = _clean(row.get("queue_id"))
        trigger_track = _clean(row.get("trigger_track_id"))
        trigger_personas = {
            normalize_persona_tag(item)
            for item in _split_multi(row.get("trigger_persona_id"))
            if item and not item.startswith("mgmt_")
        }
        if (track_id and trigger_track == track_id) or (wanted & trigger_personas):
            if queue_id and queue_id not in seen:
                seen.add(queue_id)
                matched.append(
                    {
                        "queue_id": queue_id,
                        "material_title": _clean(row.get("material_title")),
                        "priority": _clean(row.get("priority_level") or row.get("priority")),
                        "trigger_persona_id": _clean(row.get("trigger_persona_id")),
                    }
                )
    return matched


def _determine_primary_persona(
    *,
    current_persona: str,
    rectification: dict[str, Any] | None,
    track_name: str,
    active_personas: dict[str, dict[str, str]],
) -> tuple[str, list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    rect_persona = _clean((rectification or {}).get("rectification", {}).get("suggested_primary_persona"))
    rect_candidate_type = _clean((rectification or {}).get("rectification", {}).get("final_candidate_type"))
    normalized_current = normalize_persona_tag(current_persona)

    if rectification and not rect_persona and rect_candidate_type == "observation":
        issues.append(ValidationIssue("observation_primary_persona_pending", "warn", "persona_tag", "当前对象仍处于 observation，主画像保持待定。"))
        return "", issues

    if rect_persona:
        if rect_persona in active_personas:
            return rect_persona, issues
        issues.append(ValidationIssue("nonstandard_rectified_persona", "warn", "persona_tag", f"纠偏结果主画像 `{rect_persona}` 未处于 active。"))

    if normalized_current in active_personas:
        return normalized_current, issues

    for persona_id, meta in active_personas.items():
        if meta["track_name"] == track_name:
            issues.append(ValidationIssue("fallback_primary_persona", "warn", "persona_tag", f"主画像回退到 `{persona_id}`。"))
            return persona_id, issues

    issues.append(ValidationIssue("persona_missing", "block", "persona_tag", "未能确定 active 主画像。"))
    return "", issues


def _determine_secondary_personas(
    *,
    current_value: object,
    rectification: dict[str, Any] | None,
    active_personas: dict[str, dict[str, str]],
    primary_persona: str,
) -> tuple[list[str], list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    secondary: list[str] = []
    rect_secondary = (rectification or {}).get("rectification", {}).get("suggested_secondary_persona_tags") or []
    if rect_secondary:
        for item in rect_secondary:
            normalized = normalize_persona_tag(item)
            if normalized == primary_persona or not normalized:
                continue
            if normalized not in active_personas:
                issues.append(ValidationIssue("nonstandard_secondary_persona", "warn", "secondary_persona_tags", f"次级画像 `{item}` 非 active，已忽略。"))
                continue
            if normalized not in secondary:
                secondary.append(normalized)
    else:
        for item in normalize_secondary_persona_tags(current_value):
            if item == primary_persona:
                continue
            if item not in active_personas:
                issues.append(ValidationIssue("nonstandard_secondary_persona", "warn", "secondary_persona_tags", f"次级画像 `{item}` 非 active，已忽略。"))
                continue
            if item not in secondary:
                secondary.append(item)
    return secondary, issues


def evaluate_enrich_readiness(*, candidate_type: str, minimum_fact_status: str, official_source_status: str, primary_persona: str) -> bool:
    if not primary_persona:
        return False
    if candidate_type != "formal_candidate":
        return False
    if minimum_fact_status == "fail":
        return False
    if official_source_status == "missing":
        return False
    return True


def _derive_minimum_fact_status(main_row: dict[str, Any]) -> str:
    issues = evaluate_minimum_fact_set(main_row)
    if any(item.severity == "block" for item in issues):
        return "fail"
    if issues:
        return "partial"
    return "pass"


def build_enrich_results(
    *,
    account_ids: list[str] | None = None,
    rectification_path: Path | None = None,
) -> list[EnrichResult]:
    main_rows, profile_rows, _review_rows, evidence_rows = load_state()
    active_personas = load_active_personas()
    assets = load_knowledge_assets()
    learning_queue = load_learning_queue()
    rectification_map = load_rectification_results(rectification_path) if rectification_path else {}

    profile_by_id = _build_lookup(profile_rows, "account_id")
    profile_by_name = _build_lookup(profile_rows, "account_canonical_name")
    main_by_id = _build_lookup(main_rows, "account_id")
    evidence_by_account = _build_grouped_lookup(evidence_rows, "account_id")

    wanted = set(account_ids or rectification_map.keys())
    if not wanted:
        wanted = {key for key in main_by_id if key}

    results: list[EnrichResult] = []
    for account_id in wanted:
        rectification = rectification_map.get(account_id)
        main_row = main_by_id.get(account_id, {})
        account_name = _clean(main_row.get("account_canonical_name"))
        profile_row = profile_by_id.get(account_id) or profile_by_name.get(account_name) or {}
        track_name = _track_name(main_row.get("primary_track") or profile_row.get("primary_track"))
        current_level = _clean(main_row.get("静态潜客记录成熟度") or profile_row.get("静态潜客记录成熟度"))
        current_persona = _clean(main_row.get("persona_tag") or profile_row.get("persona_tag"))
        primary_persona, persona_issues = _determine_primary_persona(
            current_persona=current_persona,
            rectification=rectification,
            track_name=track_name,
            active_personas=active_personas,
        )
        secondary_personas, secondary_issues = _determine_secondary_personas(
            current_value=main_row.get("secondary_persona_tags") or profile_row.get("secondary_persona_tags"),
            rectification=rectification,
            active_personas=active_personas,
            primary_persona=primary_persona,
        )

        current_refs = _split_multi(main_row.get("knowledge_asset_refs") or profile_row.get("knowledge_asset_refs"))
        knowledge_matches, strong_assets, strong_talks = match_related_knowledge_assets(
            track_name=track_name,
            primary_persona=primary_persona,
            secondary_personas=secondary_personas,
            current_refs=current_refs,
            assets=assets,
            persona_meta=active_personas,
        )
        queue_matches = find_matching_learning_queue(
            track_name=track_name,
            persona_tags=[primary_persona, *secondary_personas],
            queue_rows=learning_queue,
        )

        rect = (rectification or {}).get("rectification", {})
        rewrite = rect.get("rewrite_suggestion") or {}
        risk_flags = rect.get("risk_flags") or []
        suggested_maturity = _clean(rect.get("suggested_maturity")) or current_level
        validation_gap = _clean(rewrite.get("validation_gap") or rect.get("rectification_action") or profile_row.get("validation_gap") or main_row.get("validation_gap"))
        account_evidence = evidence_by_account.get(account_id, [])
        evidence_official_count = sum(
            1
            for row in account_evidence
            if _clean(row.get("evidence_strength") or row.get("strength")).upper() in {"A", "S"}
            or any(token in _clean(row.get("source_locator")).lower() for token in ("官网", "year", "annual", "ir", "investor", "cninfo", "公告"))
            or any(token in _clean(row.get("source_type")).lower() for token in ("official", "annual", "ir", "website"))
        )
        profile_official_count = 0
        try:
            profile_official_count = int(profile_row.get("official_source_count") or 0)
        except Exception:
            profile_official_count = 0
        official_source_status = "available" if max(profile_official_count, evidence_official_count) >= 1 else "missing"

        enriched_main_row = {
            "account_id": account_id,
            "account_canonical_name": account_name or _clean(profile_row.get("account_canonical_name")),
            "primary_track": track_name,
            "persona_tag": primary_persona,
            "secondary_persona_tags": ",".join(secondary_personas),
            "公司产品与服务概述": _clean(main_row.get("公司产品与服务概述") or profile_row.get("公司产品与服务概述")),
            "商业模式概述": _clean(main_row.get("商业模式概述") or profile_row.get("商业模式概述")),
            "admission_reason_summary": _clean(
                rewrite.get("admission_reason_summary")
                or main_row.get("admission_reason_summary")
                or main_row.get("一话入池理由")
                or profile_row.get("admission_reason_summary")
            ),
            "validation_gap": validation_gap,
            "review_status": _clean(
                rect.get("final_review_status")
                or profile_row.get("profile_status")
                or main_row.get("review_status")
                or ("active" if primary_persona else "pending_review")
            ),
        }
        minimum_fact_status = _derive_minimum_fact_status(enriched_main_row)
        live_validation = classify_l5_candidate(enriched_main_row, account_evidence)
        candidate_type = live_validation.candidate_type
        review_status = live_validation.review_status
        required_queue_type = live_validation.required_queue_type
        issues = [*persona_issues, *live_validation.issues]
        warnings = [*secondary_issues, *live_validation.warnings]
        ready = evaluate_enrich_readiness(
            candidate_type=candidate_type,
            minimum_fact_status=minimum_fact_status,
            official_source_status=official_source_status,
            primary_persona=primary_persona,
        )
        summary = f"{account_name or account_id} enrich 完成：主画像={primary_persona or '待定'}，次级画像={len(secondary_personas)}，知识资产={len(strong_assets)}，可进入 promote={ready}。"
        results.append(
            EnrichResult(
                account_id=account_id,
                account_canonical_name=account_name or _clean(profile_row.get("account_canonical_name")),
                primary_track=track_name,
                current_level=current_level,
                suggested_maturity=suggested_maturity,
                persona_tag=primary_persona,
                secondary_persona_tags=secondary_personas,
                knowledge_asset_refs=strong_assets,
                talk_track_refs=strong_talks,
                minimum_fact_status=minimum_fact_status,
                official_source_status=official_source_status,
                candidate_type=candidate_type,
                review_status=review_status,
                required_queue_type=required_queue_type,
                validation_gap=validation_gap,
                rewrite_suggestion={key: _clean(value) for key, value in rewrite.items()},
                matched_learning_queue_items=queue_matches,
                knowledge_matches=knowledge_matches,
                issues=issues,
                warnings=warnings,
                enrich_ready_for_promote=ready,
                summary=summary,
            )
        )
    results.sort(key=lambda item: item.account_id)
    return results
