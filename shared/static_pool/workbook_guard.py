from __future__ import annotations

import os
import time
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

try:
    import fcntl
except ImportError:  # pragma: no cover
    fcntl = None


LOCK_PATH = Path("/tmp/codex_static_pool_workbook_writeback.lock")


class WorkbookLockError(RuntimeError):
    pass


@contextmanager
def workbook_write_lock(*, timeout_seconds: float = 0.0, lock_path: Path = LOCK_PATH) -> Iterator[dict[str, object]]:
    if fcntl is None:
        raise WorkbookLockError("workbook lock is unavailable on this platform.")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    fd = os.open(str(lock_path), os.O_CREAT | os.O_RDWR, 0o644)
    start = time.monotonic()
    acquired = False
    try:
        while True:
            try:
                fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
                acquired = True
                break
            except BlockingIOError:
                waited = time.monotonic() - start
                if waited >= timeout_seconds:
                    raise WorkbookLockError(
                        f"workbook write lock is busy after {waited:.3f}s: {lock_path}"
                    ) from None
                time.sleep(0.05)
        waited = time.monotonic() - start
        yield {
            "lock_path": str(lock_path),
            "acquired": acquired,
            "wait_seconds": round(waited, 6),
        }
    finally:
        if acquired:
            try:
                fcntl.flock(fd, fcntl.LOCK_UN)
            except OSError:
                pass
        os.close(fd)


def check_workbook_integrity(paths: list[Path], *, deep_scan: bool = True) -> dict[str, object]:
    entries: list[dict[str, object]] = []
    passed = True
    for path in paths:
        item = {
            "path": str(path),
            "ok": True,
            "sheet_count": 0,
            "rows_scanned": 0,
            "error_type": "",
            "error_message": "",
        }
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            item["sheet_count"] = len(wb.sheetnames)
            if deep_scan:
                rows_scanned = 0
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    for _ in ws.iter_rows(min_row=1, values_only=True):
                        rows_scanned += 1
                item["rows_scanned"] = rows_scanned
        except Exception as exc:  # noqa: BLE001
            passed = False
            item["ok"] = False
            item["error_type"] = type(exc).__name__
            item["error_message"] = str(exc)
        entries.append(item)
    return {
        "ok": passed,
        "deep_scan": deep_scan,
        "files": entries,
    }

