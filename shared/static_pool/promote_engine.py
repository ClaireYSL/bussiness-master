from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .reporting import to_jsonable
from .validators import evaluate_promotion_gate, normalize_review_status


def load_sheet_rows(path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, object]]]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [{headers[i]: row[i] for i in range(len(headers))} for row in ws.iter_rows(min_row=2, values_only=True)]
    return headers, rows


def load_main_rows(path: Path, sheet_name: str) -> list[dict[str, object]]:
    _headers, rows = load_sheet_rows(path, sheet_name)
    return rows


def attach_account_ids(
    main_rows: list[dict[str, object]],
    profile_rows: list[dict[str, object]],
) -> list[dict[str, object]]:
    profile_map = {
        str(row.get("account_canonical_name") or "").strip(): str(row.get("account_id") or "").strip()
        for row in profile_rows
        if row.get("account_canonical_name")
    }
    enriched_rows = []
    for row in main_rows:
        normalized = dict(row)
        account_id = str(normalized.get("account_id") or "").strip()
        if not account_id:
            account_id = profile_map.get(str(normalized.get("account_canonical_name") or "").strip(), "")
            if account_id:
                normalized["account_id"] = account_id
        enriched_rows.append(normalized)
    return enriched_rows


def load_main_rows_with_fallback(
    main_path: Path,
    main_sheet: str,
    shared_path: Path,
    shared_sheet: str,
) -> list[dict[str, object]]:
    # Legacy helper for archived scripts and one-off recovery tasks.
    # Current execution-layer entrypoints should read the primary main table directly.
    try:
        return load_main_rows(main_path, main_sheet)
    except Exception:
        _headers, rows = load_sheet_rows(shared_path, shared_sheet)
        normalized_rows = []
        for row in rows:
            normalized_rows.append(
                {
                    "account_id": "",
                    "account_canonical_name": row.get("公司主体"),
                    "primary_track": row.get("主线"),
                    "persona_tag": row.get("业务形态画像"),
                    "secondary_persona_tags": row.get("辅助画像标签") or row.get("次级画像"),
                    "公司产品与服务概述": row.get("公司产品与服务概述"),
                    "商业模式概述": row.get("商业模式概述"),
                    "admission_reason_summary": row.get("一话入池理由"),
                    "validation_gap": row.get("待验证项"),
                    "信息扎实度": row.get("信息扎实度"),
                    "ICP匹配概率": row.get("ICP匹配概率"),
                    "静态潜客记录成熟度": row.get("静态潜客记录成熟度"),
                    "review_status": "pending_review",
                    "knowledge_asset_refs": row.get("主要知识资产引用"),
                    "talk_track_refs": row.get("主要切入话术引用"),
                }
            )
        return normalized_rows


def build_lookup(rows: list[dict[str, object]], key: str) -> dict[str, dict[str, object]]:
    result: dict[str, dict[str, object]] = {}
    for row in rows:
        value = str(row.get(key) or "").strip()
        if value:
            result[value] = row
    return result


def build_grouped_lookup(rows: list[dict[str, object]], key: str) -> dict[str, list[dict[str, object]]]:
    result: dict[str, list[dict[str, object]]] = {}
    for row in rows:
        value = str(row.get(key) or "").strip()
        if value:
            result.setdefault(value, []).append(row)
    return result


def select_main_rows(
    main_rows: list[dict[str, object]],
    *,
    account_ids: list[str] | None = None,
    from_level: str | None = None,
    track: str | None = None,
    limit: int | None = None,
) -> list[dict[str, object]]:
    selected: list[dict[str, object]] = []
    wanted = set(account_ids or [])
    for row in main_rows:
        account_id = str(row.get("account_id") or "").strip()
        if wanted and account_id not in wanted:
            continue
        if from_level and str(row.get("静态潜客记录成熟度") or "") != from_level:
            continue
        if track and str(row.get("primary_track") or "") != track:
            continue
        selected.append(row)
        if limit and len(selected) >= limit:
            break
    return selected


def evaluate_promotion_batch(
    main_rows: list[dict[str, object]],
    profile_rows: list[dict[str, object]],
    evidence_rows: list[dict[str, object]],
    queue_rows: list[dict[str, object]],
    *,
    account_ids: list[str] | None = None,
    from_level: str | None = None,
    target_level: str | None = None,
    track: str | None = None,
    limit: int | None = None,
) -> dict[str, Any]:
    profile_by_id = build_lookup(profile_rows, "account_id")
    profile_by_name = build_lookup(profile_rows, "account_canonical_name")
    evidence_by_account = build_grouped_lookup(evidence_rows, "account_id")
    queue_by_account = build_grouped_lookup(queue_rows, "account_id")

    selected_rows = select_main_rows(
        main_rows,
        account_ids=account_ids,
        from_level=from_level,
        track=track,
        limit=limit,
    )

    results = []
    summary = {"allow": 0, "warn": 0, "block": 0}
    for main_row in selected_rows:
        account_id = str(main_row.get("account_id") or "").strip()
        account_name = str(main_row.get("account_canonical_name") or "").strip()
        profile_row = profile_by_id.get(account_id) or profile_by_name.get(account_name) or {}
        main_row = dict(main_row)
        main_row["review_status"] = normalize_review_status(str(main_row.get("review_status") or ""))
        gate = evaluate_promotion_gate(
            main_row,
            profile_row,
            evidence_by_account.get(account_id, []),
            queue_by_account.get(account_id, []),
        )
        summary[gate.decision] += 1
        results.append(
            {
                "account_id": account_id,
                "account_canonical_name": account_name,
                "from_level": main_row.get("静态潜客记录成熟度"),
                "target_level": target_level or "",
                "primary_track": main_row.get("primary_track"),
                "persona_tag": main_row.get("persona_tag"),
                "secondary_persona_tags": main_row.get("secondary_persona_tags")
                or profile_row.get("secondary_persona_tags")
                or "",
                "knowledge_asset_refs": main_row.get("knowledge_asset_refs")
                or profile_row.get("knowledge_asset_refs")
                or "",
                "talk_track_refs": main_row.get("talk_track_refs") or profile_row.get("talk_track_refs") or "",
                "promotion_gate": to_jsonable(gate),
            }
        )
    return {
        "batch_summary": summary,
        "results": results,
    }
