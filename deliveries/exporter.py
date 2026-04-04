from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Any, Optional
import csv
import tempfile

from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from shared.db import get_session_factory
from shared.models import QueryJob, ResearchReport

DEFAULT_OUTPUT_DIR_NAME = "business_leads_delivery"
QUERY_EXPORT_TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "query_export_template.xlsx"
QUERY_EXPORT_TEMPLATE_DATA_START_ROW = 3
QUERY_EXPORT_HEADERS = [
    "公司名称",
    "行业",
    "客户行业",
    "客户省份",
    "规模",
    "城市",
    "成立年份",
    "注册资本",
    "法人",
    "官网",
    "统一信用代码",
    "客户产品/服务",
    "商业模式",
    "客户客群",
    "是否上市",
    "是否IPO",
    "客户收入规模",
    "客户利润",
    "客户营收增长情况",
    "已上线系统",
    "数字化项目动态",
    "数据现状",
    "BI切入机会",
    "相似客户",
    "主要竞品",
    "业务标签",
    "招聘代表岗位",
    "在招职位数",
    "客户招聘信息",
    "IT团队规模",
    "企业近一年重大事件",
    "联系人1姓名",
    "联系人1职位",
    "联系人1手机号",
    "联系人2姓名",
    "联系人2职位",
    "联系人2手机号",
]
QUERY_EXPORT_TEMPLATE_HEADERS = [
    "客户名称",
    "客户行业",
    "客户省份",
    "城市",
    "客户产品/服务",
    "商业模式",
    "客户客群",
    "是否上市",
    "是否IPO",
    "客户收入规模",
    "客户利润",
    "客户营收增长情况",
    "数字化系统现状",
    "相关岗位的招聘信息",
    "企业近一年重大事件",
    "联系人1姓名",
    "联系人1电话",
    "联系人2姓名",
    "联系人2电话",
    ".......",
    "加微信话术（25字以内）",
    "电话话术（详细）",
]


def _ensure_output_dir(output_dir: str | Path | None) -> Path:
    if output_dir is not None:
        path = Path(output_dir)
    else:
        path = Path(tempfile.gettempdir()) / DEFAULT_OUTPUT_DIR_NAME
    path.mkdir(parents=True, exist_ok=True)
    return path


def _write_csv_rows(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _normalize_cell_value(value: Any) -> str:
    if isinstance(value, list):
        return "、".join(str(entry).strip() for entry in value if str(entry).strip())
    return str(value or "").strip()


def _first_non_empty(*values: Any) -> str:
    for value in values:
        normalized = _normalize_cell_value(value)
        if normalized:
            return normalized
    return ""


def _join_labeled_lines(items: list[tuple[str, Any]]) -> str:
    lines: list[str] = []
    for label, value in items:
        normalized = _normalize_cell_value(value)
        if normalized:
            lines.append(f"{label}：{normalized}")
    return "\n".join(lines)


def _build_digital_status_text(export_fields: dict[str, Any]) -> str:
    return _join_labeled_lines(
        [
            ("已上线系统", export_fields.get("已上线系统")),
            ("数据现状", export_fields.get("数据现状")),
            ("项目动态", export_fields.get("数字化项目动态")),
            ("BI切入机会", export_fields.get("BI切入机会")),
        ]
    )


def _build_hiring_text(export_fields: dict[str, Any]) -> str:
    position_count = export_fields.get("在招职位数")
    if position_count in (None, ""):
        position_count = None
    return _join_labeled_lines(
        [
            ("代表岗位", export_fields.get("招聘代表岗位")),
            ("在招职位数", position_count),
            ("招聘信息", export_fields.get("客户招聘信息")),
        ]
    )


def _build_wechat_script_text(report: Optional[ResearchReport], export_fields: dict[str, Any]) -> str:
    first_contact_script = dict(report.first_contact_script or {}) if report is not None else {}
    return _first_non_empty(
        first_contact_script.get("wechat_add_line"),
        first_contact_script.get("opening_line"),
        export_fields.get("BI切入机会"),
    )


def _build_phone_script_text(report: Optional[ResearchReport], export_fields: dict[str, Any]) -> str:
    first_contact_script = dict(report.first_contact_script or {}) if report is not None else {}
    call_talking_points = [
        str(item).strip()
        for item in (first_contact_script.get("call_talking_points") or [])
        if str(item).strip()
    ]
    qualification_questions = [
        str(item).strip()
        for item in (first_contact_script.get("qualification_questions") or [])
        if str(item).strip()
    ]
    script_text = _join_labeled_lines(
        [
            ("切入时机", first_contact_script.get("why_now")),
            ("开场白", first_contact_script.get("opening_line")),
            ("沟通要点", "；".join(call_talking_points)),
            ("建议追问", "；".join(qualification_questions)),
            ("收尾推进", first_contact_script.get("closing_transition")),
        ]
    )
    if script_text:
        return script_text
    return _join_labeled_lines(
        [
            ("BI切入机会", export_fields.get("BI切入机会")),
            ("项目动态", export_fields.get("数字化项目动态")),
            ("重大事件", export_fields.get("企业近一年重大事件")),
        ]
    )


def _copy_template_row_style(worksheet: Any, source_row: int, target_row: int, max_column: int) -> None:
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
    for column_index in range(1, max_column + 1):
        source_cell = worksheet.cell(source_row, column_index)
        target_cell = worksheet.cell(target_row, column_index)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)


def _ensure_template_rows(worksheet: Any, required_last_row: int) -> None:
    if required_last_row <= worksheet.max_row:
        return
    source_row = QUERY_EXPORT_TEMPLATE_DATA_START_ROW
    current_max_row = worksheet.max_row
    additional_rows = required_last_row - current_max_row
    worksheet.insert_rows(current_max_row + 1, amount=additional_rows)
    for row_index in range(current_max_row + 1, required_last_row + 1):
        _copy_template_row_style(worksheet, source_row, row_index, worksheet.max_column)


def _write_xlsx_rows(path: Path, rows: list[dict[str, Any]]) -> None:
    if not QUERY_EXPORT_TEMPLATE_PATH.exists():
        raise ValueError("导出模板不存在，请检查 deliveries/templates/query_export_template.xlsx")

    workbook = load_workbook(QUERY_EXPORT_TEMPLATE_PATH)
    worksheet = workbook[workbook.sheetnames[0]]
    last_required_row = QUERY_EXPORT_TEMPLATE_DATA_START_ROW + max(len(rows), 1) - 1
    _ensure_template_rows(worksheet, last_required_row)

    for row_offset, row in enumerate(rows, start=QUERY_EXPORT_TEMPLATE_DATA_START_ROW):
        for column_index, header in enumerate(QUERY_EXPORT_TEMPLATE_HEADERS, start=1):
            worksheet.cell(row=row_offset, column=column_index, value=_normalize_cell_value(row.get(header)))

    workbook.save(path)


def _research_competitors_text(report: Optional[ResearchReport]) -> str:
    if report is None:
        return ""

    competitor_names: list[str] = []
    seen: set[str] = set()
    for competitor in report.competitors or []:
        if not isinstance(competitor, dict):
            continue
        name = str(competitor.get("name") or "").strip()
        if not name or name in seen:
            continue
        seen.add(name)
        competitor_names.append(name)

    return "、".join(competitor_names)


def _load_query_job_export_sources(
    job_id: int,
    *,
    session_factory: Optional[sessionmaker[Session]] = None,
) -> list[dict[str, Any]]:
    factory = session_factory or get_session_factory()

    with factory() as session:
        job = session.get(QueryJob, job_id)
        if job is None:
            raise ValueError("查询批次不存在")

        rows: list[dict[str, Any]] = []
        for item in job.items:
            if item.status != "completed" or item.result is None:
                continue
            company_snapshot = item.result.company_snapshot or {}
            export_fields = company_snapshot.get("export_fields") or {}
            if not export_fields:
                continue
            rows.append(
                {
                    "export_fields": export_fields,
                    "research_report": item.research_report,
                }
            )

    return rows


def _build_query_job_rows_from_sources(sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for payload in sources:
        export_fields = payload["export_fields"]
        report = payload["research_report"]
        row: dict[str, Any] = {}
        for header in QUERY_EXPORT_HEADERS:
            value = export_fields.get(header)
            if header == "主要竞品" and not value:
                value = _research_competitors_text(report)
            row[header] = _normalize_cell_value(value)
        rows.append(row)
    return rows


def _build_query_job_template_rows_from_sources(sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for payload in sources:
        export_fields = payload["export_fields"]
        report = payload["research_report"]
        row = {
            "客户名称": _first_non_empty(export_fields.get("公司名称")),
            "客户行业": _first_non_empty(export_fields.get("客户行业"), export_fields.get("行业")),
            "客户省份": _first_non_empty(export_fields.get("客户省份")),
            "城市": _first_non_empty(export_fields.get("城市")),
            "客户产品/服务": _first_non_empty(export_fields.get("客户产品/服务")),
            "商业模式": _first_non_empty(export_fields.get("商业模式")),
            "客户客群": _first_non_empty(export_fields.get("客户客群")),
            "是否上市": _first_non_empty(export_fields.get("是否上市")),
            "是否IPO": _first_non_empty(export_fields.get("是否IPO")),
            "客户收入规模": _first_non_empty(export_fields.get("客户收入规模")),
            "客户利润": _first_non_empty(export_fields.get("客户利润")),
            "客户营收增长情况": _first_non_empty(export_fields.get("客户营收增长情况")),
            "数字化系统现状": _build_digital_status_text(export_fields),
            "相关岗位的招聘信息": _build_hiring_text(export_fields),
            "企业近一年重大事件": _first_non_empty(export_fields.get("企业近一年重大事件")),
            "联系人1姓名": _first_non_empty(export_fields.get("联系人1姓名")),
            "联系人1电话": _first_non_empty(export_fields.get("联系人1手机号")),
            "联系人2姓名": _first_non_empty(export_fields.get("联系人2姓名")),
            "联系人2电话": _first_non_empty(export_fields.get("联系人2手机号")),
            ".......": "",
            "加微信话术（25字以内）": _build_wechat_script_text(report, export_fields),
            "电话话术（详细）": _build_phone_script_text(report, export_fields),
        }
        rows.append(row)
    return rows


def prepare_query_job_rows(
    job_id: int,
    *,
    session_factory: Optional[sessionmaker[Session]] = None,
) -> list[dict[str, Any]]:
    sources = _load_query_job_export_sources(job_id, session_factory=session_factory)
    return _build_query_job_rows_from_sources(sources)


def prepare_query_job_template_rows(
    job_id: int,
    *,
    session_factory: Optional[sessionmaker[Session]] = None,
) -> list[dict[str, Any]]:
    sources = _load_query_job_export_sources(job_id, session_factory=session_factory)
    return _build_query_job_template_rows_from_sources(sources)


def export_query_job(
    job_id: int,
    *,
    file_format: str,
    session_factory: Optional[sessionmaker[Session]] = None,
    output_dir: str | Path | None = None,
) -> str:
    normalized_format = file_format.lower()
    if normalized_format not in {"csv", "xlsx"}:
        raise ValueError("仅支持 csv 或 xlsx 导出")

    sources = _load_query_job_export_sources(job_id, session_factory=session_factory)
    rows = _build_query_job_rows_from_sources(sources)
    if not rows:
        raise ValueError("当前批次没有可导出的已完成结果")

    output_path = _ensure_output_dir(output_dir)
    file_path = output_path / f"query_job_{job_id}.{normalized_format}"
    if normalized_format == "csv":
        _write_csv_rows(file_path, QUERY_EXPORT_HEADERS, rows)
    else:
        template_rows = _build_query_job_template_rows_from_sources(sources)
        _write_xlsx_rows(file_path, template_rows)
    return str(file_path)
